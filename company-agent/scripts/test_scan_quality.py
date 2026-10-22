"""Offline tests for the data-quality fixes after the first Nordstrom run (2026-10-15):
junk-row filter, detail-page locations, Workday detection and dates, company merge on the final
domain, and dedupe's junk removal + re-keying + excel_upsert row removal.

Run: uv run python scripts/test_scan_quality.py   (no network; scratch copies only)
"""
from __future__ import annotations

import datetime as dt
import json
import os
import shutil
import sqlite3
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import ats  # noqa: E402
import crawl  # noqa: E402
import memory  # noqa: E402
from common import ROOT, looks_like_job  # noqa: E402

results: list[tuple[str, bool]] = []


def check(name: str, ok: bool) -> None:
    results.append((name, bool(ok)))
    print(("PASS " if ok else "FAIL ") + name)


# 1. junk rows
N = "https://careers.nordstrom.com"
for title, url in [
    ("29", f"{N}/jobs/page/29"), ("Open Jobs", f"{N}/jobs/page/2"), ("Go to last page", f"{N}/jobs/page/133"),
    ("Go to previous page", f"{N}/careers/page/145"), ("Feel Good About the Possibilities", f"{N}/careers/page/1"),
    ("Job Location", f"{N}/x-y/job/57ADC6B00B966A257FE3D10B1B38F388"), ("Next", f"{N}/jobs?page=3"),
    ("Careers", "https://acme.com/careers"), ("Search Jobs", "https://acme.com/jobs/search"), ("QA", "https://acme.com/jobs/qa"),
]:
    check(f"junk refused: {title!r} {url[-28:]}", not looks_like_job(title, url))
for title, url in [
    ("Software Engineer", "https://humancloud.dev/careers/software-engineer"),
    ("Restaurant Host - Marketplace Café - King of Prussia", f"{N}/restaurant-host/job/FEDD2BA86FFAF24A8F1274081524EB29"),
    ("Sr Financial Analyst (3 open roles total)", f"{N}/sr-financial-analyst/job/ABC123"),
    ("Analyst, Banking Alliances", "https://job-boards.greenhouse.io/razorpaysoftwareprivatelimited/jobs/4740728005"),
]:
    check(f"real job kept: {title[:40]!r}", looks_like_job(title, url))

listing = b"""<html><body><h1>Open Jobs</h1>
<div><a href="/jobs/page/2">2</a><a href="/jobs/page/3">Open Jobs</a></div>
<div><h3>Warehouse Associate</h3><p>Dubuque, IA</p><a href="/warehouse-associate/job/AAA111">View</a></div>
<div><a href="/beauty-sales/job/BBB222">Beauty Sales - Lancome</a></div></body></html>"""
cards = crawl.listing_cards(crawl.parse_html(listing), f"{N}/jobs/", "nordstrom.com")
titles = sorted(c["title"] for c in cards.values())
check("listing_cards keeps 2 real cards, drops pagination", titles == ["Beauty Sales - Lancome", "Warehouse Associate"])

from common import normalize_url  # noqa: E402
check("normalize_url: percent-encoded and raw path are one key",
      normalize_url(f"{N}/host-caf%C3%A9-kop/job/FEDD") == normalize_url(f"{N}/host-café-kop/job/FEDD"))

# 2. location from the detail page
detail_ld = b"""<html><head><script type="application/ld+json">{"@context":"https://schema.org","@type":"JobPosting",
"title":"Beauty Sales","jobLocation":{"@type":"Place","address":{"addressLocality":"Columbia","addressRegion":"MD",
"addressCountry":{"@type":"Country","name":"US"}}}}</script></head><body><h1>Beauty Sales</h1></body></html>"""
check("detail page: JSON-LD jobLocation first", crawl.page_location(crawl.parse_html(detail_ld)) == "Columbia, MD, US")
detail_txt = b"<html><body><h1>Asset Protection - Coordinator</h1><p>Farmington, CT</p><p>Full time</p></body></html>"
check("detail page: page text when no JSON-LD", crawl.page_location(crawl.parse_html(detail_txt)) == "Farmington, CT")

# 3. Workday
d = ats.detect(["https://nordstrom.wd501.myworkdayjobs.com/nordstrom_careers/job/Seattle-WA/X_R-887467",
                "https://acme.wd5.myworkdayjobs.com/en-US/External_Careers", "https://x.wd1.myworkdayjobs.com/wday/cxs/x/y/jobs"])
check("workday detect: tenant.wdN/site, locale skipped, API path ignored",
      [(x["ats"], x["slug"]) for x in d] == [("workday", "nordstrom.wd501/nordstrom_careers"), ("workday", "acme.wd5/External_Careers")])
check("workday board url", ats.board_url("workday", "nordstrom.wd501/nordstrom_careers") == "https://nordstrom.wd501.myworkdayjobs.com/nordstrom_careers")
t = dt.date(2026, 10, 15)
check("workday posted dates", [ats._workday_posted(x, t) for x in ("Posted Today", "Posted Yesterday", "Posted 3 Days Ago", "Posted 30+ Days Ago")]
      == ["2026-10-15", "2026-10-14", "2026-10-12", ""])
check("workday source type known", ats.SOURCE["workday"] == "ATS: Workday" and "workday" in ats.FETCHERS)

# 4 + 5 on scratch copies
tmp = Path(tempfile.mkdtemp(prefix="cq-"))
try:
    db = tmp / "m.db"
    con = memory.connect(db)
    a = memory.upsert_company(con, "Acme", "acme.dev")
    b = con.execute("INSERT INTO companies (name, domain, first_scanned) VALUES ('Acme','acme.co.in','2026-01-01')").lastrowid
    con.execute("INSERT INTO jobs (company_id, job_key, url, title, source_type, date_discovered, last_checked, status) VALUES (?,?,?,?,?,?,?,?)",
                (b, "https://acme.co.in/careers/x", "https://acme.co.in/careers/x", "Old Role", "Careers Page", "2026-01-01", "2026-01-01", "Open"))
    memory.add_route(con, b, "https://acme.dev/careers", "careers", 200)
    memory.add_route(con, a, "https://acme.dev/careers", "careers", 200)
    moved = memory.merge_company(con, "acme.co.in", "acme.dev")
    n_comp = con.execute("SELECT COUNT(*) FROM companies").fetchone()[0]
    check("merge: one company left, job history moved", n_comp == 1 and moved["jobs"] == 1
          and con.execute("SELECT company_id FROM jobs").fetchone()[0] == a)
    check("merge: alias resolves to the final-domain company", memory.company_by_domain(con, "acme.co.in")["id"] == a)
    check("merge: first_scanned kept from the older record", con.execute("SELECT first_scanned FROM companies").fetchone()[0] == "2026-01-01")
    check("upsert on alias domain updates, never duplicates",
          memory.upsert_company(con, "Acme", "acme.co.in") == a and con.execute("SELECT COUNT(*) FROM companies").fetchone()[0] == 1)
    con.close()

    # dedupe + excel_upsert end to end, in a scratch copy of the project (never the real data)
    proj = tmp / "proj"
    (proj / "scripts").mkdir(parents=True)
    for f in (ROOT / "scripts").glob("*.py"):
        shutil.copy(f, proj / "scripts" / f.name)
    shutil.copytree(ROOT / "db", proj / "db")
    (proj / "runs").mkdir()
    env = {**os.environ, "PYTHONDONTWRITEBYTECODE": "1"}
    py = sys.executable

    def run(*args):
        return subprocess.run([py, *args], cwd=proj, env=env, capture_output=True, text=True)

    con = memory.connect(proj / "memory.db")
    memory.upsert_company(con, "Shop", "shop.com")
    con.close()
    today = dt.date.today().isoformat()

    def row(title, url, src, loc=""):
        return {"Company Name": "Shop", "Official Website": "https://shop.com/", "Career Page": "https://careers.shop.com/",
                "Job Title": title, "Location / Work Mode": loc, "Source Type": src, "Job / Post URL": url,
                "Date Discovered": today, "Last Checked": today, "Status": "Open"}

    old = [row("2", "https://careers.shop.com/jobs/page/2", "Careers Page"),
           row("Open Jobs", "https://careers.shop.com/jobs/page/3", "Careers Page"),
           row("Stock Associate - Mall A", "https://careers.shop.com/stock-a/job/AAA", "Careers Page"),
           row("Cashier - Mall B", "https://careers.shop.com/cashier-b/job/BBB", "Careers Page")]
    # an earlier run's junk is stored directly, as the first Nordstrom run did
    con = sqlite3.connect(proj / "memory.db")
    for r in old:
        con.execute("INSERT INTO jobs (company_id, job_key, url, title, location, source_type, date_discovered, last_checked, status) VALUES (1,?,?,?,?,?,?,?,?)",
                    (r["Job / Post URL"].replace("https://careers.shop.com", "https://careers.shop.com"), r["Job / Post URL"], r["Job Title"], "", r["Source Type"], "2026-10-01", "2026-10-01", "Open"))
    con.commit()
    con.close()
    # build the "before" sheet from the raw old rows (junk included) like the first run did
    (proj / "runs" / "sheet-before.json").write_text(json.dumps({"meta": {"company": "Shop", "domain": "shop.com"}, "rows": old}))
    run("scripts/excel_upsert.py", "--rows", "runs/sheet-before.json")
    feed = [row("Stock Associate - Mall A", "https://shop.wd1.myworkdayjobs.com/ext/job/A_R-1", "ATS: Workday", "City A, ST"),
            row("Shoe Sales - Mall C", "https://shop.wd1.myworkdayjobs.com/ext/job/C_R-3", "ATS: Workday", "City C, ST")]
    (proj / "runs" / "feed.json").write_text(json.dumps({"meta": {"company": "Shop", "domain": "shop.com",
                                                                   "complete_sources": ["ATS: Workday", "Careers Page"]}, "rows": feed}))
    out = run("scripts/dedupe.py", "runs/feed.json")
    final = sorted((proj / "runs").glob("shop-*.final.rows.json"))[-1]
    payload = json.loads(final.read_text())
    by_title = {r["Job Title"]: r for r in payload["rows"]}
    check("dedupe: junk rows removed from memory and listed", sorted(x["title"] for x in payload["meta"]["removed"]) == ["2", "Open Jobs"])
    check("dedupe: careers row re-keyed to the ATS URL, Date Discovered kept",
          by_title["Stock Associate - Mall A"]["Job / Post URL"].endswith("A_R-1") and by_title["Stock Associate - Mall A"]["Date Discovered"] == "2026-10-01")
    # Two-miss rule (improvement plan item 4): one complete scan without the job leaves it Open.
    check("dedupe: real job missing once from the feed stays Open (missing once), not deleted",
          by_title["Cashier - Mall B"]["Status"] == "Open" and payload["meta"]["missing_once_titles"] == ["Cashier - Mall B"])
    check("dedupe: no junk titles in the final file", not any(t in by_title for t in ("2", "Open Jobs")))
    up = run("scripts/excel_upsert.py", "--rows", str(final.relative_to(proj)))
    from openpyxl import load_workbook
    ws = load_workbook(proj / "Company_Hiring_Intelligence.xlsx")["Shop"]
    sheet = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    titles = sorted(r[3] for r in sheet if r[3])
    check("excel: junk rows removed, re-keyed row updated in place, no duplicates",
          titles == ["Cashier - Mall B", "Shoe Sales - Mall C", "Stock Associate - Mall A"])
    links_ok = all((c.hyperlink is None) or c.hyperlink.ref == c.coordinate for r in ws.iter_rows(min_row=2) for c in r)
    check("excel: hyperlinks point at their own cells after row removal", links_ok)
    con = memory.connect(proj / "memory.db")  # the rerun is on a later day (same-day reruns count one miss)
    con.execute("UPDATE jobs SET last_miss_date=date(last_miss_date, '-1 day') WHERE last_miss_date IS NOT NULL")
    con.commit()
    con.close()
    again = run("scripts/dedupe.py", "runs/feed.json")
    check("rerun: 0 new rows", "new: 0," in again.stdout)
    final2 = json.loads(sorted((proj / "runs").glob("shop-*.final.rows.json"))[-1].read_text())
    check("rerun: job missing from a second complete scan is Closed",
          {r["Job Title"]: r["Status"] for r in final2["rows"]}["Cashier - Mall B"] == "Closed")
    if not all(ok for _, ok in results[-8:]):
        print(out.stdout, out.stderr, up.stdout, up.stderr, again.stdout, again.stderr)
finally:
    shutil.rmtree(tmp, ignore_errors=True)

bad = [n for n, ok in results if not ok]
print(f"\n{len(results) - len(bad)}/{len(results)} passed")
sys.exit(1 if bad else 0)
