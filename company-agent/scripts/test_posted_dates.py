"""Offline tests for approximate Posted Dates (2026-10-15): the relative-age parser
(common.approx_posted), the validator's Posted Date Basis checks, the Stale margin, dedupe's
basis handling, the migration of a 14-column sheet to the new 15th column, and the aggregate-URL
dedupe key (LinkedIn / Indeed listing pages shared by many jobs).

Run: uv run python scripts/test_posted_dates.py   (no network; scratch copies only)
"""
from __future__ import annotations

import datetime as dt
import json
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import dedupe  # noqa: E402
import excel_upsert  # noqa: E402
import memory  # noqa: E402
import progress  # noqa: E402
from common import COLUMNS, EXACT_BASIS, approx_posted, basis_margin, fill_posted_date, is_aggregate_url, job_key, normalize_url  # noqa: E402
from validate_rows import validate_row  # noqa: E402

results: list[tuple[str, bool]] = []


def check(name: str, ok: bool) -> None:
    results.append((name, bool(ok)))
    print(("PASS " if ok else "FAIL ") + name)


tmp = Path(tempfile.mkdtemp(prefix="pd-"))
progress.STATE = tmp / ".current.json"
_connect = memory.connect
DB = tmp / "m.db"
memory.connect = lambda path=DB: _connect(path)
dedupe.RUNS_DIR = tmp / "runs"
dedupe.RUNS_DIR.mkdir()

SEEN = "2026-10-15"
try:
    # ---- parser: every phrase, seen 2026-10-15 ---------------------------------------------------
    cases = [
        ("just now", "2026-10-15", 0, ""),
        ("today", "2026-10-15", 0, ""),
        ("yesterday", "2026-10-14", 0, ""),
        ("Posted yesterday", "2026-10-14", 0, ""),
        ("5 minutes ago", "2026-10-15", 0, ""),
        ("1 min ago", "2026-10-15", 0, ""),
        ("3 hours ago", "2026-10-15", 0, ""),
        ("an hour ago", "2026-10-15", 0, ""),
        ("30 hours ago", "2026-10-14", 0, ""),
        ("4 days ago", "2026-10-11", 0, ""),
        ("1 day", "2026-10-14", 0, ""),
        ("1 day ago", "2026-10-14", 0, ""),
        ("a week ago", "2026-10-08", 7, ""),
        ("3 weeks ago", "2026-09-24", 7, ""),
        ("1 month ago", "2026-09-15", 30, " (month precision)"),
        ("2 months ago", "2026-08-16", 30, " (month precision)"),
        ("30+ days ago", "2026-09-15", 0, " (upper bound: posted on or before)"),
        ("Reposted 2 weeks ago", "2026-10-01", 7, ""),
        ("Reposted 3 days ago", "2026-10-12", 0, ""),
    ]
    for text, date, margin, tag in cases:
        hit = approx_posted(text, SEEN)
        phrase = text[7:] if text.lower().startswith("posted ") else text
        check(f"parser: {text!r} -> {date}",
              hit is not None and hit["date"] == date and hit["margin"] == margin
              and hit["basis"] == f'approx: "{phrase}" seen {SEEN}{tag}')
    hit = approx_posted("Incident Engineer … Bengaluru 3 weeks ago 58 applicants", SEEN)
    check("parser: age inside snippet text is found and quoted alone",
          hit and hit["date"] == "2026-09-24" and hit["basis"] == f'approx: "3 weeks ago" seen {SEEN}')
    hit = approx_posted("1 month ago 58 applicants", SEEN)
    check("parser: '1 month ago 58 applicants' -> month precision", hit and hit["date"] == "2026-09-15" and "(month precision)" in hit["basis"])
    for text in ["", None, "gibberish", "work 5 days a week", "Apply today", "2026-10-01", "Senior Engineer, 3 years experience",
                 "x" * 500]:
        check(f"parser: unparseable {str(text)[:30]!r} -> None", approx_posted(text, SEEN) is None)
    check("parser: bad seen date -> None", approx_posted("3 weeks ago", "not-a-date") is None)

    # ---- fill_posted_date ------------------------------------------------------------------------
    r = fill_posted_date({"Posted Date": "", "_posted_age": "3 weeks ago", "Last Checked": SEEN})
    check("fill: age -> estimate + approx basis", r["Posted Date"] == "2026-09-24" and r["Posted Date Basis"].startswith('approx: "3 weeks ago"'))
    r = fill_posted_date({"Posted Date": "2026-10-01"})
    check("fill: real date -> exact (from page)", r["Posted Date Basis"] == EXACT_BASIS)
    r = fill_posted_date({"Posted Date": "", "_posted_age": "soon", "Last Checked": SEEN})
    check("fill: unreadable age -> both blank", r["Posted Date"] == "" and r["Posted Date Basis"] == "")
    r = fill_posted_date({"Posted Date": ""})
    check("fill: no age -> both blank", r["Posted Date"] == "" and r["Posted Date Basis"] == "")

    # ---- validator -------------------------------------------------------------------------------
    base = {"Company Name": "Acme", "Job Title": "Engineer", "Job / Post URL": "https://www.linkedin.com/jobs/view/1",
            "Source Type": "LinkedIn (public index)", "Date Discovered": SEEN, "Status": "Unverified"}
    good = fill_posted_date({**base, "_posted_age": "3 weeks ago", "Last Checked": SEEN})
    check("validator: script-made estimate passes", validate_row(good, 1) == [])
    check("validator: exact basis passes", validate_row({**base, "Posted Date": "2026-10-01", "Posted Date Basis": EXACT_BASIS}, 1) == [])
    check("validator: old 14-column row (no basis) passes", validate_row({**base, "Posted Date": "2026-10-01"}, 1) == [])
    check("validator: raw _posted_age without a date passes", validate_row({**base, "_posted_age": "3 weeks ago"}, 1) == [])
    check("validator: invented estimate (date not from the basis) is rejected",
          any("does not match its basis" in e for e in validate_row({**good, "Posted Date": "2026-09-01"}, 1)))
    check("validator: free-text basis is rejected",
          any("Posted Date Basis" in e for e in validate_row({**base, "Posted Date": "2026-09-01", "Posted Date Basis": "about a month"}, 1)))
    check("validator: basis without a date is rejected",
          any("is blank" in e for e in validate_row({**base, "Posted Date Basis": EXACT_BASIS}, 1)))
    check("validator: basis quoting an unreadable age is rejected",
          bool(validate_row({**base, "Posted Date": SEEN, "Posted Date Basis": f'approx: "soon" seen {SEEN}'}, 1)))
    check("validator: non-text _posted_age is rejected", bool(validate_row({**base, "_posted_age": 21}, 1)))
    # Typed columns were appended after it on 2026-10-16 (test_typed_fields.py covers them).
    check("columns: Posted Date Basis is column 15", COLUMNS.index("Posted Date Basis") == 14)

    # ---- Stale margin ----------------------------------------------------------------------------
    T = "2026-10-15"

    def ago(n):
        return (dt.date.fromisoformat(T) - dt.timedelta(days=n)).isoformat()

    wk = approx_posted("9 weeks ago", ago(5))  # posted 68 days before T
    mo = approx_posted("2 months ago", ago(5))  # 65 days before T
    check("margin: exact/blank 0, weeks 7, months 30, upper bound 0",
          basis_margin(EXACT_BASIS) == 0 and basis_margin("") == 0 and basis_margin(wk["basis"]) == 7
          and basis_margin(mo["basis"]) == 30 and basis_margin(approx_posted("30+ days ago", T)["basis"]) == 0)
    check("stale: exact date 61 days old is Stale", dedupe.is_stale(T, ago(61), EXACT_BASIS, T))
    check("stale: week estimate 63 days old is not Stale (needs > 67)",
          not dedupe.is_stale(T, approx_posted("9 weeks ago", T)["date"], approx_posted("9 weeks ago", T)["basis"], T))
    check("stale: week estimate 68 days old is Stale", dedupe.is_stale(T, wk["date"], wk["basis"], T))
    check("stale: month estimate 65 days old is not Stale (needs > 90)", not dedupe.is_stale(T, mo["date"], mo["basis"], T))
    m4 = approx_posted("3 months ago", ago(5))  # 95 days
    check("stale: month estimate 95 days old is Stale", dedupe.is_stale(T, m4["date"], m4["basis"], T))
    check("stale: not seen for 61 days is Stale whatever the date", dedupe.is_stale(ago(61), "", "", T))

    # ---- dedupe: basis stored, estimate never replaces an exact date ----------------------------
    today = dt.date.today().isoformat()
    con = memory.connect()
    memory.upsert_company(con, "Acme", "acme.com")
    con.close()

    def prow(n, **kw):
        return {"Company Name": "Acme", "Official Website": "https://acme.com/", "Career Page": "",
                "Job Title": f"Role {n}", "Location / Work Mode": "", "Source Type": "LinkedIn (public index)",
                "Job / Post URL": f"https://www.linkedin.com/jobs/view/{n}", "Date Discovered": today,
                "Last Checked": today, "Status": "Unverified", **kw}

    def run_dedupe(rows):
        f = tmp / "batch.rows.json"
        f.write_text(json.dumps({"meta": {"company": "Acme", "domain": "acme.com"}, "rows": rows}))
        dedupe.main([str(f)])
        final = sorted(dedupe.RUNS_DIR.glob("acme-*.final.rows.json"))[-1]
        return {r["Job Title"]: r for r in json.loads(final.read_text())["rows"]}, final

    out, final = run_dedupe([prow(1, **{"_posted_age": "3 weeks ago"}), prow(2, **{"Posted Date": "2026-10-01"}), prow(3)])
    exp = approx_posted("3 weeks ago", today)
    check("dedupe: estimate and basis reach the final file",
          out["Role 1"]["Posted Date"] == exp["date"] and out["Role 1"]["Posted Date Basis"] == exp["basis"])
    check("dedupe: real date marked exact; no date stays blank",
          out["Role 2"]["Posted Date Basis"] == EXACT_BASIS and out["Role 3"]["Posted Date"] == "" and out["Role 3"]["Posted Date Basis"] == "")
    check("dedupe: final file passes the validator", all(validate_row(r, i) == [] for i, r in enumerate(out.values(), 1)))
    out, _ = run_dedupe([prow(2, **{"_posted_age": "1 month ago"}), prow(1, **{"_posted_age": "4 weeks ago"})])
    check("dedupe: an estimate never replaces an exact date", out["Role 2"]["Posted Date"] == "2026-10-01" and out["Role 2"]["Posted Date Basis"] == EXACT_BASIS)
    exp = approx_posted("4 weeks ago", today)
    check("dedupe: a newer estimate replaces an older one", out["Role 1"]["Posted Date"] == exp["date"] and out["Role 1"]["Posted Date Basis"] == exp["basis"])
    out, _ = run_dedupe([prow(1)])
    check("dedupe: a re-find without an age keeps the stored estimate", out["Role 1"]["Posted Date"] == exp["date"])
    out, final = run_dedupe([prow(4, **{"_posted_age": "9 weeks ago"}), prow(5, **{"_posted_age": "10 weeks ago"})])
    check("dedupe: 9 weeks (63 d) not Stale, 10 weeks (70 d) Stale",
          out["Role 4"]["Status"] == "Unverified" and out["Role 5"]["Status"] == "Stale")

    # ---- migration: a 14-column sheet gets the 15th column on the next upsert -----------------
    wbp = tmp / "wb.xlsx"
    old_cols = COLUMNS[:14]
    wb = Workbook()
    ws = wb.active
    ws.title = "Acme"
    ws.append(old_cols)
    old = {c: "" for c in old_cols}
    old.update({"Company Name": "Acme", "Job Title": "Old Role", "Job / Post URL": "https://acme.com/jobs/old",
                "Source Type": "Careers Page", "Posted Date": "2026-09-01", "Date Discovered": "2026-09-02",
                "Last Checked": "2026-09-02", "Status": "Open"})
    ws.append([old[c] or None for c in old_cols])
    wb.save(wbp)
    con = memory.connect()
    con.execute("UPDATE companies SET sheet_name='Acme' WHERE domain='acme.com'")
    con.commit()
    con.close()
    backup = tmp / "wb.backup.xlsx"
    shutil.copy2(wbp, backup)
    argv = sys.argv
    sys.argv = ["excel_upsert.py", "--rows", str(final), "--workbook", str(wbp)]
    try:
        excel_upsert.main()
    finally:
        sys.argv = argv
    ws = load_workbook(wbp)["Acme"]
    header = [c.value for c in ws[1]]
    check("migration: header gains 'Posted Date Basis' as column 15, first 14 unchanged",
          header == COLUMNS and header[:14] == old_cols)
    oldrow = [c.value for c in ws[2]]
    check("migration: the old row keeps its 14 cells and a blank 15th",
          oldrow[:14] == [old[c] or None for c in old_cols] and oldrow[14] is None)
    rows = {ws.cell(row=r, column=4).value: ws.cell(row=r, column=15).value for r in range(3, ws.max_row + 1)}
    check("migration: new rows carry their basis in column 15",
          rows.get("Role 4", "").startswith('approx: "9 weeks ago"') and len(rows) == 5)
    check("migration: filter covers every column", ws.auto_filter.ref.startswith(f"A1:{get_column_letter(len(COLUMNS))}"))
    check("migration: backup copy untouched (14 columns)", load_workbook(backup)["Acme"].max_column == 14)

    # ---- aggregate LinkedIn / Indeed URLs (2026-10-15): key by company + title + location ------
    agg_cases = [
        ("https://in.linkedin.com/jobs/globex-jobs", True),
        ("https://www.linkedin.com/jobs/search?keywords=react&location=Pune", True),
        ("https://www.linkedin.com/company/globex/jobs", True),
        ("https://in.indeed.com/cmp/Globex/jobs", True),
        ("https://www.linkedin.com/jobs/view/3912345678", False),
        ("https://in.linkedin.com/jobs/view/react-developer-at-globex-3912345678/", False),
        ("https://www.linkedin.com/jobs/search?keywords=react&currentJobId=3912345678", False),
        ("https://in.indeed.com/viewjob?jk=abc123def", False),
        ("https://globex.com/careers", False),
        ("https://boards.greenhouse.io/globex/jobs/123", False),
    ]
    check("aggregate: listing URLs detected, per-job and non-index URLs not",
          all(is_aggregate_url(u) == want for u, want in agg_cases))
    AGG = "https://in.linkedin.com/jobs/globex-jobs"
    con = memory.connect()
    memory.upsert_company(con, "Globex", "globex.com")
    con.commit()
    con.close()

    def grow(title, loc="Pune, Maharashtra", url=AGG, **kw):
        return {"Company Name": "Globex", "Official Website": "https://globex.com/", "Career Page": "",
                "Job Title": title, "Location / Work Mode": loc, "Source Type": "LinkedIn (public index)",
                "Job / Post URL": url, "Date Discovered": today, "Last Checked": today, "Status": "Unverified", **kw}

    def run_globex(rows, complete=()):
        f = tmp / "globex.rows.json"
        f.write_text(json.dumps({"meta": {"company": "Globex", "domain": "globex.com", "complete_sources": list(complete)}, "rows": rows}))
        dedupe.main([str(f)])
        fin = sorted(dedupe.RUNS_DIR.glob("globex-*.final.rows.json"))[-1]
        return json.loads(fin.read_text()), fin

    def jobs_db():
        con = memory.connect()
        out = {r["title"]: dict(r) for r in con.execute("SELECT j.* FROM jobs j JOIN companies c ON c.id=j.company_id WHERE c.domain='globex.com'")}
        con.close()
        return out

    k1 = job_key(grow("Backend Engineer"))
    check("aggregate: job_key is company|title|location, punctuation and case normalised",
          k1 == "globex|backend engineer|pune maharashtra" and job_key(grow("BACKEND  engineer!", "pune - maharashtra")) == k1)
    check("aggregate: per-job URL keeps its normalised URL key",
          job_key(grow("X", url="https://www.linkedin.com/jobs/view/111/?trk=abc")) == normalize_url("https://www.linkedin.com/jobs/view/111/?trk=abc"))

    data, gfinal = run_globex([grow("Backend Engineer"), grow("Frontend Engineer"), grow("Data Analyst")])
    on_agg = [r for r in data["rows"] if r["Job / Post URL"] == AGG]
    check("aggregate: one URL with 3 different titles gives 3 rows (was 1)",
          sorted(r["Job Title"] for r in on_agg) == ["Backend Engineer", "Data Analyst", "Frontend Engineer"])
    data, _ = run_globex([grow("QA Engineer"), grow("QA Engineer"), grow("qa engineer", "Pune Maharashtra")])
    check("aggregate: the same title and location twice gives 1 row",
          sum(1 for r in data["rows"] if r["Job Title"].lower() == "qa engineer") == 1)
    data, _ = run_globex([grow("Support Engineer", "Bengaluru"), grow("Support Engineer", "Mumbai")])
    check("aggregate: same title in two locations stays 2 rows",
          sum(1 for r in data["rows"] if r["Job Title"] == "Support Engineer") == 2)
    data, _ = run_globex([grow("Platform Engineer", url="https://www.linkedin.com/jobs/view/222"),
                          grow("Platform Engineer", url="https://www.linkedin.com/jobs/view/333")])
    con = memory.connect()
    pj = [dict(r) for r in con.execute("SELECT job_key, url FROM jobs WHERE url LIKE '%/jobs/view/%' AND title='Platform Engineer'")]
    con.close()
    check("aggregate: per-job URLs are unchanged (URL key, two URLs stay two jobs)",
          len(pj) == 2 and all(r["job_key"] == normalize_url(r["url"]) for r in pj))
    data, _ = run_globex([{**grow("Backend Engineer"), "Company Name": "Globex Technologies Pvt Ltd"}])
    check("aggregate: a different company spelling from the scout still matches the stored job",
          sum(1 for r in data["rows"] if r["Job Title"] == "Backend Engineer") == 1)

    # Miss counter across runs: the key must match from run to run for "Closed after two complete misses".
    LI = "LinkedIn (public index)"
    trio = [grow("Miss A", "Delhi"), grow("Miss B", "Delhi"), grow("Miss C", "Delhi")]
    run_globex(trio, [LI])
    j = jobs_db()
    check("miss counter: a re-run with all three keeps 3 jobs at missing_count 0",
          run_globex(trio, [LI])[0] and all(jobs_db()[t]["missing_count"] == 0 for t in ("Miss A", "Miss B", "Miss C"))
          and len([t for t in jobs_db() if t.startswith("Miss")]) == 3)
    data, _ = run_globex(trio[:2], [LI])
    j = jobs_db()
    check("miss counter: one complete miss counts 1 for that job only, still Unverified",
          j["Miss C"]["missing_count"] == 1 and j["Miss C"]["status"] == "Unverified" and j["Miss A"]["missing_count"] == 0
          and [t for t in data["meta"]["missing_once_titles"] if t.startswith("Miss")] == ["Miss C"])
    run_globex(trio, [LI])
    check("miss counter: seen again resets to 0", jobs_db()["Miss C"]["missing_count"] == 0)
    run_globex(trio[:2], [LI])
    con = memory.connect()  # the next scan is on a later day (a same-day re-run does not count again)
    con.execute("UPDATE jobs SET last_miss_date=date(last_miss_date, '-1 day') WHERE last_miss_date IS NOT NULL")
    con.commit()
    con.close()
    data, _ = run_globex(trio[:2], [LI])
    j = jobs_db()
    check("miss counter: two complete misses in a row close it; the others stay open",
          j["Miss C"]["status"] == "Closed" and j["Miss A"]["status"] == "Unverified" and "Miss C" in data["meta"]["closed_titles"])

    # Old memory: one job stored under the aggregate URL as its key (before the fix) is re-keyed, not duplicated.
    con = memory.connect()
    gid = con.execute("SELECT id FROM companies WHERE domain='globex.com'").fetchone()["id"]
    OLD = "https://in.linkedin.com/jobs/globex-old"
    con.execute("INSERT INTO jobs (company_id, job_key, url, title, location, source_type, date_discovered, last_checked, status)"
                " VALUES (?,?,?,?,?,?,?,?,?)", (gid, normalize_url(OLD), OLD, "Legacy Role", "Pune", LI, "2026-10-01", "2026-10-01", "Unverified"))
    con.commit()
    con.close()
    data, gfinal = run_globex([grow("Legacy Role", "Pune", url=OLD), grow("Second Legacy", "Pune", url=OLD)])
    j = jobs_db()
    check("old memory: an aggregate-URL key is re-keyed (history kept, renamed_keys reported)",
          j["Legacy Role"]["job_key"] == "globex|legacy role|pune" and j["Legacy Role"]["date_discovered"] == "2026-10-01"
          and data["meta"]["renamed_keys"].get(normalize_url(OLD)) == "globex|legacy role|pune")
    check("old memory: the second job behind the old URL is added, not merged",
          sorted(r["Job Title"] for r in data["rows"] if r["Job / Post URL"] == OLD) == ["Legacy Role", "Second Legacy"])

    # Workbook: an old sheet holding one row for the aggregate URL; the upsert keeps one row per job, twice over.
    gwb = tmp / "globex.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Globex"
    ws.append(COLUMNS)
    oldrow = {c_: "" for c_ in COLUMNS}
    oldrow.update({"Company Name": "Globex", "Job Title": "Legacy Role", "Location / Work Mode": "Pune", "Job / Post URL": OLD,
                   "Source Type": LI, "Date Discovered": "2026-10-01", "Last Checked": "2026-10-01", "Status": "Unverified"})
    ws.append([oldrow[c_] or None for c_ in COLUMNS])
    wb.save(gwb)
    con = memory.connect()
    con.execute("UPDATE companies SET sheet_name='Globex' WHERE domain='globex.com'")
    con.commit()
    con.close()

    def upsert(fin):
        argv = sys.argv
        sys.argv = ["excel_upsert.py", "--rows", str(fin), "--workbook", str(gwb)]
        try:
            excel_upsert.main()
        finally:
            sys.argv = argv
        ws = load_workbook(gwb)["Globex"]
        return [[c.value for c in r] for r in ws.iter_rows(min_row=2)]

    first = upsert(gfinal)
    second = upsert(gfinal)
    ti, ui, li = COLUMNS.index("Job Title"), COLUMNS.index("Job / Post URL"), COLUMNS.index("Location / Work Mode")
    agg_jobs = [(r[ti], r[li]) for r in second if r[ui] == AGG]
    agg_rows = {t for t, _ in agg_jobs}
    check("workbook: one sheet row per job behind an aggregate URL (old row updated in place)",
          sorted(r[ti] for r in second if r[ui] == OLD) == ["Legacy Role", "Second Legacy"] and second[0][ti] == "Legacy Role"
          and len(agg_jobs) == len(set(agg_jobs)) == 9 and {"Backend Engineer", "Frontend Engineer", "Data Analyst"} <= agg_rows)
    check("workbook: a re-upsert adds nothing and drops nothing", len(first) == len(second) == len(json.loads(gfinal.read_text())["rows"]))

    # ---- memory.db migration: an old jobs table gains posted_date_basis -------------------------
    olddb = tmp / "old.db"
    import sqlite3
    c = sqlite3.connect(olddb)
    c.execute("CREATE TABLE jobs (id INTEGER PRIMARY KEY, company_id INTEGER, job_key TEXT, url TEXT, title TEXT, location TEXT,"
              " source_type TEXT, posted_by TEXT, posted_date TEXT, date_discovered TEXT, last_checked TEXT,"
              " description_summary TEXT, application_url TEXT, status TEXT, raw_json TEXT)")
    c.commit()
    c.close()
    c = _connect(olddb)
    cols = {r["name"] for r in c.execute("PRAGMA table_info(jobs)")}
    c.close()
    check("memory: old jobs table gains posted_date_basis (and missing_count)", {"posted_date_basis", "missing_count"} <= cols)
finally:
    shutil.rmtree(tmp, ignore_errors=True)

bad = [n for n, ok in results if not ok]
print(f"\n{len(results) - len(bad)}/{len(results)} passed")
sys.exit(1 if bad else 0)
