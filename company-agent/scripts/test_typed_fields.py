"""Offline tests for the typed job fields (2026-10-16): Work Mode, City, Country, Function, Seniority,
Salary Min / Max / Currency. Covers the parsers, the validator, the ATS feed hints (Greenhouse,
Lever, Ashby with compensation, Workable, SmartRecruiters), the JSON-LD JobPosting reader, dedupe into
memory.db (including an old database without the new columns) and a 15-column sheet migrating in place.

Run: uv run python scripts/test_typed_fields.py
"""
from __future__ import annotations

import datetime as dt
import json
import shutil
import sqlite3
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import ats  # noqa: E402
import crawl  # noqa: E402
import dedupe  # noqa: E402
import excel_upsert  # noqa: E402
import memory  # noqa: E402
import progress  # noqa: E402
from common import COLUMNS  # noqa: E402
from jobfields import (TYPED_COLUMNS, city_country, fill_typed, function_of, salary, seniority_of,  # noqa: E402
                       work_mode)
from validate_rows import validate_row  # noqa: E402

results: list[tuple[str, bool]] = []


def check(name: str, ok: bool) -> None:
    results.append((name, bool(ok)))
    print(("PASS " if ok else "FAIL ") + name)


tmp = Path(tempfile.mkdtemp(prefix="tf-"))
progress.STATE = tmp / ".current.json"
_connect = memory.connect
DB = tmp / "m.db"
memory.connect = lambda path=DB: _connect(path)
dedupe.RUNS_DIR = tmp / "runs"
dedupe.RUNS_DIR.mkdir()


class FakeResp:
    def __init__(self, data, status=200):
        self._data, self.status_code = data, status

    def json(self):
        return self._data

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(self.status_code)


def main() -> None:
    # ---- columns ---------------------------------------------------------------------------------
    check("columns: 15 old columns unchanged, 8 typed columns appended last",
          COLUMNS[14] == "Posted Date Basis" and COLUMNS[15:] == TYPED_COLUMNS and len(COLUMNS) == 23
          and TYPED_COLUMNS == ["Work Mode", "City", "Country", "Function", "Seniority", "Salary Min", "Salary Max", "Currency"])

    # ---- Work Mode: stated or Unknown, never from silence ----------------------------------------
    cases = [
        ((None, "Pune, India · Full-time"), "Unknown"),
        ((None, ""), "Unknown"),
        ((None, "Remote · Full-time"), "Remote"),
        ((None, "Bengaluru (Hybrid)"), "Hybrid"),
        ((None, "Pune · OnSite · FullTime"), "Onsite"),
        ((None, "Mumbai · On-site"), "Onsite"),
        ((None, "Remote or Hybrid, India"), "Hybrid"),
        ((None, "Remote; London (On-site)"), "Unknown"),
        (("remote", ""), "Remote"), (("on-site", ""), "Onsite"), (("OnSite", ""), "Onsite"),
        (("hybrid", ""), "Hybrid"), (("unspecified", "Pune"), "Unknown"), ((None, "Remote-first culture"), "Remote"),
        ((None, "Promote sales"), "Unknown"),
    ]
    bad = [(a, want, work_mode(*a)) for a, want in cases if work_mode(*a) != want]
    check(f"work mode: {len(cases)} cases (stated hint or word, else Unknown)", not bad)
    if bad:
        print("   ", bad)

    # ---- City / Country ---------------------------------------------------------------------------
    cc = [
        ((None, "Pune, India · Full-time"), ("Pune", "India")),
        ((None, "Bengaluru, Karnataka, IN"), ("Bengaluru", "India")),
        ((None, "San Francisco, CA"), ("San Francisco", "")),
        ((None, "Austin, TX, United States"), ("Austin", "United States")),
        ((None, "Remote"), ("", "")),
        ((None, "Remote - India"), ("", "India")),
        ((None, "India"), ("", "India")),
        ((None, "Bengaluru"), ("Bengaluru", "")),
        ((None, "London, UK; Pune, India"), ("London; Pune", "United Kingdom; India")),
        ((None, "Multiple Locations"), ("", "")),
        ((None, "Berlin, DE"), ("Berlin", "")),  # DE is also Delaware: ambiguous, so no country
        ((None, "Berlin, Germany"), ("Berlin", "Germany")),
        ((None, "Toronto, ON, CA"), ("Toronto", "Canada")),
        (([["Pune", "IN"]], "anything"), ("Pune", "India")),
        (([["Pune", "Atlantis"]], ""), ("Pune", "")),
        (([[None, None]], "Delhi, India"), ("Delhi", "India")),
    ]
    bad = [(a, want, city_country(*a)) for a, want in cc if city_country(*a) != want]
    check(f"city/country: {len(cc)} cases (structured first; a country only when named)", not bad)
    if bad:
        print("   ", bad)

    # ---- Function from the ATS department -------------------------------------------------------
    fn = [("Engineering", "Engineering"), ("Software Development", "Engineering"), ("Data Engineering", "Data"),
          ("Product Management", "Product"), ("Product Design", "Design"), ("Business Development", "Sales"),
          ("Sales Engineering", "Sales"), ("Marketing", "Marketing"), ("Customer Support", "Customer Success"),
          ("Finance & Accounting", "Finance"), ("People Operations", "People"), ("Talent Acquisition", "People"),
          ("Legal", "Legal"), ("IT", "IT"), ("Operations", "Operations"), ("Kitchen", "Other"), ("", ""), ("Learning and Development", "People"),
          ("GTM Offline First", "Sales"), ("Technical Account Management", "Customer Success"),
          (None, ""), (["", None], ""), (["Kitchen", "Platform"], "Engineering"), (["Kitchen"], "Other")]
    bad = [(d, want, function_of(d)) for d, want in fn if function_of(d) != want]
    check(f"function: {len(fn)} departments (unmapped -> Other, none -> blank)", not bad)
    if bad:
        print("   ", bad)

    # ---- Seniority from a level word in the title -----------------------------------------------
    sn = [("Software Engineer", ""), ("Senior Software Engineer", "Senior"), ("Sr. Backend Developer", "Senior"),
          ("Staff Engineer", "Staff"), ("Principal Data Scientist", "Staff"), ("Tech Lead, Payments", "Lead"),
          ("Engineering Manager", "Manager"), ("Product Manager", ""), ("Senior Product Manager", "Senior"),
          ("Account Manager", ""), ("Director of Engineering", "Director"), ("Head of Design", "Director"),
          ("VP Engineering", "VP"), ("Vice President, Sales", "VP"), ("Chief Technology Officer", "Exec"),
          ("CTO", "Exec"), ("Software Engineering Intern", "Intern"), ("Summer Internship - Data", "Intern"),
          ("Junior QA Engineer", "Entry"), ("Graduate Engineer Trainee", "Entry"), ("Mid-level Designer", "Mid"),
          ("Internal Auditor", ""), ("General Manager, India", "Exec"), ("Co-op Student", "Intern"),
          ("Lead Generation Specialist", "Lead"), ("", "")]
    bad = [(t, want, seniority_of(t)) for t, want in sn if seniority_of(t) != want]
    check(f"seniority: {len(sn)} titles (level word only; no word -> blank)", not bad)
    if bad:
        print("   ", bad)

    # ---- Salary: only a stated yearly salary ------------------------------------------------------
    check("salary: Ashby yearly range", salary({"min": 120000, "max": 150000.0, "currency": "usd", "interval": "1 YEAR"}) == (120000, 150000, "USD"))
    check("salary: JSON-LD single value fills both", salary({"min": "1,500,000", "max": None, "currency": "INR", "interval": "YEAR"}) == (1500000, 1500000, "INR"))
    check("salary: hourly / monthly / no currency / no hint -> blank",
          salary({"min": 50, "max": 70, "currency": "USD", "interval": "1 HOUR"}) == ("", "", "")
          and salary({"min": 5000, "max": 7000, "currency": "EUR", "interval": "MONTH"}) == ("", "", "")
          and salary({"min": 1, "max": 2, "currency": "", "interval": "YEAR"}) == ("", "", "")
          and salary(None) == ("", "", ""))
    check("salary: min above max or zero -> blank",
          salary({"min": 9, "max": 5, "currency": "USD", "interval": "YEAR"}) == ("", "", "")
          and salary({"min": 0, "max": 0, "currency": "USD", "interval": "YEAR"}) == ("", "", ""))

    # ---- validator --------------------------------------------------------------------------------
    today = dt.date.today().isoformat()
    base = {"Company Name": "Acme", "Official Website": "https://acme.com/", "Job Title": "Senior Engineer",
            "Location / Work Mode": "Pune, India", "Source Type": "ATS: Lever", "Job / Post URL": "https://jobs.lever.co/acme/1",
            "Date Discovered": today, "Last Checked": today, "Status": "Open"}
    good = fill_typed(dict(base))
    check("validator: filled row passes", validate_row(good, 1) == [] and good["Work Mode"] == "Unknown" and good["Seniority"] == "Senior")
    check("validator: invented Remote (no word, no hint) blocked", bool(validate_row({**good, "Work Mode": "Remote"}, 1)))
    check("validator: Remote backed by a script hint passes", validate_row({**good, "Work Mode": "Remote", "_work_mode": "remote"}, 1) == [])
    check("validator: Work Mode off the list blocked", bool(validate_row({**good, "Work Mode": "WFH"}, 1)))
    check("validator: Function off the list blocked", bool(validate_row({**good, "Function": "Ninjas"}, 1)))
    check("validator: Seniority the title does not name blocked", bool(validate_row({**good, "Seniority": "Director"}, 1)))
    check("validator: salary without currency blocked", bool(validate_row({**good, "Salary Min": 1, "Salary Max": 2}, 1)))
    check("validator: salary min above max blocked", bool(validate_row({**good, "Salary Min": 5, "Salary Max": 2, "Currency": "USD"}, 1)))
    check("validator: stated salary passes", validate_row({**good, "Salary Min": 100, "Salary Max": 200, "Currency": "USD"}, 1) == [])

    # ---- ATS feeds carry the hints ------------------------------------------------------------------
    calls = []
    feeds = {
        "boards-api.greenhouse.io": {"jobs": [{"title": "Senior Backend Engineer", "location": {"name": "Remote - India"},
                                               "absolute_url": "https://boards.greenhouse.io/acme/jobs/1", "first_published": "2026-10-01T00:00:00Z",
                                               "content": "Build things.", "departments": [{"name": "Engineering"}]}], "meta": {"total": 1}},
        "api.lever.co": [{"text": "Product Designer", "categories": {"location": "Pune, India", "team": "Design", "department": "Product & Design", "commitment": "Full-time"},
                          "workplaceType": "hybrid", "hostedUrl": "https://jobs.lever.co/acme/2", "applyUrl": "https://jobs.lever.co/acme/2/apply",
                          "createdAt": 1760000000000, "descriptionPlain": "Design things."}],
        "api.ashbyhq.com": {"jobs": [{"title": "Staff Data Scientist", "location": "San Francisco", "workplaceType": "OnSite", "employmentType": "FullTime",
                                      "department": "Data", "team": "ML", "jobUrl": "https://jobs.ashbyhq.com/acme/3", "publishedAt": "2026-10-02T00:00:00Z",
                                      "descriptionPlain": "Model things.",
                                      "address": {"postalAddress": {"addressLocality": "San Francisco", "addressRegion": "CA", "addressCountry": "USA"}},
                                      "compensation": {"summaryComponents": [
                                          {"compensationType": "EquityPercentage", "interval": "NONE", "minValue": 0.1, "maxValue": 0.2},
                                          {"compensationType": "Salary", "interval": "1 YEAR", "currencyCode": "USD", "minValue": 200000, "maxValue": 250000}]}},
                                     {"title": "Account Executive", "location": "London", "workplaceType": "Remote", "department": "Sales",
                                      "jobUrl": "https://jobs.ashbyhq.com/acme/4", "publishedAt": "2026-10-03", "descriptionPlain": "Sell."}]},
        "apply.workable.com": {"name": "Acme", "jobs": [{"title": "Junior Accountant", "city": "Mumbai", "state": "", "country": "India", "telecommuting": False,
                                                         "department": "Finance", "url": "https://apply.workable.com/acme/j/5", "published_on": "2026-10-04"}]},
        "api.smartrecruiters.com": {"totalFound": 1, "content": [{"id": "6", "name": "Support Lead", "releasedDate": "2026-10-05",
                                                                  "location": {"city": "Chennai", "country": "in", "remote": False, "hybrid": True},
                                                                  "department": {"label": "Customer Care"}, "function": {"label": "Customer Service"},
                                                                  "typeOfEmployment": {"label": "Full-time"}}]},
    }

    def fake_get(url, **kw):
        calls.append((url, kw.get("params")))
        host = url.split("/")[2]
        return FakeResp(feeds.get(host, {}), 200 if host in feeds else 404)

    real_get = ats._get
    ats._get = fake_get
    try:
        got = {}
        for name, slug in (("greenhouse", "acme"), ("lever", "acme"), ("ashby", "acme"), ("workable", "acme"), ("smartrecruiters", "acme")):
            rows, _ = ats.fetch(name, slug)
            got[name] = [fill_typed(r) for r in rows]
    finally:
        ats._get = real_get
    t = lambda r: tuple(r[c] for c in TYPED_COLUMNS)  # noqa: E731
    check("greenhouse: department -> Engineering, 'Remote - India' -> Remote / India, Senior",
          t(got["greenhouse"][0]) == ("Remote", "", "India", "Engineering", "Senior", "", "", ""))
    check("lever: workplaceType hybrid, team Design, Pune / India",
          t(got["lever"][0]) == ("Hybrid", "Pune", "India", "Design", "", "", "", ""))
    check("ashby: compensation requested (includeCompensation=true)",
          any("ashbyhq" in u and (p or {}).get("includeCompensation") == "true" for u, p in calls))
    check("ashby: OnSite, San Francisco / United States, Data, Staff, salary 200000-250000 USD",
          t(got["ashby"][0]) == ("Onsite", "San Francisco", "United States", "Data", "Staff", 200000, 250000, "USD"))
    check("ashby: no compensation -> salary blank; Remote; Sales; city from text",
          t(got["ashby"][1]) == ("Remote", "London", "", "Sales", "", "", "", ""))
    check("workable: telecommuting false stays Unknown (not Onsite); Finance; Entry",
          t(got["workable"][0]) == ("Unknown", "Mumbai", "India", "Finance", "Entry", "", "", ""))
    check("smartrecruiters: hybrid flag, Chennai / India (country code), Customer Success, Lead",
          t(got["smartrecruiters"][0]) == ("Hybrid", "Chennai", "India", "Customer Success", "Lead", "", "", ""))

    # ---- JSON-LD JobPosting on a careers page ------------------------------------------------------
    ld = {"@context": "https://schema.org", "@type": "JobPosting", "title": "Senior Data Engineer", "datePosted": "2026-10-06",
          "employmentType": "FULL_TIME", "jobLocationType": "TELECOMMUTE", "url": "https://acme.com/careers/sde",
          "jobLocation": {"@type": "Place", "address": {"addressLocality": "Pune", "addressCountry": {"name": "IN"}}},
          "baseSalary": {"@type": "MonetaryAmount", "currency": "INR",
                         "value": {"@type": "QuantitativeValue", "minValue": 2500000, "maxValue": 3500000, "unitText": "YEAR"}}}
    ld2 = {"@type": "JobPosting", "title": "Office Manager", "url": "https://acme.com/careers/om",
           "jobLocation": [{"address": {"addressLocality": "Austin", "addressRegion": "TX", "addressCountry": "US"}}],
           "baseSalary": {"currency": "USD", "value": {"value": 30, "unitText": "HOUR"}}}
    page = crawl.parse_html(f'<html><body><script type="application/ld+json">{json.dumps({"@graph": [ld, ld2]})}</script></body></html>'.encode())
    jl = [fill_typed(r) for r in crawl.jsonld_jobs(page, "https://acme.com/careers")]
    check("json-ld: datePosted kept as Posted Date", jl and jl[0]["Posted Date"] == "2026-10-06")
    check("json-ld: TELECOMMUTE -> Remote, Pune / India, Senior, yearly INR salary",
          t(jl[0]) == ("Remote", "Pune", "India", "", "Senior", 2500000, 3500000, "INR"))
    check("json-ld: no jobLocationType -> Unknown; hourly salary left blank; Manager",
          len(jl) == 2 and t(jl[1]) == ("Unknown", "Austin", "United States", "", "Manager", "", "", ""))

    # ---- dedupe into an OLD memory.db (no typed columns) -------------------------------------------
    con = sqlite3.connect(DB)
    con.executescript((Path(__file__).resolve().parent.parent / "db" / "schema.sql").read_text())
    for c, _ in memory.TYPED_DB_COLUMNS:  # the shape of a memory.db made before 2026-10-16
        con.execute(f"ALTER TABLE jobs DROP COLUMN {c}")
    con.commit()
    con.close()
    cols = {r[1] for r in sqlite3.connect(DB).execute("PRAGMA table_info(jobs)")}
    check("old db: built without the typed columns", "work_mode" not in cols and "posted_date_basis" in cols)
    con = memory.connect()
    cols = {r["name"] for r in con.execute("PRAGMA table_info(jobs)")}
    check("old db: memory.connect adds the 8 typed columns", all(c in cols for c, _ in memory.TYPED_DB_COLUMNS))
    memory.upsert_company(con, "Acme", "acme.com")
    cid = memory.company_by_domain(con, "acme.com")["id"]
    # a job stored before the typed columns existed (all NULL)
    con.execute("""INSERT INTO jobs (company_id, job_key, url, title, location, source_type, date_discovered, last_checked, status, raw_json)
                   VALUES (?,?,?,?,?,?,?,?,?,?)""", (cid, "https://acme.com/careers/old", "https://acme.com/careers/old", "Lead Engineer",
                                                    "Remote · Full-time", "Careers Page", "2026-09-01", "2026-09-01", "Open", "{}"))
    con.commit()
    con.close()

    def row(src, title, url, **kw):
        return {"Company Name": "Acme", "Official Website": "https://acme.com/", "Career Page": "https://acme.com/careers", "Job Title": title,
                "Location / Work Mode": "", "Source Type": src, "Job / Post URL": url, "Date Discovered": today, "Last Checked": today,
                "Status": "Open", **kw}

    batch = [row("ATS: Ashby", r["Job Title"], r["Job / Post URL"], **{k: v for k, v in r.items() if k.startswith("_") or k == "Location / Work Mode"})
             for r in got["ashby"]]
    batch.append(row("Job Board (public index)", "Senior Engineer (WFH)", "https://board.example/jobs/9", Status="Unverified",
                     **{"Location / Work Mode": "Work from home", "Work Mode": "Onsite"}))
    f = tmp / "batch.rows.json"
    f.write_text(json.dumps({"meta": {"company": "Acme", "domain": "acme.com"}, "rows": batch}))
    errs = sum((validate_row(r, i) for i, r in enumerate(batch, 1)), [])
    check("batch: an LLM-written Work Mode the text contradicts is refused by the validator", any("Work Mode" in e for e in errs))
    batch[-1].pop("Work Mode")
    f.write_text(json.dumps({"meta": {"company": "Acme", "domain": "acme.com"}, "rows": batch}))
    dedupe.main([str(f)])
    final = sorted(dedupe.RUNS_DIR.glob("acme-*.final.rows.json"))[-1]
    out = {r["Job Title"]: r for r in json.loads(final.read_text())["rows"]}
    check("dedupe: typed cells reach the final file",
          t(out["Staff Data Scientist"]) == ("Onsite", "San Francisco", "United States", "Data", "Staff", 200000, 250000, "USD"))
    check("dedupe: public row's Work Mode from its own text (work from home -> Remote)", out["Senior Engineer (WFH)"]["Work Mode"] == "Remote")
    check("dedupe: job stored before the migration gets typed cells from its title and location",
          t(out["Lead Engineer"]) == ("Remote", "", "", "", "Lead", "", "", ""))
    check("dedupe: final file passes the validator", all(validate_row(r, i) == [] for i, r in enumerate(out.values(), 1)))
    con = memory.connect()
    j = con.execute("SELECT work_mode, function, salary_min, currency FROM jobs WHERE title='Staff Data Scientist'").fetchone()
    con.close()
    check("memory.db: typed columns stored", tuple(j) == ("Onsite", "Data", 200000, "USD"))

    # ---- workbook: a 15-column sheet gains the 8 columns in place --------------------------------
    wbp = tmp / "wb.xlsx"
    old_cols = COLUMNS[:15]
    wb = Workbook()
    ws = wb.active
    ws.title = "Acme"
    ws.append(old_cols)
    old = {c: "" for c in old_cols}
    old.update({"Company Name": "Acme", "Job Title": "Lead Engineer", "Job / Post URL": "https://acme.com/careers/old",
                "Location / Work Mode": "Remote · Full-time", "Source Type": "Careers Page", "Date Discovered": "2026-09-01",
                "Last Checked": "2026-09-01", "Status": "Open"})
    ws.append([old[c] or None for c in old_cols])
    other = wb.create_sheet("Other Co")
    other.append(old_cols)
    wb.save(wbp)
    con = memory.connect()
    con.execute("UPDATE companies SET sheet_name='Acme' WHERE domain='acme.com'")
    con.commit()
    con.close()
    backup = tmp / "wb.backup.xlsx"
    shutil.copy2(wbp, backup)
    argv = sys.argv
    sys.argv = ["excel_upsert.py", "--rows", str(final), "--workbook", str(wbp)]
    try:
        excel_upsert.main()
    finally:
        sys.argv = argv
    wb = load_workbook(wbp)
    ws = wb["Acme"]
    header = [c.value for c in ws[1]]
    check("sheet: first 15 headers unchanged, 8 typed headers appended", header[:15] == old_cols and header[15:] == TYPED_COLUMNS)
    check("sheet: other sheets untouched", [c.value for c in wb["Other Co"][1]] == old_cols)
    vals = {ws.cell(row=r, column=4).value: [ws.cell(row=r, column=i).value for i in range(16, 24)] for r in range(2, ws.max_row + 1)}
    check("sheet: old row updated in place with typed cells", vals.get("Lead Engineer") == ["Remote", None, None, None, "Lead", None, None, None]
          and list(vals).count("Lead Engineer") == 1)
    check("sheet: salary written as numbers", vals.get("Staff Data Scientist")[5:] == [200000, 250000, "USD"])
    check("sheet: backup untouched (15 columns)", load_workbook(backup)["Acme"].max_column == 15)

    shutil.rmtree(tmp, ignore_errors=True)
    passed = sum(ok for _, ok in results)
    print(f"\n{passed}/{len(results)} passed")
    sys.exit(0 if passed == len(results) else 1)


if __name__ == "__main__":
    main()
