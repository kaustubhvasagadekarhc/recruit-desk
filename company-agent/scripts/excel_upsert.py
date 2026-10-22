"""Idempotent upsert of job rows into Company_Hiring_Intelligence.xlsx, one sheet per company.

A row is matched by its dedupe key (common.job_key: the normalised Job / Post URL; company + title +
location when there is no URL or the URL is an aggregate LinkedIn / Indeed listing page shared by
many jobs, so those jobs keep one sheet row each).
A match is updated in place; anything else is appended. Rerunning never adds duplicates.
Columns are only ever appended (Posted Date Basis, 2026-10-15): a sheet with the older 14 columns
gets the new header on its next upsert and its old rows keep their cells.
dedupe.py's meta.renamed_keys (a job that moved to a better source's URL) updates the old row in
place, and meta.removed_keys (junk rows: navigation links, not jobs) removes those rows.

CLI: uv run python scripts/excel_upsert.py --rows runs/<company>-<date>.final.rows.json
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import entities  # noqa: E402
import last_record  # noqa: E402
import memory  # noqa: E402
import progress  # noqa: E402
from common import COLUMNS, WORKBOOK_PATH, bare_domain, job_key, looks_like_job, normalize_url, safe_sheet_name  # noqa: E402
from validate_rows import load_rows, validate_payload  # noqa: E402

WIDTHS = {
    "Company Name": 22, "Official Website": 28, "Career Page": 34, "Job Title": 40,
    "Location / Work Mode": 30, "Source Type": 20, "Job / Post URL": 46, "Posted By": 18,
    "Posted Date": 12, "Date Discovered": 14, "Last Checked": 13, "Description Summary": 60,
    "Application URL": 46, "Status": 11, "Posted Date Basis": 40,
    "Work Mode": 11, "City": 18, "Country": 16, "Function": 16, "Seniority": 11,
    "Salary Min": 12, "Salary Max": 12, "Currency": 9,
}
URL_COLS = {"Official Website", "Career Page", "Job / Post URL", "Application URL"}
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
STATUS_FILL = {
    "Open": PatternFill("solid", fgColor="E2EFDA"),
    "Closed": PatternFill("solid", fgColor="EDEDED"),
    "Unverified": PatternFill("solid", fgColor="FFF2CC"),
    "Stale": PatternFill("solid", fgColor="FCE4D6"),
}


def prepare_sheet(ws) -> None:
    for i, col in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i, value=col)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS[col]
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def sheet_for(wb, company: str, domain: str) -> tuple:
    con = memory.connect()
    c = memory.company_by_domain(con, domain) if domain else None
    name = c["sheet_name"] if c and c["sheet_name"] else None
    if name and name in wb.sheetnames:
        return wb[name], con, c
    name = safe_sheet_name(company, set(wb.sheetnames) | set(entities.ENTITY_SHEETS))   # Companies and Jobs are reserved
    ws = wb.create_sheet(name)
    if c:
        con.execute("UPDATE companies SET sheet_name=? WHERE id=?", (name, c["id"]))
        con.commit()
    return ws, con, c


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--rows", required=True)
    ap.add_argument("--workbook", default=str(WORKBOOK_PATH))
    a = ap.parse_args()

    payload = json.loads(Path(a.rows).read_text())
    errs = validate_payload(payload)
    if errs:
        progress.fail_exit("refusing to write invalid rows:\n" + "\n".join(errs[:30]))
    rows = load_rows(payload)
    if not rows:
        progress.fail_exit("no rows to write")
    wbp = Path(a.workbook)
    lock = wbp.with_name("~$" + wbp.name)
    if lock.exists():
        progress.fail_exit(f"{wbp.name} looks open in Excel ({lock.name} exists); close it and rerun")

    meta = payload.get("meta", {}) if isinstance(payload, dict) else {}
    company = meta.get("company") or rows[0]["Company Name"]
    domain = bare_domain(meta.get("domain") or rows[0].get("Official Website") or "")

    if wbp.exists():
        wb = load_workbook(wbp)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    ws, _, _ = sheet_for(wb, company, domain)
    prepare_sheet(ws)

    index = {}
    junk_keys = set(meta.get("removed_keys") or [])
    dupe_rows: list[int] = []
    for r in range(2, ws.max_row + 1):
        vals = {col: ws.cell(row=r, column=i).value or "" for i, col in enumerate(COLUMNS, 1)}
        if any(vals.values()):
            if job_key(vals) in index:  # a second row for the same job: keep the first, drop this one
                dupe_rows.append(r)
                continue
            index[job_key(vals)] = r
            if vals["Source Type"] == "Careers Page" and not looks_like_job(vals["Job Title"], vals["Job / Post URL"]):
                junk_keys.add(job_key(vals))
    for old, new in (meta.get("renamed_keys") or {}).items():
        # keys stored before the URL-encoding fix; an aggregate-URL key (company|title|location) is not a URL
        old, new = (normalize_url(x) if x.startswith("http") else x for x in (old, new))
        if old in index and new not in index:
            index[new] = index.pop(old)
    progress.sub(f"{len(rows)} rows validated; writing sheet '{ws.title}'")
    updated = appended = 0
    newest = None   # (job key, record) of the last NEW job of this batch
    new_rows = []   # every NEW job of this batch, for the person's search history
    next_row = max([1, *index.values()]) + 1
    for row in rows:
        k = job_key(row)
        r = index.get(k)
        if r is None:
            r = next_row
            next_row += 1
            index[k] = r
            appended += 1
            newest = (k, {"job_title": row.get("Job Title"), "company": row.get("Company Name"), "location": row.get("Location / Work Mode"),
                          "status": row.get("Status"), "source_type": row.get("Source Type"), "posted_date": row.get("Posted Date"),
                          "listing_url": row.get("Job / Post URL"), "date_discovered": row.get("Date Discovered")})
            new_rows.append(newest)
        else:
            updated += 1
        for i, col in enumerate(COLUMNS, 1):
            v = row.get(col, "") or ""
            cell = ws.cell(row=r, column=i, value=v if v != "" else None)
            cell.alignment = Alignment(vertical="top", wrap_text=col == "Description Summary")
            if col in URL_COLS and v:
                cell.hyperlink = v
                cell.font = Font(color="0563C1", underline="single")
            elif col in URL_COLS:
                cell.hyperlink = None
            if col == "Status":
                cell.fill = STATUS_FILL.get(v, PatternFill())
    # Junk rows (proved not to be jobs) leave the sheet, bottom-up so row numbers stay valid.
    junk_keys |= {normalize_url(k) for k in junk_keys if k.startswith("http")}
    gone = sorted({index.pop(k) for k in junk_keys if k in index} | set(dupe_rows), reverse=True)
    for r in gone:
        ws.delete_rows(r)
    if gone:
        # openpyxl moves cells but not their hyperlink refs: point each link at its cell again.
        for line in ws.iter_rows(min_row=2):
            for cell in line:
                if cell.hyperlink:
                    cell.hyperlink.ref = cell.coordinate
        progress.sub(f"removed {len(gone)} junk row(s) from sheet '{ws.title}'")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(ws.max_row, 2)}"
    ent_con = memory.connect()
    n_comp, n_jobs = entities.refresh(wb, ent_con)
    ent_con.close()
    progress.sub(f"entity sheets: {n_comp} companies, {n_jobs} jobs")

    # Atomic save: a crash never leaves a half-written workbook.
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=wbp.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, wbp)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    if newest:   # a plain local file; Recruit Desk stores it as the last record while its person is signed in
        last_record.write(last_record.folder(wbp.parent) / last_record.FILE_NAME, *newest)
        last_record.append_rows(last_record.folder(wbp.parent), new_rows)
    print(f"sheet '{ws.title}': updated {updated}, appended {appended}, removed {len(gone)}, data rows {len(index)} -> {wbp.name}")
    progress.result(f"sheet '{ws.title}': {appended} inserted, {updated} updated" + (f", {len(gone)} junk rows removed" if gone else "") + f", {len(index)} rows in sheet")


if __name__ == "__main__":
    main()
