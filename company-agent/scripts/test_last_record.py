"""Offline tests for last_record.json: the last NEW job of an upsert, minimal fields, no network (2026-10-19). Temp copies only.

Run: uv run python scripts/test_last_record.py
"""
from __future__ import annotations

import hashlib
import json
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import excel_upsert  # noqa: E402
import last_record  # noqa: E402
import memory  # noqa: E402
from common import COLUMNS, job_key  # noqa: E402

results: list[tuple[str, bool]] = []


def check(name: str, ok: bool) -> None:
    results.append((name, bool(ok)))
    print(("PASS " if ok else "FAIL ") + name)


tmp = Path(tempfile.mkdtemp(prefix="lr-"))
orig_connect = memory.connect
memory.connect = lambda path=None: orig_connect(tmp / "memory.db")   # the real memory.db is never touched
wbp, lr = tmp / "wb.xlsx", tmp / "last_record.json"
_con = memory.connect()
memory.upsert_company(_con, "Acme", "acme.com", official_website="https://acme.com/", sheet_name="Acme")   # reruns must hit the same sheet
_con.commit()
_con.close()


def row(title, url, **kw):
    r = {c: "" for c in COLUMNS}
    r.update({"Company Name": "Acme", "Official Website": "https://acme.com/", "Career Page": "https://acme.com/careers", "Job Title": title,
              "Location / Work Mode": "Pune, India", "Source Type": "ATS: Greenhouse", "Job / Post URL": url, "Date Discovered": "2026-10-18",
              "Last Checked": "2026-10-18", "Status": "Open", "Description Summary": "Build things with a team, mail jobs@acme.com", "Application URL": url})
    r.update(kw)
    return r


def upsert(rows):
    f = tmp / "rows.json"
    f.write_text(json.dumps({"meta": {"company": "Acme", "domain": "acme.com"}, "rows": [{k: v for k, v in r.items() if v != ""} for r in rows]}))
    argv = sys.argv
    sys.argv = ["excel_upsert.py", "--rows", str(f), "--workbook", str(wbp)]
    try:
        excel_upsert.main()
    finally:
        sys.argv = argv


A, B, C = row("Backend Engineer", "https://boards.greenhouse.io/acme/jobs/1"), row("SRE", "https://boards.greenhouse.io/acme/jobs/2"), row("Designer", "https://boards.greenhouse.io/acme/jobs/3")
upsert([A, B])
d = json.loads(lr.read_text(encoding="utf-8"))
check("the last NEW job of the batch is written", d["agent"] == "company" and d["record"]["job_title"] == "SRE")
check("only the listed fields, no description or application link", set(d["record"]) <= set(last_record.FIELDS) and "jobs@acme.com" not in lr.read_text(encoding="utf-8")
      and "Build things" not in lr.read_text(encoding="utf-8"))
check("the key is only there as a hash", d["subject_hash"] == hashlib.sha256(f"company:{job_key(B)}".encode()).hexdigest()
      and "greenhouse.io/acme/jobs/2\"," not in d["subject_hash"])
check("fields carry the job's company, location, status and listing link", (d["record"]["company"], d["record"]["location"], d["record"]["status"], d["record"]["listing_url"])
      == ("Acme", "Pune, India", "Open", "https://boards.greenhouse.io/acme/jobs/2"))
first = lr.read_text(encoding="utf-8")
import time
time.sleep(1.1)   # a rewrite would carry a later created_at
upsert([A, B])
from openpyxl import load_workbook
check("a rerun is a pure update: still one sheet, two jobs", load_workbook(wbp).sheetnames == ["Acme", "Companies", "Jobs"] and load_workbook(wbp)["Acme"].max_row == 3)
check("a rerun that only updates leaves the file alone", lr.read_text(encoding="utf-8") == first)
upsert([A, B, C])
check("a later new job replaces it", json.loads(lr.read_text(encoding="utf-8"))["record"]["job_title"] == "Designer")
check("clean drops contact-looking values and unlisted fields", last_record.clean({"job_title": "Mail a@b.co", "company": "Acme", "email": "x@y.zz", "status": "Call +91 98765 43210"}) == {"company": "Acme"})
check("a missing or damaged file reads as none", last_record.read(tmp / "nope.json") is None)

# every NEW job also goes to new_rows.jsonl (the person's search history)
nr_path = tmp / "new_rows.jsonl"
nr = [json.loads(line) for line in nr_path.read_text(encoding="utf-8").splitlines()]
check("history rows: one line per NEW job, none for a rerun", len(nr) == 3 and {r["record"]["job_title"] for r in nr} == {"Backend Engineer", "SRE", "Designer"})
check("history rows: listed fields and a hashed key only", all(set(r["record"]) <= set(last_record.FIELDS) and len(r["subject_hash"]) == 64 and r["agent"] == "company" for r in nr)
      and "jobs@acme.com" not in nr_path.read_text(encoding="utf-8") and "Build things" not in nr_path.read_text(encoding="utf-8"))
check("history rows: contact-looking or empty records add nothing", last_record.append_rows(tmp, [("k", {"job_title": "Mail a@b.co"}), ("k2", {})]) == 0 and len(nr_path.read_text(encoding="utf-8").splitlines()) == 3)
big = tmp / "big"
big.mkdir()
(big / last_record.ROWS_FILE).write_text("x" * (last_record.MAX_ROWS_FILE_BYTES + 1))
check("history rows: the file stops growing past its limit", last_record.append_rows(big, [("k", {"job_title": "SRE"})]) == 0)

# Recruit Desk keeps the Excel file in its data folder but reads the hand-off files from the assistant's own folder: RECRUIT_DESK_HANDOFF_DIR
import os
hand = tmp / "handoff"
hand.mkdir()
check("without the variable the hand-off folder is next to the workbook", last_record.folder(tmp) == tmp)
os.environ["RECRUIT_DESK_HANDOFF_DIR"] = str(hand)
try:
    check("with the variable it is that folder", last_record.folder(tmp) == hand)
    wbp_data = tmp / "data-folder"
    wbp_data.mkdir()
    wbp_saved = wbp
    wbp = wbp_data / "wb.xlsx"
    upsert([row("Platform Engineer", "https://boards.greenhouse.io/acme/jobs/9")])
    check("the Excel file is where the data folder says", wbp.exists())
    check("the last record and the history rows are in the hand-off folder, not next to the Excel file",
          (hand / last_record.FILE_NAME).exists() and (hand / last_record.ROWS_FILE).exists()
          and not (wbp_data / last_record.FILE_NAME).exists() and not (wbp_data / last_record.ROWS_FILE).exists())
finally:
    del os.environ["RECRUIT_DESK_HANDOFF_DIR"]

memory.connect = orig_connect
shutil.rmtree(tmp, ignore_errors=True)
bad = [n for n, ok in results if not ok]
print(f"\n{len(results) - len(bad)}/{len(results)} passed")
sys.exit(1 if bad else 0)
