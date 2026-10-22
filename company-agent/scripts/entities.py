"""Entity sheets: one row per company and one row per job (2026-10-18, data-structure plan item 7).

uv run python scripts/entities.py [--dry-run]

Builder or human only (not in the agent's allow list). Backs up the workbook and memory.db into backups/ and adds
two export sheets built from the company sheets and memory.db:

- 'Companies': one row per company, keyed by domain (ATS, scan times, job counts by status).
- 'Jobs': one row per job across every company, keyed by the job key, with the company's domain as the foreign key.

Both are pure exports like the Evidence sheets of the other agents: the company sheets stay the place where rows
are written, and every excel_upsert rebuilds the two sheets from them. Nothing in a company sheet is changed.
Idempotent: a rerun gives the same sheets.
"""
from __future__ import annotations

import argparse
import json
import shutil
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
import progress  # noqa: E402
from common import COLUMNS, DB_PATH, ROOT, SCHEMA_PATH, WORKBOOK_PATH, bare_domain, job_key  # noqa: E402

COMPANIES_SHEET, JOBS_SHEET = "Companies", "Jobs"
ENTITY_SHEETS = (COMPANIES_SHEET, JOBS_SHEET)
BACKUPS = ROOT / "backups"

COMPANY_COLUMNS = [("Domain", 24), ("Company Name", 24), ("Official Website", 30), ("Career Page", 36), ("ATS", 14),
                   ("ATS Slug", 18), ("Sheet", 24), ("Jobs", 8), ("Open", 8), ("Closed", 8), ("Unverified", 10),
                   ("Stale", 8), ("First Scanned", 12), ("Last Scanned", 12)]
# The job columns that repeat the company are left out of Jobs; Domain points at the Companies row instead.
JOB_COLUMNS_FROM_SHEET = [c for c in COLUMNS if c not in ("Official Website", "Career Page")]
JOB_COLUMNS = [("Job Key", 46), ("Domain", 24), *[(c, 18) for c in JOB_COLUMNS_FROM_SHEET]]
JOB_WIDTHS = {"Company Name": 22, "Job Title": 40, "Location / Work Mode": 28, "Job / Post URL": 46, "Application URL": 40,
              "Description Summary": 50, "Posted Date Basis": 30}
URL_COLUMNS = {"Official Website", "Career Page", "Job / Post URL", "Application URL"}


def _company_sheets(wb) -> list:
    """Company sheets: current or older header (the first 14 columns), not an entity sheet."""
    out = []
    for ws in wb.worksheets:
        if ws.title in ENTITY_SHEETS:
            continue
        hdr = [c.value for c in ws[1]]
        if hdr[:14] == COLUMNS[:14]:
            out.append((ws, [h for h in hdr if h]))
    return out


def collect(wb, con) -> tuple[list[list], list[list]]:
    """(company rows, job rows) from every company sheet of wb and from memory.db."""
    by_sheet = {r["sheet_name"]: r for r in con.execute("SELECT * FROM companies WHERE sheet_name IS NOT NULL")}
    companies: dict[str, dict] = {r["domain"]: dict(r) for r in con.execute("SELECT * FROM companies")}
    jobs: list[list] = []
    counts: dict[str, dict[str, int]] = {}
    for ws, hdr in _company_sheets(wb):
        known = by_sheet.get(ws.title)
        for vals in ws.iter_rows(min_row=2, values_only=True):
            row = {h: vals[i] if i < len(vals) and vals[i] is not None else "" for i, h in enumerate(hdr)}
            if not row.get("Job Title") and not row.get("Job / Post URL"):
                continue
            domain = known["domain"] if known else bare_domain(str(row.get("Official Website") or ""))
            if domain and domain not in companies:
                companies[domain] = {"domain": domain, "name": row.get("Company Name"), "official_website": row.get("Official Website"),
                                     "career_page": row.get("Career Page"), "ats_type": None, "ats_slug": None,
                                     "sheet_name": ws.title, "first_scanned": None, "last_scanned": None}
            status = str(row.get("Status") or "Unverified")
            c = counts.setdefault(domain, {})
            c[status] = c.get(status, 0) + 1
            jobs.append([job_key(row), domain, *[row.get(h) for h in JOB_COLUMNS_FROM_SHEET]])
    comp_rows = []
    for domain, c in companies.items():
        n = counts.get(domain, {})
        comp_rows.append([domain, c.get("name"), c.get("official_website"), c.get("career_page"), c.get("ats_type"),
                          c.get("ats_slug"), c.get("sheet_name"), sum(n.values()), n.get("Open", 0), n.get("Closed", 0),
                          n.get("Unverified", 0), n.get("Stale", 0), (c.get("first_scanned") or "")[:10] or None,
                          (c.get("last_scanned") or "")[:10] or None])
    comp_rows.sort(key=lambda r: (-r[8], str(r[1]).lower()))
    status_i = 2 + JOB_COLUMNS_FROM_SHEET.index("Status")
    jobs.sort(key=lambda r: (r[1], {"Open": 0, "Unverified": 1, "Stale": 2, "Closed": 3}.get(r[status_i], 4), str(r[2 + JOB_COLUMNS_FROM_SHEET.index("Job Title")]).lower()))
    return comp_rows, jobs


def _protect(ws) -> None:
    """Mark an export sheet as read-only (no password; sorting and filtering stay allowed): it is rebuilt on every upsert, so an
    edit made in it would be lost."""
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.sheet_properties.tabColor = "808080"


def _write(wb, name: str, columns: list[tuple[str, int]], rows: list[list]) -> None:
    if name in wb.sheetnames:
        existing = [c.value for c in wb[name][1]]
        if existing != [h for h, _ in columns]:
            progress.fail_exit(f"a sheet named '{name}' already exists and is not an entity sheet (a company called '{name}'?); "
                               "rename it before the entity sheets can be built. Nothing was changed")
        del wb[name]
    ws = wb.create_sheet(name)
    ws.append([h for h, _ in columns])
    for i, (h, width) in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = JOB_WIDTHS.get(h, width) if name == JOBS_SHEET else width
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    url_cols = [i for i, (h, _) in enumerate(columns, start=1) if h in URL_COLUMNS]
    for r in rows:
        ws.append([v if v not in ("", None) else None for v in r])
        for i in url_cols:
            cell = ws.cell(ws.max_row, i)
            if cell.value and str(cell.value).startswith("http"):
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(columns)).column_letter}{max(ws.max_row, 1)}"
    _protect(ws)


def refresh(wb, con) -> tuple[int, int]:
    """Rebuild the Companies and Jobs sheets in wb (not saved), at the end of the workbook. Returns (companies, jobs)."""
    comp_rows, jobs = collect(wb, con)
    _write(wb, COMPANIES_SHEET, COMPANY_COLUMNS, comp_rows)
    _write(wb, JOBS_SHEET, JOB_COLUMNS, jobs)
    return len(comp_rows), len(jobs)


def migrate(wb_path: Path = WORKBOOK_PATH, db_path: Path = DB_PATH, dry_run: bool = False, stamp: str | None = None,
            backups: Path = BACKUPS) -> dict:
    if not wb_path.exists():
        progress.fail_exit(f"no workbook at {wb_path}")
    stamp = stamp or datetime.now().strftime("%Y%m%d-%H%M%S")
    made = []
    if not dry_run:
        backups.mkdir(exist_ok=True)
        for p in (wb_path, db_path):
            if p.exists():
                b = backups / f"{p.stem}.before-entities-{stamp}{p.suffix}"
                if b.exists():
                    progress.fail_exit(f"backup {b.name} already exists; move it aside first")
                shutil.copy2(p, b)
                made.append(b.name)
        progress.sub(f"backed up: {', '.join(made)}")
    wb = load_workbook(wb_path)
    con = sqlite3.connect(db_path)
    con.row_factory = sqlite3.Row
    con.executescript(SCHEMA_PATH.read_text())
    companies, jobs = refresh(wb, con)
    con.close()
    if not dry_run:
        wb.save(wb_path)
    return {"companies": companies, "jobs": jobs, "backups": made, "dry_run": dry_run}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true", help="print the counts, change nothing, take no backup")
    a = ap.parse_args()
    r = migrate(dry_run=a.dry_run)
    progress.result(f"{r['companies']} companies, {r['jobs']} jobs"
                    + (" (dry run, nothing saved)" if a.dry_run else f" · backups {', '.join(r['backups'])}"))
    print(json.dumps(r, indent=2))


if __name__ == "__main__":
    main()
