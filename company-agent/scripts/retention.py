"""Retention (security review step 9): find Stale (or Rejected) job rows older than N days.

  uv run python scripts/retention.py [--days 180]            # dry run (default): list only, change nothing
  uv run python scripts/retention.py [--days 180] --apply    # human only, from a terminal: delete them

A row's age is its Last Checked date (the last time any scan saw it). --apply removes the rows
from every company sheet of Company_Hiring_Intelligence.xlsx and the matching jobs from memory.db.
The agent's PreToolUse hook blocks --apply; only the human runs it.
"""
from __future__ import annotations

import argparse
import os
import sqlite3
import sys
import tempfile
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import COLUMNS, DB_PATH, WORKBOOK_PATH, job_key  # noqa: E402

PURGE_STATUSES = {"Stale", "Rejected"}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--days", type=int, default=180)
    ap.add_argument("--apply", action="store_true", help="delete the listed rows (human only)")
    ap.add_argument("--workbook", default=str(WORKBOOK_PATH))
    ap.add_argument("--db", default=str(DB_PATH))
    a = ap.parse_args()
    cutoff = (date.today() - timedelta(days=a.days)).isoformat()
    wb_path, db_path = Path(a.workbook), Path(a.db)
    if not wb_path.exists():
        print(f"no workbook at {wb_path}; nothing to do")
        return

    from openpyxl import load_workbook

    wb = load_workbook(wb_path)
    col = {c: i + 1 for i, c in enumerate(COLUMNS)}
    hits, undated = [], 0
    for ws in wb.worksheets:
        if [c.value for c in ws[1]][: len(COLUMNS)] != COLUMNS:
            continue
        for r in range(2, ws.max_row + 1):
            status = (ws.cell(r, col["Status"]).value or "").strip()
            if status not in PURGE_STATUSES:
                continue
            seen = str(ws.cell(r, col["Last Checked"]).value or ws.cell(r, col["Date Discovered"]).value or "")[:10]
            if not seen:
                undated += 1
                continue
            if seen < cutoff:
                vals = {c: ws.cell(r, i).value or "" for c, i in col.items()}
                hits.append((ws, r, job_key(vals), vals["Job Title"], status, seen))

    mode = "APPLY" if a.apply else "dry run"
    print(f"retention {mode}: {len(hits)} Stale/Rejected rows last checked before {cutoff} ({a.days} days)")
    for ws, r, _, title, status, seen in hits:
        print(f"  {ws.title} · row {r} · {title} · {status} · last checked {seen}")
    if undated:
        print(f"  {undated} Stale/Rejected rows have no date and were skipped")
    if not a.apply:
        print("nothing changed; a human adds --apply in a terminal to delete these rows")
        return
    if not hits:
        return

    for ws in {h[0] for h in hits}:
        for r in sorted((h[1] for h in hits if h[0] is ws), reverse=True):
            ws.delete_rows(r)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=wb_path.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, wb_path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    con = sqlite3.connect(db_path)
    for _, _, key, *_ in hits:
        con.execute("DELETE FROM jobs WHERE job_key=? AND status IN ('Stale','Rejected')", (key,))
    con.commit()
    print(f"deleted {len(hits)} rows from {wb_path.name} and their jobs from memory.db")


if __name__ == "__main__":
    main()
