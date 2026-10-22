"""Offline tests for the Companies and Jobs entity sheets (2026-10-18). Temp copies only.

Run: uv run python scripts/test_entities.py
"""
from __future__ import annotations

import io
import shutil
import sys
import tempfile
from contextlib import redirect_stdout
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import entities  # noqa: E402
import memory  # noqa: E402
from common import COLUMNS, job_key, safe_sheet_name  # noqa: E402

results: list[tuple[str, bool]] = []


def check(name: str, ok: bool) -> None:
    results.append((name, bool(ok)))
    print(("PASS " if ok else "FAIL ") + name)


tmp = Path(tempfile.mkdtemp(prefix="ent-"))
db, wbp = tmp / "memory.db", tmp / "wb.xlsx"
con = memory.connect(db)
memory.upsert_company(con, "Acme", "acme.com", official_website="https://acme.com", career_page="https://acme.com/careers",
                      ats_type="greenhouse", ats_slug="acme", sheet_name="Acme")
memory.upsert_company(con, "Beta", "beta.io", official_website="https://beta.io", sheet_name="Beta Inc")
con.commit()


def job(company, site, title, url, status, loc="Pune, India"):
    r = {c: "" for c in COLUMNS}
    r.update({"Company Name": company, "Official Website": site, "Job Title": title, "Location / Work Mode": loc,
              "Source Type": "ATS: Greenhouse", "Job / Post URL": url, "Date Discovered": "2026-10-15",
              "Last Checked": "2026-10-16", "Status": status})
    return r


wb = Workbook()
wb.remove(wb.active)
for name, rows in (("Acme", [job("Acme", "https://acme.com", "Backend Engineer", "https://boards.greenhouse.io/acme/jobs/1", "Open"),
                             job("Acme", "https://acme.com", "Designer", "https://boards.greenhouse.io/acme/jobs/2", "Closed")]),
                   ("Beta Inc", [job("Beta", "https://beta.io", "SRE", "https://beta.io/jobs/9", "Open")])):
    ws = wb.create_sheet(name)
    ws.append(COLUMNS)
    for r in rows:
        ws.append([r[c] or None for c in COLUMNS])
old = wb.create_sheet("Old Co")   # a 14-column sheet from before the typed columns
old.append(COLUMNS[:14])
r = job("Old", "https://old.example", "Analyst", "https://old.example/j/1", "Stale")
old.append([r[c] or None for c in COLUMNS[:14]])
wb.save(wbp)

wb = load_workbook(wbp)
n_comp, n_jobs = entities.refresh(wb, con)
check("3 companies and 4 jobs counted (a sheet with no memory row still gets a Companies row)", (n_comp, n_jobs) == (3, 4))
check("entity sheets sit after the company sheets", wb.sheetnames == ["Acme", "Beta Inc", "Old Co", "Companies", "Jobs"])
comp = {r[0]: r for r in wb["Companies"].iter_rows(min_row=2, values_only=True)}
check("Companies keyed by domain with counts by status", comp["acme.com"][7:12] == (2, 1, 1, 0, 0))
check("Companies carries ATS and slug", comp["acme.com"][4:6] == ("greenhouse", "acme"))
check("a sheet unknown to memory.db is keyed by its website domain", comp["old.example"][7] == 1 and comp["old.example"][11] == 1)
hdr = [c.value for c in wb["Jobs"][1]]
jobs = [dict(zip(hdr, r)) for r in wb["Jobs"].iter_rows(min_row=2, values_only=True)]
check("Jobs header: key, domain, then the sheet columns without website and career page",
      hdr[:3] == ["Job Key", "Domain", "Company Name"] and "Official Website" not in hdr and "Status" in hdr)
check("Jobs key is the dedupe key and Domain the foreign key",
      {(j["Job Key"], j["Domain"]) for j in jobs} >= {(job_key(job("Acme", "", "Backend Engineer", "https://boards.greenhouse.io/acme/jobs/1", "Open")), "acme.com")})
check("Open jobs sort before Closed within a company", [j["Status"] for j in jobs if j["Domain"] == "acme.com"] == ["Open", "Closed"])

before = [[c.value for c in row] for row in wb["Acme"].iter_rows()]
entities.refresh(wb, con)
entities.refresh(wb, con)
check("rebuilding twice keeps one sheet each, same rows", wb.sheetnames.count("Jobs") == 1 and wb["Jobs"].max_row == 5 and wb["Companies"].max_row == 4)
check("company sheets are untouched", before == [[c.value for c in row] for row in wb["Acme"].iter_rows()])

check("entity sheets are marked read-only (no password) with sorting and filtering allowed",
      wb["Jobs"].protection.sheet and not wb["Jobs"].protection.autoFilter and not wb["Jobs"].protection.sort)
foreign = Workbook()
fs = foreign.active
fs.title = "Jobs"
fs.append(["Role", "Pay"])
fs.append(["Chef", 10])
try:
    entities.refresh(foreign, con)
    clobbered = True
except SystemExit:
    clobbered = False
check("a company sheet that happens to be called Jobs is refused, never deleted", not clobbered and foreign["Jobs"]["A2"].value == "Chef")
check("Companies and Jobs are reserved sheet names",
      safe_sheet_name("Jobs", set(wb.sheetnames) | set(entities.ENTITY_SHEETS)) == "Jobs (2)")

out = io.StringIO()
with redirect_stdout(out):
    dry = entities.migrate(wbp, db, dry_run=True, backups=tmp / "bk")
check("dry run changes nothing and takes no backup", dry["backups"] == [] and "Jobs" not in load_workbook(wbp).sheetnames and not (tmp / "bk").exists())
with redirect_stdout(out):
    real = entities.migrate(wbp, db, stamp="T1", backups=tmp / "bk")
check("migrate backs up the workbook and the database, then writes the sheets",
      len(real["backups"]) == 2 and (tmp / "bk" / "wb.before-entities-T1.xlsx").exists() and "Jobs" in load_workbook(wbp).sheetnames)
try:
    with redirect_stdout(out):
        entities.migrate(wbp, db, stamp="T1", backups=tmp / "bk")
    refused = False
except SystemExit:
    refused = True
check("an existing backup is never overwritten", refused)

shutil.rmtree(tmp, ignore_errors=True)
bad = [n for n, ok in results if not ok]
print(f"\n{len(results) - len(bad)}/{len(results)} passed")
sys.exit(1 if bad else 0)
