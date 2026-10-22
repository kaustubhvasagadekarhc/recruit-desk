"""One-time backfill of the structured fields and the Evidence sheet (2026-10-16, plan items 1 and 4).

uv run scripts/backfill_structured.py [--dry-run]

Builder or human only (not in the agent's allow list): back up Candidate_Research.xlsx and memory.db first.
For every JD sheet it appends any missing headers in place (no header or cell of an older column is changed) and
fills the new typed columns of every existing row from that row's own cells (common.structured_fields), then
re-exports the Evidence sheet from memory.db and refreshes Evidence Count. The next upsert of a JD does the same
for its rows, so this is only needed for sheets written before the change. Idempotent.
"""
from __future__ import annotations

import argparse
import json
import sqlite3
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment

sys.path.insert(0, str(Path(__file__).parent))
from common import COLUMNS, DB_PATH, ENTITY_SHEETS, EVIDENCE_SHEET, HEADERS, SCHEMA_PATH, WORKBOOK_PATH, header_len, structured_fields  # noqa: E402
from excel_upsert import evidence_table, prepare_sheet, refresh_evidence_counts, write_evidence_sheet  # noqa: E402
import progress  # noqa: E402

NEW_FIELDS = [f for f, _, _ in COLUMNS[COLUMNS.index(next(c for c in COLUMNS if c[0] == "first_name")):] if f != "evidence_count"]


def backfill(wb_path: Path = WORKBOOK_PATH, db_path: Path = DB_PATH, dry_run: bool = False) -> dict:
    if not wb_path.exists():
        progress.fail_exit(f"no workbook at {wb_path}")
    wb = load_workbook(wb_path)
    con = sqlite3.connect(db_path)
    con.executescript(SCHEMA_PATH.read_text())
    field_of = {h: f for f, h, _ in COLUMNS}
    sheets, rows_filled = [], 0
    for ws in list(wb.worksheets):
        if ws.title in (EVIDENCE_SHEET, *ENTITY_SHEETS) or header_len([c.value for c in ws[1]]) is None:
            continue
        prepare_sheet(wb, ws.title)
        col = {h: i + 1 for i, h in enumerate(HEADERS)}
        for r in range(2, ws.max_row + 1):
            if not ws.cell(r, col["Candidate Name"]).value:
                continue
            row = {field_of[h]: ws.cell(r, i).value for h, i in col.items()}
            for f, v in structured_fields(row).items():
                if f in NEW_FIELDS:
                    ws.cell(r, col[next(h for ff, h, _ in COLUMNS if ff == f)], v if v != "" else None).alignment = Alignment(vertical="top")
            rows_filled += 1
        ws.auto_filter.ref = f"A1:{ws.cell(1, len(HEADERS)).column_letter}{ws.max_row}"
        sheets.append(ws.title)
    ev = evidence_table(con)
    write_evidence_sheet(wb, ev)
    refresh_evidence_counts(wb, ev)
    con.close()
    if not dry_run:
        wb.save(wb_path)
    return {"sheets": sheets, "rows_filled": rows_filled, "evidence_rows": len(ev), "dry_run": dry_run}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    r = backfill(dry_run=a.dry_run)
    progress.result(f"{r['rows_filled']} rows filled on {len(r['sheets'])} sheet(s) · Evidence sheet {r['evidence_rows']} rows"
                    + (" (dry run, nothing saved)" if a.dry_run else ""))
    print(json.dumps(r, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
