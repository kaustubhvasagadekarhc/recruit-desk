"""Idempotent upsert of validated candidates into Candidate_Research.xlsx.

uv run scripts/excel_upsert.py --input runs/<file>.json [--run <run_id>] [--dry-run] [--contacts]

One worksheet per JD (name from memory.db). A row is matched by its dedupe key
(normalised Profile URL, else name plus company): matches are updated in place,
new candidates are appended. Recruiter Status and Date Found are never
overwritten on an existing row, so recruiter edits survive reruns.
Also records candidates and scores in memory.db.

Privacy (plan item 10, 2026-10-15):
- Email / Mobile is written only when the run passes --contacts; otherwise it is left blank (data minimisation).
- Every row's sources go to the memory.db sources table (dedupe key, url, source type, query id, fetched_at,
  snippet hash): the primary profile and, if any, the corroboration page. `memory.py where <key>` prints it.
- People in the suppressed table (`memory.py forget`) are never written.
- Sheets written before 2026-10-15 get the new Evidence Sources column appended to their header in place.

Key factors (2026-10-15): four columns appended at the end, in this order: Open to Work, Stack Match,
Experience Match, Location Match (filled by score.py). Older sheets (14 or 15 columns) get the missing headers
appended in place on their next upsert; their existing rows stay blank there until rescored and upserted.

Structured fields and Evidence (2026-10-16, data-structure plan items 1 and 4): 16 typed columns appended at the end
(First Name ... Cap Adjustment pts, Evidence Count; see common.COLUMNS), filled by common.structured_fields() from the
row itself; the score parts plus Cap Adjustment pts add up to Match Score. Each row's cited facts and quotes go to the
memory.db evidence table (replaced per candidate and JD on every upsert), and the whole table, plus any source with no
fact yet, is re-exported to the 'Evidence' sheet. Evidence Count (rows in the Evidence sheet for that candidate, all
JDs) is refreshed on every JD sheet. Older sheets (14, 15 or 19 columns) get the new headers appended in place.

GitHub and Stack Exchange (2026-10-16, plan item 9): GitHub Languages, GitHub Last Active and SO Top Tags appended at
the end, from the batch's dev_signals (scripts/dev_evidence.py); a batch without them keeps the cells an earlier fetch
wrote. Their facts go to the Evidence sheet as "GitHub API" / "Stack Exchange API" rows and the API pages to sources.
35-column sheets get the three headers appended in place.
"""
from __future__ import annotations

import argparse
import json
import sqlite3
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
from common import (COLUMNS, DB_PATH, DEV_FIELDS, EVIDENCE_COLUMNS, EVIDENCE_HEADERS, EVIDENCE_SHEET, HEADERS,  # noqa: E402
                    ENTITY_SHEETS, SCHEMA_PATH, STATUSES, WORKBOOK_PATH, drop_suppressed, evidence_rows, header_len,
                    dev_sources, row_key, snippet_hash, structured_fields, validate_payload)
import entities  # noqa: E402
import last_record  # noqa: E402
import progress  # noqa: E402

PRESERVE = {"Recruiter Status", "Date Found"}


def now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def prepare_sheet(wb, name: str):
    if name in wb.sheetnames:
        ws = wb[name]
        n = header_len([c.value for c in ws[1]])
        if n is None:
            progress.fail_exit(f"sheet {name!r} has an unexpected header row; refusing to write")
        if n < len(HEADERS):   # older sheet: append the newer columns to the header in place
            for i in range(n, len(HEADERS)):
                cell = ws.cell(1, i + 1, HEADERS[i])
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                ws.column_dimensions[cell.column_letter].width = COLUMNS[i][2]
            progress.sub(f"sheet {name!r}: added column(s) {', '.join(HEADERS[n:])}")
        return ws
    ws = wb.create_sheet(name)
    ws.append(HEADERS)
    for i, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    status_col = ws.cell(1, HEADERS.index("Recruiter Status") + 1).column_letter
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    dv.add(f"{status_col}2:{status_col}5000")
    ws.add_data_validation(dv)
    return ws


def record_sources(con, key: str, c: dict, jd_id: str) -> None:
    """Audit trail: the primary source and the corroboration page (if any), once per (key, url, snippet)."""
    found = c.get("fetched_at") or c.get("date_found") or now()
    rows = [(c["profile_url"], c.get("source") or "unknown", c.get("query_id"), found,
             snippet_hash(c.get("evidence_notes_raw") or c.get("evidence_notes") or ""))]
    cor = c.get("corroboration")
    if isinstance(cor, dict) and cor.get("url"):
        rows.append((cor["url"], cor.get("source_type") or "Web page", None, cor.get("fetched_at") or now(),
                     snippet_hash(cor.get("snippet") or "")))
    for kind, sig in dev_sources(c):
        rows.append((sig["url"], kind, None, sig.get("fetched_at") or now(),
                     snippet_hash(json.dumps(sig.get("facts"), sort_keys=True))))
    for url, kind, qid, at, h in rows:
        con.execute("INSERT OR IGNORE INTO sources (dedupe_key, url, source_type, query_id, fetched_at, snippet_hash, jd_id) "
                    "VALUES (?,?,?,?,?,?,?)", (key, url, kind, qid, at, h, jd_id))


def record_evidence(con, key: str, c: dict, jd_id: str) -> None:
    """Replace this candidate's evidence rows for this JD with the ones the batch cites."""
    con.execute("DELETE FROM evidence WHERE dedupe_key=? AND jd_id IS ?", (key, jd_id))
    for e in evidence_rows(c):
        con.execute("INSERT OR IGNORE INTO evidence (dedupe_key, jd_id, fact, quote, url, source_type, fetched_at) "
                    "VALUES (?,?,?,?,?,?,?)", (key, jd_id, e["fact"], e["quote"], e["url"], e["source_type"], e["fetched_at"]))


def evidence_table(con) -> list[tuple]:
    """Evidence sheet rows: every evidence row, plus every source with no evidence row yet (Fact and Quote blank)."""
    names = dict(con.execute("SELECT dedupe_key, name FROM candidates"))
    out = [(k, f, q, u, t, (at or "")[:10], names.get(k, ""), jd) for k, f, q, u, t, at, jd in con.execute(
        "SELECT dedupe_key, fact, quote, url, source_type, fetched_at, jd_id FROM evidence")]
    have = {(r[0], r[3]) for r in out}
    for k, u, t, at, jd in con.execute("SELECT dedupe_key, url, source_type, fetched_at, jd_id FROM sources"):
        if (k, u) not in have:
            have.add((k, u))
            out.append((k, "", "", u, t, (at or "")[:10], names.get(k, ""), jd))
    return sorted(out, key=lambda r: (r[0], r[7] or "", r[5], r[3], r[1], r[2]))


def write_evidence_sheet(wb, rows: list[tuple]):
    """(Re)build the Evidence sheet from rows; it is a pure export, so recruiters edit nothing there."""
    if EVIDENCE_SHEET in wb.sheetnames:
        del wb[EVIDENCE_SHEET]
    ws = wb.create_sheet(EVIDENCE_SHEET)
    ws.append(EVIDENCE_HEADERS)
    for i, (_, width) in enumerate(EVIDENCE_COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    url_col = EVIDENCE_HEADERS.index("Source URL") + 1
    for r in rows:
        ws.append([v if v not in ("", None) else None for v in r])
        cell = ws.cell(ws.max_row, url_col)
        if cell.value and str(cell.value).startswith("http"):
            cell.hyperlink = cell.value
            cell.font = Font(color="0563C1", underline="single")
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(EVIDENCE_HEADERS)).column_letter}{max(ws.max_row, 1)}"
    return ws


def refresh_evidence_counts(wb, rows: list[tuple]) -> None:
    """Evidence Count on every current-header JD sheet: Evidence rows for that candidate (all JDs)."""
    counts: dict[str, int] = {}
    for r in rows:
        counts[r[0]] = counts.get(r[0], 0) + 1
    for ws in wb.worksheets:
        hdr = [c.value for c in ws[1]]
        if ws.title in (EVIDENCE_SHEET, *ENTITY_SHEETS) or header_len(hdr) != len(HEADERS):
            continue
        ci = {h: i + 1 for i, h in enumerate(HEADERS)}
        for r in range(2, ws.max_row + 1):
            name = ws.cell(r, ci["Candidate Name"]).value
            if name:
                k = row_key(ws.cell(r, ci["Profile URL"]).value or "", name, ws.cell(r, ci["Current Role / Company"]).value or "")
                ws.cell(r, ci["Evidence Count"], counts.get(k, 0))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--run", type=int)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--contacts", action="store_true", help="keep Email / Mobile (only when the JD run asked for contacts)")
    a = ap.parse_args()
    payload = json.loads(Path(a.input).read_text())
    if not a.contacts:
        blanked = sum(bool((c.get("email_mobile") or "").strip()) for c in payload.get("candidates") or [])
        for c in payload.get("candidates") or []:
            c["email_mobile"], c["contact_source_url"] = "", ""
        if blanked:
            progress.sub(f"{blanked} Email / Mobile values left blank (no --contacts on this run)")
    errs = validate_payload(payload)
    if errs:
        print("VALIDATION FAILED:\n- " + "\n- ".join(errs), file=sys.stderr)
        progress.fail_exit(f"validation failed: {len(errs)} problems, nothing written", 2)
    progress.sub(f"validated {len(payload['candidates'])} rows: all pass")

    con = sqlite3.connect(DB_PATH)
    con.executescript(SCHEMA_PATH.read_text())
    jd = con.execute("SELECT id, sheet_name FROM jds WHERE id=?", (payload["jd_id"],)).fetchone()
    if not jd:
        progress.fail_exit(f"unknown jd_id {payload['jd_id']!r}; run memory.py add-jd first")
    sheet = jd[1]
    jd_title = con.execute("SELECT title FROM jds WHERE id=?", (payload["jd_id"],)).fetchone()[0]
    payload["candidates"], dropped = drop_suppressed(payload["candidates"], DB_PATH)
    if dropped:
        progress.sub(f"{dropped} suppressed people not written (they asked to be forgotten)")

    wb = load_workbook(WORKBOOK_PATH) if WORKBOOK_PATH.exists() else Workbook()
    if "Sheet" in wb.sheetnames and wb["Sheet"].max_row == 1 and wb["Sheet"]["A1"].value is None:
        del wb["Sheet"]
    ws = prepare_sheet(wb, sheet)
    col = {h: i + 1 for i, h in enumerate(HEADERS)}
    index = {}
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, col["Candidate Name"]).value
        if name:
            k = row_key(ws.cell(r, col["Profile URL"]).value or "", name, ws.cell(r, col["Current Role / Company"]).value or "")
            index.setdefault(k, r)

    inserted = updated = 0
    newest = None   # (dedupe key, record) of the last NEW candidate of this batch
    new_rows = []   # every NEW candidate of this batch, for the person's search history
    total = len(payload["candidates"])
    if not total:
        progress.result("nothing to write: every candidate in the batch is suppressed")
        print(json.dumps({"sheet": sheet, "inserted": 0, "updated": 0, "rows_in_sheet": ws.max_row - 1,
                          "workbook": str(WORKBOOK_PATH), "dry_run": a.dry_run, "suppressed_dropped": dropped}, indent=2))
        return
    for i, c in enumerate(payload["candidates"], 1):
        c.setdefault("recruiter_status", "New")
        k = row_key(c["profile_url"], c["candidate_name"], c.get("current_role_company", ""))
        r = index.get(k)
        if r is None:
            r = ws.max_row + 1
            index[k] = r
            inserted += 1
            fresh = True
        else:
            updated += 1
            fresh = False
        c.update(structured_fields(c))
        if fresh:
            newest = (k, {"name": c["candidate_name"], "title": c.get("current_title"), "company": c.get("current_company"),
                          "location": c.get("location"), "match_score": c.get("match_score"), "profile_url": c.get("profile_url"),
                          "jd": jd_title, "date_found": c.get("date_found")})
            new_rows.append(newest)
        for field, header, _ in COLUMNS:
            if field == "evidence_count":
                continue   # refreshed for every row below, from the Evidence sheet
            if not fresh and header in PRESERVE and ws.cell(r, col[header]).value:
                continue
            v = c.get(field, "")
            if field in DEV_FIELDS and v in (None, "") and not fresh:
                continue   # a batch without GitHub / Stack Exchange data keeps what an earlier fetch wrote
            if isinstance(v, list):
                v = ", ".join(map(str, v))
            cell = ws.cell(r, col[header], v if v != "" else None)
            cell.alignment = Alignment(wrap_text=header in {"Score Breakdown", "Evidence / Notes", "Key Skills", "Stack Match"}, vertical="top")
            if header == "Profile URL" and v:
                cell.hyperlink = v
                cell.font = Font(color="0563C1", underline="single")
        con.execute("INSERT INTO candidates VALUES (?,?,?,?,?,?) ON CONFLICT(dedupe_key) DO UPDATE SET "
                    "name=excluded.name, company=excluded.company, profile_url=excluded.profile_url, last_seen=excluded.last_seen",
                    (k, c["candidate_name"], c.get("current_role_company"), c["profile_url"], now(), now()))
        con.execute("INSERT OR REPLACE INTO candidate_scores VALUES (?,?,?,?,?,?,?)",
                    (k, payload["jd_id"], c["match_score"], c["score_breakdown"], c["evidence_confidence"], a.run, now()))
        record_sources(con, k, c, payload["jd_id"])
        record_evidence(con, k, c, payload["jd_id"])
        if progress.every(i, total):
            progress.sub(f"writing {i}/{total} · {inserted} inserted, {updated} updated so far")
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(HEADERS)).column_letter}{ws.max_row}"
    ev = evidence_table(con)
    write_evidence_sheet(wb, ev)
    refresh_evidence_counts(wb, ev)
    progress.sub(f"Evidence sheet: {len(ev)} rows")
    people, matches = entities.refresh(wb, con)
    progress.sub(f"Candidates sheet: {people} people · Matches sheet: {matches} rows")

    if a.dry_run:
        con.rollback()
    else:
        wb.save(WORKBOOK_PATH)
        if newest:   # a plain local file; Recruit Desk stores it as the last record while its person is signed in
            last_record.write(last_record.folder(WORKBOOK_PATH.parent) / last_record.FILE_NAME, *newest)
            last_record.append_rows(last_record.folder(WORKBOOK_PATH.parent), new_rows)
        if a.run:
            con.execute("UPDATE runs SET found=found+?, inserted=inserted+?, updated=updated+? WHERE id=?",
                        (len(payload["candidates"]), inserted, updated, a.run))
        con.commit()
    rows = ws.max_row - 1
    progress.result(f"{inserted} inserted, {updated} updated · sheet '{sheet}' now {rows} rows"
                    + (" (dry run, nothing saved)" if a.dry_run else ""))
    print(json.dumps({"sheet": sheet, "inserted": inserted, "updated": updated,
                      "rows_in_sheet": ws.max_row - 1, "workbook": str(WORKBOOK_PATH), "dry_run": a.dry_run}, indent=2))


if __name__ == "__main__":
    main()
