"""Retention (security review step 9): find Rejected or Stale candidate rows older than N days.

  uv run scripts/retention.py [--days 180] [--new-days 90]   # dry run (default): list only, change nothing
  uv run scripts/retention.py [--days 180] --apply    # human only, from a terminal: delete them

A row's age is its Date Found. --apply removes the rows from every JD sheet of
Candidate_Research.xlsx and their scores and feedback from memory.db (a candidate with no score
left for any JD is removed too). The agent's PreToolUse hook blocks --apply; only the human runs it.

The dry run also lists rows still "New" (never reviewed) whose Date Found is older than --new-days (90):
a human should review or reject them. --apply never deletes these; they are only listed (plan item 10).
"""
from __future__ import annotations

import argparse
import os
import sqlite3
import sys
import tempfile
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, os.path.dirname(__file__))
from common import DB_PATH, HEADERS, WORKBOOK_PATH, delete_sheet_rows, header_ok, row_key  # noqa: E402

PURGE_STATUSES = {"Rejected", "Stale"}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--days", type=int, default=180)
    ap.add_argument("--new-days", type=int, default=90, help="list New rows not reviewed for this many days")
    ap.add_argument("--apply", action="store_true", help="delete the listed rows (human only)")
    ap.add_argument("--workbook", default=str(WORKBOOK_PATH))
    ap.add_argument("--db", default=str(DB_PATH))
    a = ap.parse_args()
    cutoff = (date.today() - timedelta(days=a.days)).isoformat()
    new_cutoff = (date.today() - timedelta(days=a.new_days)).isoformat()
    WB, DB = Path(a.workbook), Path(a.db)
    if not WB.exists():
        print(f"no workbook at {WB}; nothing to do")
        return

    from openpyxl import load_workbook

    wb = load_workbook(WB)
    con = sqlite3.connect(DB)
    jd_by_sheet = {s: j for j, s in con.execute("SELECT id, sheet_name FROM jds")}
    hits, stale_new, undated = [], [], 0
    for ws in wb.worksheets:
        if not header_ok([c.value for c in ws[1]]):
            continue
        col = {h: i + 1 for i, h in enumerate(HEADERS)}
        for r in range(2, ws.max_row + 1):
            status = (ws.cell(r, col["Recruiter Status"]).value or "").strip()
            if status in ("", "New") and ws.cell(r, col["Candidate Name"]).value:
                found = str(ws.cell(r, col["Date Found"]).value or "")[:10]
                if found and found < new_cutoff:
                    stale_new.append((ws.title, r, ws.cell(r, col["Candidate Name"]).value, found))
                continue
            if status not in PURGE_STATUSES:
                continue
            found = str(ws.cell(r, col["Date Found"]).value or "")[:10]
            if not found:
                undated += 1
                continue
            if found < cutoff:
                name = ws.cell(r, col["Candidate Name"]).value or ""
                key = row_key(ws.cell(r, col["Profile URL"]).value or "", name, ws.cell(r, col["Current Role / Company"]).value or "")
                hits.append((ws, r, key, jd_by_sheet.get(ws.title), name, status, found))

    mode = "APPLY" if a.apply else "dry run"
    print(f"retention {mode}: {len(hits)} Rejected/Stale rows with Date Found before {cutoff} ({a.days} days)")
    for ws, r, key, jd, name, status, found in hits:
        print(f"  {ws.title} · row {r} · {name} · {status} · found {found}")
    if undated:
        print(f"  {undated} Rejected/Stale rows have no Date Found and were skipped")
    print(f"not reviewed: {len(stale_new)} rows still New with Date Found before {new_cutoff} ({a.new_days} days); "
          "review or reject them (listed only, never deleted by --apply)")
    for title, r, name, found in stale_new:
        print(f"  {title} · row {r} · {name} · New · found {found}")
    if not a.apply:
        print("nothing changed; a human adds --apply in a terminal to delete these rows")
        return
    if not hits:
        return

    for ws in {h[0] for h in hits}:
        delete_sheet_rows(ws, [h[1] for h in hits if h[0] is ws])
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=WB.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, WB)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    for _, _, key, jd, *_ in hits:
        if jd:
            con.execute("DELETE FROM candidate_scores WHERE dedupe_key=? AND jd_id=?", (key, jd))
            con.execute("DELETE FROM feedback WHERE dedupe_key=? AND jd_id=?", (key, jd))
    for key in {h[2] for h in hits}:
        con.execute("DELETE FROM candidates WHERE dedupe_key=? AND dedupe_key NOT IN (SELECT dedupe_key FROM candidate_scores)", (key,))
    con.commit()
    print(f"deleted {len(hits)} rows from {WB.name} and their scores and feedback from memory.db")


if __name__ == "__main__":
    main()
