"""Export candidates as CRM contacts (data-structure plan item 8, 2026-10-16).

uv run scripts/export_crm.py --format hubspot|salesforce [--out exports/<file>.csv] [--status Shortlisted ...]

Reads every JD sheet of Candidate_Research.xlsx (read only) and writes one CSV row per person, merged across JDs
(best Match Score, all JD sheets listed). The record key is the email when the row holds one (the upsert only
keeps Email / Mobile when the person self-published it and the run asked for contacts), otherwise the profile URL.
Nothing is looked up or guessed: no email harvesting, blanks stay blank, and placeholder text never reaches a CRM
field. Column names follow the HubSpot contact import and the Salesforce Data Import Wizard (Contacts).
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
from common import ENTITY_SHEETS, EVIDENCE_SHEET, HEADERS, ROOT, STATUSES, WORKBOOK_PATH, header_len, row_key, structured_fields  # noqa: E402
import progress  # noqa: E402

EXPORT_DIR = ROOT / "exports"
FORMATS = {
    "hubspot": ["Record Key", "Record Key Type", "First Name", "Last Name", "Email", "Phone Number", "Job Title",
                "Company Name", "City", "Country/Region", "Profile URL", "Lead Status", "Match Score", "Matched JDs"],
    "salesforce": ["Record Key", "Record Key Type", "First Name", "Last Name", "Email", "Phone", "Title", "Account Name",
                   "Mailing City", "Mailing Country", "Profile URL", "Lead Source", "Description", "Match Score", "Matched JDs"],
}
_EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_PHONE = re.compile(r"\+?\d[\d ()-]{7,}\d")


def split_contact(text: str) -> tuple[str, str]:
    """(email, phone) found in an Email / Mobile cell; '' for each that is not there."""
    t = text or ""
    email = _EMAIL.search(t)
    phone = _PHONE.search(_EMAIL.sub(" ", t))
    return (email.group(0).lower() if email else ""), (phone.group(0).strip() if phone else "")


def city_country(location: str) -> tuple[str, str]:
    """'Pune, Maharashtra, India' -> ('Pune', 'India'); one part -> (part, ''); remote/unknown text is not a city."""
    parts = [p.strip() for p in (location or "").split(",") if p.strip()]
    if not parts or re.search(r"remote|unknown|not (stated|shown)", parts[0], re.I):
        return "", ""
    return parts[0], (parts[-1] if len(parts) > 1 else "")


def read_people(wb_path: Path, statuses: set[str] | None = None) -> list[dict]:
    """One merged dict per person across every JD sheet (header must be a known one)."""
    wb = load_workbook(wb_path, read_only=True)
    people: dict[str, dict] = {}
    for ws in wb.worksheets:
        if ws.title in (EVIDENCE_SHEET, *ENTITY_SHEETS):
            continue
        rows = ws.iter_rows(values_only=True)
        hdr = list(next(rows, []) or [])
        n = header_len(hdr)
        if n is None:
            continue
        for vals in rows:
            r = {HEADERS[i]: vals[i] for i in range(min(n, len(vals)))}
            name = r.get("Candidate Name")
            if not name:
                continue
            if statuses and (r.get("Recruiter Status") or "New") not in statuses:
                continue
            k = row_key(r.get("Profile URL") or "", name, r.get("Current Role / Company") or "")
            p = people.setdefault(k, {"jds": [], "rows": []})
            p["jds"].append(ws.title)
            p["rows"].append(r)
    wb.close()
    return [dict(v, key=k) for k, v in people.items()]


def to_contact(person: dict, fmt: str) -> dict:
    rows = person["rows"]
    best = max(rows, key=lambda r: r.get("Match Score (0-100)") or 0)
    s = structured_fields({"candidate_name": best.get("Candidate Name"), "current_role_company": best.get("Current Role / Company"),
                           "experience": best.get("Experience"), "location": best.get("Location"),
                           "first_name": best.get("First Name"), "last_name": best.get("Last Name"),
                           "current_title": best.get("Title"), "current_company": best.get("Company")})
    email = phone = ""
    for r in rows:   # a self-published contact on any of the person's rows
        e, ph = split_contact(str(r.get("Email / Mobile") or ""))
        email, phone = email or e, phone or ph
    url = best.get("Profile URL") or ""
    city, country = city_country(best.get("Location") or "")
    key, key_type = (email, "email") if email else (url, "profile_url")
    first, last = s["first_name"], s["last_name"]
    jds = "; ".join(dict.fromkeys(person["jds"]))
    score = best.get("Match Score (0-100)")
    if fmt == "hubspot":
        return {"Record Key": key, "Record Key Type": key_type, "First Name": first, "Last Name": last, "Email": email,
                "Phone Number": phone, "Job Title": s["current_title"], "Company Name": s["current_company"], "City": city,
                "Country/Region": country, "Profile URL": url, "Lead Status": best.get("Recruiter Status") or "New",
                "Match Score": score, "Matched JDs": jds}
    if not last:   # Salesforce requires Last Name: a one-word name goes there, never an invented surname
        first, last = "", first
    return {"Record Key": key, "Record Key Type": key_type, "First Name": first, "Last Name": last, "Email": email,
            "Phone": phone, "Title": s["current_title"], "Account Name": s["current_company"], "Mailing City": city,
            "Mailing Country": country, "Profile URL": url, "Lead Source": "Candidate Research",
            "Description": f"Recruiter Status: {best.get('Recruiter Status') or 'New'}; profile: {url}",
            "Match Score": score, "Matched JDs": jds}


def export(fmt: str, out: Path, wb_path: Path = WORKBOOK_PATH, statuses: set[str] | None = None) -> dict:
    if not wb_path.exists():
        progress.fail_exit(f"no workbook at {wb_path}")
    contacts = [to_contact(p, fmt) for p in read_people(wb_path, statuses)]
    contacts = [c for c in contacts if c["Record Key"]]   # no key (no email, no URL) = cannot be imported safely
    contacts.sort(key=lambda c: (-(c["Match Score"] or 0), c["Last Name"], c["First Name"]))
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FORMATS[fmt])
        w.writeheader()
        for c in contacts:
            w.writerow({k: ("" if v is None else v) for k, v in c.items()})
    by_email = sum(c["Record Key Type"] == "email" for c in contacts)
    return {"format": fmt, "out": str(out), "contacts": len(contacts), "keyed_by_email": by_email,
            "keyed_by_profile_url": len(contacts) - by_email}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--format", required=True, choices=sorted(FORMATS))
    ap.add_argument("--out", help="CSV path (default exports/candidates-<format>-<date>.csv)")
    ap.add_argument("--status", nargs="*", choices=STATUSES, help="only rows with these Recruiter Status values")
    a = ap.parse_args()
    out = Path(a.out) if a.out else EXPORT_DIR / f"candidates-{a.format}-{date.today().isoformat()}.csv"
    r = export(a.format, out, statuses=set(a.status) if a.status else None)
    progress.result(f"{r['contacts']} contacts → {r['out']} ({r['keyed_by_email']} keyed by email, "
                    f"{r['keyed_by_profile_url']} by profile URL)")


if __name__ == "__main__":
    main()
