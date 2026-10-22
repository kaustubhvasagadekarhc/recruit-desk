"""Entity sheets: one row per candidate, one row per candidate-JD match (2026-10-18, data-structure plan item 7).

uv run scripts/entities.py [--dry-run]

Builder or human only (not in the agent's allow list). Backs up Candidate_Research.xlsx and memory.db next to
them (*.backup-<date>-before-entities.*), then adds two export sheets built from the JD sheets:

- 'Candidates': one row per person, however many JDs they match (key = the dedupe key).
- 'Matches': one row per (person, JD) with the score, its parts, Open to Work and the Recruiter Status.

Both are pure exports like 'Evidence': the JD sheets stay the place where rows are written and where recruiters edit
Recruiter Status, and every upsert rebuilds the two sheets from them. Nothing in a JD sheet is changed. Email /
Mobile is not copied (export_crm.py handles contacts). Idempotent: a rerun gives the same sheets.
"""
from __future__ import annotations

import argparse
import json
import shutil
import sqlite3
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).parent))
from common import (COLUMNS, DB_PATH, ENTITY_SHEETS, EVIDENCE_SHEET, HEADERS, SCHEMA_PATH, WORKBOOK_PATH,  # noqa: E402
                    header_len, row_key, structured_fields)
import progress  # noqa: E402

CANDIDATES_SHEET, MATCHES_SHEET = ENTITY_SHEETS
PTS = [h for _, h, _ in COLUMNS if h.endswith(" pts")]

CANDIDATE_COLUMNS = [("Candidate Key", 36), ("Candidate Name", 24), ("First Name", 14), ("Last Name", 16), ("Title", 30),
                     ("Company", 22), ("Location", 22), ("Work Mode", 10), ("Experience Years", 10), ("Profile URL", 44),
                     ("JDs Matched", 9), ("Best Match", 9), ("Best Match JD", 28), ("Evidence Count", 9),
                     ("First Seen", 12), ("Last Seen", 12)]
MATCH_COLUMNS = [("Candidate Key", 36), ("Candidate Name", 24), ("JD Id", 28), ("JD Title", 30), ("Sheet", 26),
                 ("Match Score (0-100)", 10), ("Evidence Confidence", 11), *[(h, 8) for h in PTS], ("Open to Work", 10),
                 ("Recruiter Status", 14), ("Date Found", 12)]
FIELD_OF = {h: f for f, h, _ in COLUMNS}


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def collect(wb, con) -> tuple[list[list], list[list]]:
    """(candidate rows, match rows) read from every current or older JD sheet of wb and from memory.db."""
    jds = {r[0]: (r[1], r[2]) for r in con.execute("SELECT sheet_name, id, title FROM jds")}
    seen = {r[0]: (r[1] or "", r[2] or "") for r in con.execute("SELECT dedupe_key, first_seen, last_seen FROM candidates")}
    ev_counts: dict[str, int] = {}
    if EVIDENCE_SHEET in wb.sheetnames:
        rows = wb[EVIDENCE_SHEET].iter_rows(min_row=2, values_only=True)
        for r in rows:
            if r and r[0]:
                ev_counts[r[0]] = ev_counts.get(r[0], 0) + 1
    people: dict[str, dict] = {}
    matches: list[list] = []
    for ws in wb.worksheets:
        if ws.title == EVIDENCE_SHEET or ws.title in ENTITY_SHEETS:
            continue
        hdr = [c.value for c in ws[1]]
        n = header_len(hdr)
        if n is None:
            continue
        jd_id, jd_title = jds.get(ws.title, ("", ""))
        for vals in ws.iter_rows(min_row=2, values_only=True):
            cells = {HEADERS[i]: vals[i] for i in range(min(n, len(vals)))}
            if not cells.get("Candidate Name"):
                continue
            derived = structured_fields({FIELD_OF[h]: v for h, v in cells.items()})
            row = {h: cells.get(h) if cells.get(h) not in (None, "") else derived.get(FIELD_OF[h], "") for h in HEADERS}
            key = row_key(row["Profile URL"] or "", row["Candidate Name"], row["Current Role / Company"] or "")
            score = _num(row["Match Score (0-100)"])
            otw = str(row["Open to Work"] or "").split(":")[0].strip()
            matches.append([key, row["Candidate Name"], jd_id, jd_title, ws.title, row["Match Score (0-100)"],
                            row["Evidence Confidence"], *[row[h] for h in PTS], otw or None,
                            row["Recruiter Status"], row["Date Found"]])
            p = people.setdefault(key, {"row": row, "jds": set(), "best": (-1.0, "", None)})
            p["jds"].add(ws.title)
            if score is not None and score > p["best"][0]:
                p["best"] = (score, jd_title or ws.title, row["Match Score (0-100)"])
                p["row"] = row   # the best match's row carries the freshest, fullest cells
    cands = []
    for key, p in people.items():
        r = p["row"]
        first, last = seen.get(key, ("", ""))
        cands.append([key, r["Candidate Name"], r["First Name"], r["Last Name"], r["Title"], r["Company"], r["Location"],
                      r["Work Mode"], r["Experience Years"], r["Profile URL"], len(p["jds"]), p["best"][2], p["best"][1] or None,
                      ev_counts.get(key, 0), first[:10] or None, last[:10] or None])
    cands.sort(key=lambda c: (-(_num(c[11]) or -1), str(c[1]).lower()))
    matches.sort(key=lambda m: (str(m[4]), -(_num(m[5]) or -1), str(m[1]).lower()))
    return cands, matches


def _protect(ws) -> None:
    """Read-only marker (no password; sorting and filtering stay allowed): the sheet is rebuilt on every upsert, so an edit
    made in it, Recruiter Status included, would be lost. Edit the JD sheet instead."""
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.sheet_properties.tabColor = "808080"


def _write(wb, name: str, columns: list[tuple[str, int]], rows: list[list]) -> None:
    if name in wb.sheetnames:
        if [c.value for c in wb[name][1]] != [h for h, _ in columns]:
            progress.fail_exit(f"a sheet named '{name}' already exists and is not an entity sheet; rename it before the entity "
                               "sheets can be built. Nothing was changed")
        del wb[name]
    ws = wb.create_sheet(name)
    ws.append([h for h, _ in columns])
    for i, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    url_col = next((i for i, (h, _) in enumerate(columns, start=1) if h == "Profile URL"), None)
    for r in rows:
        ws.append([v if v not in ("", None) else None for v in r])
        if url_col:
            cell = ws.cell(ws.max_row, url_col)
            if cell.value and str(cell.value).startswith("http"):
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(columns)).column_letter}{max(ws.max_row, 1)}"
    _protect(ws)


def refresh(wb, con) -> tuple[int, int]:
    """Rebuild the Candidates and Matches sheets in wb (not saved). Returns (people, matches)."""
    cands, matches = collect(wb, con)
    _write(wb, CANDIDATES_SHEET, CANDIDATE_COLUMNS, cands)
    _write(wb, MATCHES_SHEET, MATCH_COLUMNS, matches)
    # order: JD sheets, Candidates, Matches, Evidence
    for name in (CANDIDATES_SHEET, MATCHES_SHEET, EVIDENCE_SHEET):
        if name in wb.sheetnames:
            wb.move_sheet(name, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index(name))
    return len(cands), len(matches)


def migrate(wb_path: Path = WORKBOOK_PATH, db_path: Path = DB_PATH, dry_run: bool = False, stamp: str | None = None) -> dict:
    if not wb_path.exists():
        progress.fail_exit(f"no workbook at {wb_path}")
    stamp = stamp or date.today().isoformat()
    backups = []
    if not dry_run:
        for p in (wb_path, db_path):
            if p.exists():
                b = p.with_name(f"{p.stem}.backup-{stamp}-before-entities{p.suffix}")
                if b.exists():
                    progress.fail_exit(f"backup {b.name} already exists; move it aside first")
                shutil.copy2(p, b)
                backups.append(b.name)
        progress.sub(f"backed up: {', '.join(backups)}")
    wb = load_workbook(wb_path)
    # A dry run must not create or alter memory.db: read it as it is, or an empty copy when there is none.
    con = sqlite3.connect(db_path if (db_path.exists() or not dry_run) else ":memory:")
    if not (dry_run and db_path.exists()):
        con.executescript(SCHEMA_PATH.read_text())
    people, matches = refresh(wb, con)
    con.close()
    if not dry_run:
        wb.save(wb_path)
    return {"candidates": people, "matches": matches, "backups": backups, "dry_run": dry_run}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true", help="print the counts, change nothing, take no backup")
    a = ap.parse_args()
    r = migrate(dry_run=a.dry_run)
    progress.result(f"{r['candidates']} people, {r['matches']} matches"
                    + (" (dry run, nothing saved)" if a.dry_run else f" · backups {', '.join(r['backups'])}"))
    print(json.dumps(r, indent=2))


if __name__ == "__main__":
    main()
