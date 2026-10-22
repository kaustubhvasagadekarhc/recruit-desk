"""SQLite memory for the Candidate Agent.

uv run scripts/memory.py init
uv run scripts/memory.py add-jd --file jds/<id>.json      # parsed JD: id, title, source, raw_text, criteria, weights
uv run scripts/memory.py get-jd --id <jd_id>
uv run scripts/memory.py synonyms add "<role>" "<synonym>" [...]
uv run scripts/memory.py synonyms list "<role>"
uv run scripts/memory.py plan-queries --jd <jd_id> [--max 6] [--open-to-work]   # planned Exa queries (+ optional open-to-work extra)
uv run scripts/memory.py run-start --jd <jd_id> --queries '["q1","q2"]'   # prints run id
uv run scripts/memory.py run-finish --run <id> [--notes "..."]
uv run scripts/memory.py seen --jd <jd_id>                 # candidates already scored for this JD
uv run scripts/memory.py feedback --jd <jd_id> --key <dedupe_key> [--status S] [--score N] [--note "..."]
uv run scripts/memory.py sync-feedback --jd <jd_id>        # pull Recruiter Status edits from the workbook
uv run scripts/memory.py feedback-list --jd <jd_id>        # read feedback (use this, never the sqlite3 CLI)
uv run scripts/memory.py history [--jd <jd_id>]
uv run scripts/memory.py where <key|profile URL>          # audit trail for a subject-access request
uv run scripts/memory.py forget <key|profile URL> --yes   # erase the person everywhere and suppress them (human asks)
"""
from __future__ import annotations

import argparse
import json
import sqlite3
import sys
from datetime import datetime, timezone

sys.path.insert(0, __import__("os").path.dirname(__file__))
from common import (DB_PATH, DEFAULT_WEIGHTS, SCHEMA_PATH, STATUSES, WORKBOOK_PATH, dedupe_key,  # noqa: E402
                    ENTITY_SHEETS, EVIDENCE_SHEET, delete_sheet_rows, header_ok, key_hash, open_to_work_query, plan_queries, row_key, safe_sheet_name, suppression_hashes)
import last_record  # noqa: E402
import progress  # noqa: E402


def now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def connect() -> sqlite3.Connection:
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    con.executescript(SCHEMA_PATH.read_text())
    return con


def add_jd(con, jd: dict) -> dict:
    for f in ("id", "title", "criteria"):
        if not jd.get(f):
            progress.fail_exit(f"JD file is missing '{f}'")
    weights = jd.get("weights") or DEFAULT_WEIGHTS
    if sum(weights.values()) != 100:
        progress.fail_exit(f"weights must sum to 100, got {sum(weights.values())}")
    row = con.execute("SELECT sheet_name FROM jds WHERE id=?", (jd["id"],)).fetchone()
    if row:
        sheet = row["sheet_name"]
        con.execute(
            "UPDATE jds SET title=?, source=?, raw_text=?, criteria_json=?, weights_json=?, updated_at=? WHERE id=?",
            (jd["title"], jd.get("source"), jd.get("raw_text"), json.dumps(jd["criteria"]), json.dumps(weights), now(), jd["id"]),
        )
    else:
        taken = {r[0] for r in con.execute("SELECT sheet_name FROM jds")} | {EVIDENCE_SHEET, *ENTITY_SHEETS}   # Evidence, Candidates and Matches are reserved
        sheet = safe_sheet_name(jd.get("sheet_name") or jd["title"], taken)
        con.execute(
            "INSERT INTO jds VALUES (?,?,?,?,?,?,?,?,?)",
            (jd["id"], jd["title"], sheet, jd.get("source"), jd.get("raw_text"), json.dumps(jd["criteria"]), json.dumps(weights), now(), now()),
        )
    for syn in jd.get("criteria", {}).get("role_synonyms", []):
        con.execute("INSERT OR IGNORE INTO role_synonyms VALUES (?,?)", (jd["title"], syn))
    con.commit()
    return {"jd_id": jd["id"], "sheet_name": sheet, "weights": weights}


def get_jd(con, jd_id: str) -> dict:
    r = con.execute("SELECT * FROM jds WHERE id=?", (jd_id,)).fetchone()
    if not r:
        sys.exit(f"unknown jd_id {jd_id!r}; run add-jd first")
    d = dict(r)
    d["criteria"] = json.loads(d.pop("criteria_json"))
    d["weights"] = json.loads(d.pop("weights_json"))
    return d


def queries_for(con, jd_id: str, limit: int = 6) -> list[str]:
    """Planned Exa queries for a JD, using its criteria and the role_synonyms memory for its title."""
    jd = get_jd(con, jd_id)
    syns = [r[0] for r in con.execute("SELECT synonym FROM role_synonyms WHERE role=? ORDER BY synonym", (jd["title"],))]
    return plan_queries(jd["title"], jd["criteria"], syns, limit)


def sync_feedback(con, jd_id: str) -> int:
    """Copy Recruiter Status values a human changed in the workbook into feedback."""
    from openpyxl import load_workbook

    jd = get_jd(con, jd_id)
    if not WORKBOOK_PATH.exists():
        return 0
    wb = load_workbook(WORKBOOK_PATH, read_only=True)
    if jd["sheet_name"] not in wb.sheetnames:
        return 0
    ws = wb[jd["sheet_name"]]
    rows = ws.iter_rows(values_only=True)
    hdr = list(next(rows))
    iname, icomp, iurl, istat = (hdr.index(h) for h in ("Candidate Name", "Current Role / Company", "Profile URL", "Recruiter Status"))
    n = 0
    for r in rows:
        status = r[istat]
        if not r[iname] or status is None or status not in STATUSES:
            continue
        key = row_key(r[iurl] or "", r[iname], r[icomp] or "")
        last = con.execute(
            "SELECT recruiter_status FROM feedback WHERE dedupe_key=? AND jd_id=? AND recruiter_status IS NOT NULL ORDER BY id DESC LIMIT 1",
            (key, jd_id),
        ).fetchone()
        if status == "New" and not last:
            continue   # never reviewed: nothing to record. A revert to New after a recorded status is recorded, so it stops counting as Shortlisted
        if not last or last[0] != status:
            con.execute(
                "INSERT INTO feedback (dedupe_key, jd_id, recruiter_status, note, created_at) VALUES (?,?,?,?,?)",
                (key, jd_id, status, "synced from workbook", now()),
            )
            n += 1
    con.commit()
    return n


def resolve_key(con, ref: str) -> str:
    """A dedupe key as given (url:… / nc:…), or the key of a profile URL."""
    ref = (ref or "").strip()
    if ref.startswith(("url:", "nc:")):
        return ref
    return dedupe_key(ref)


def _workbook_rows(wb_path, key: str) -> list[tuple[str, int]]:
    """(sheet, row) of every workbook row whose dedupe key is key."""
    from openpyxl import load_workbook

    if not wb_path.exists():
        return []
    wb = load_workbook(wb_path, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        hdr = list(next(rows, []) or [])
        if not header_ok(hdr):
            continue
        iname, icomp, iurl = (hdr.index(h) for h in ("Candidate Name", "Current Role / Company", "Profile URL"))
        for r, vals in enumerate(rows, start=2):
            if vals[iname] and row_key(vals[iurl] or "", vals[iname], vals[icomp] or "") == key:
                out.append((ws.title, r))
    wb.close()
    return out


def _evidence_sheet_rows(wb_path, keys: set[str], sheet: str = EVIDENCE_SHEET) -> list[int]:
    """Rows of an export sheet (Evidence, Candidates or Matches) whose Candidate Key is one of keys."""
    from openpyxl import load_workbook

    wb = load_workbook(wb_path, read_only=True)
    out = []
    if sheet in wb.sheetnames:
        rows = wb[sheet].iter_rows(values_only=True)
        hdr = list(next(rows, []) or [])
        if "Candidate Key" in hdr:
            i = hdr.index("Candidate Key")
            out = [r for r, vals in enumerate(rows, start=2) if vals and vals[i] in keys]
    wb.close()
    return out


def open_to_work_counts(sheet: str, wb_path=None) -> dict:
    """{"Yes": n, "Likely": n, "rows": n} from the Open to Work column of one JD sheet (zeros if missing)."""
    from openpyxl import load_workbook

    p = wb_path or WORKBOOK_PATH
    out = {"Yes": 0, "Likely": 0, "rows": 0}
    if not p.exists():
        return out
    wb = load_workbook(p, read_only=True)
    if sheet in wb.sheetnames:
        rows = wb[sheet].iter_rows(values_only=True)
        hdr = list(next(rows, []) or [])
        if "Open to Work" in hdr:
            i = hdr.index("Open to Work")
            for vals in rows:
                if not vals or not vals[0]:
                    continue
                out["rows"] += 1
                v = str(vals[i] or "").split(":")[0].strip()
                if v in ("Yes", "Likely"):
                    out[v] += 1
    wb.close()
    return out


def where(con, ref: str, wb_path=None) -> dict:
    """Everything held about one person: memory rows, workbook rows and the source trail."""
    key = resolve_key(con, ref)
    cand = con.execute("SELECT * FROM candidates WHERE dedupe_key=?", (key,)).fetchone()
    q = lambda sql: [dict(r) for r in con.execute(sql, (key,))]  # noqa: E731
    hashes = set(suppression_hashes(cand["profile_url"] or "", cand["name"], cand["company"] or "").values()) if cand else {key_hash(key)}
    sup = [dict(r) for r in con.execute(
        f"SELECT kind, created_at, reason FROM suppressed WHERE key_hash IN ({','.join('?' * len(hashes))})", tuple(hashes))]
    return {
        "dedupe_key": key,
        "candidate": dict(cand) if cand else None,
        "scores": q("SELECT jd_id, match_score, confidence, breakdown, run_id, updated_at FROM candidate_scores WHERE dedupe_key=?"),
        "feedback": q("SELECT jd_id, recruiter_status, score_override, note, created_at FROM feedback WHERE dedupe_key=?"),
        "sources": q("SELECT url, source_type, query_id, fetched_at, snippet_hash, jd_id FROM sources WHERE dedupe_key=? ORDER BY fetched_at"),
        "evidence": q("SELECT jd_id, fact, quote, url, source_type, fetched_at FROM evidence WHERE dedupe_key=? ORDER BY fetched_at"),
        "workbook_rows": [{"sheet": s_, "row": r} for s_, r in _workbook_rows(wb_path or WORKBOOK_PATH, key)],
        "suppressed": sup,
    }


def forget(con, ref: str, wb_path=None, reason: str = "") -> dict:
    """Delete the person from memory.db and every workbook sheet, then suppress them (hashes only)."""
    import os
    import tempfile
    from openpyxl import load_workbook

    wb_path = wb_path or WORKBOOK_PATH
    key = resolve_key(con, ref)
    cand = con.execute("SELECT * FROM candidates WHERE dedupe_key=?", (key,)).fetchone()
    keys = {key}
    if cand:   # the person may also be held under their name|company key
        keys.add(row_key("", cand["name"], cand["company"] or ""))
        hashes = suppression_hashes(cand["profile_url"] or "", cand["name"], cand["company"] or "")
    else:
        hashes = {"url" if key.startswith("url:") else "nc": key_hash(key)}
    removed_rows = 0
    last_removed = last_record.forget(last_record.folder(wb_path.parent) / last_record.FILE_NAME, keys)   # the app then removes the stored copy
    rows_removed = last_record.forget_rows(last_record.folder(wb_path.parent), keys)   # queued history rows; a tombstone makes the app remove stored copies
    if wb_path.exists():
        targets = {k: _workbook_rows(wb_path, k) for k in keys}
        hits = sorted({h for v in targets.values() for h in v})
        ev_hits = _evidence_sheet_rows(wb_path, keys)
        ent_hits = {s: _evidence_sheet_rows(wb_path, keys, s) for s in ENTITY_SHEETS}
        if hits or ev_hits or any(ent_hits.values()):
            wb = load_workbook(wb_path)
            for sheet in {h[0] for h in hits}:
                delete_sheet_rows(wb[sheet], [h[1] for h in hits if h[0] == sheet])
            if ev_hits:
                delete_sheet_rows(wb[EVIDENCE_SHEET], ev_hits)
            for s, rows_ in ent_hits.items():
                if rows_:
                    delete_sheet_rows(wb[s], rows_)
            fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=wb_path.parent)
            os.close(fd)
            try:
                wb.save(tmp)
                os.replace(tmp, wb_path)
            finally:
                if os.path.exists(tmp):
                    os.remove(tmp)
            removed_rows = len(hits)
    counts = {}
    for t in ("candidate_scores", "feedback", "sources", "evidence", "candidates"):
        counts[t] = sum(con.execute(f"DELETE FROM {t} WHERE dedupe_key=?", (k,)).rowcount for k in keys)
    for kind, h in hashes.items():
        con.execute("INSERT OR IGNORE INTO suppressed VALUES (?,?,?,?)", (h, kind, now(), reason or "forget request"))
    con.commit()
    leftovers = sorted(p.name for p in wb_path.parent.glob("*backup*") if p.is_file() and p != wb_path)
    return {"dedupe_key": key, "workbook_rows_deleted": removed_rows, "deleted": counts, "suppressed": sorted(hashes),
            "last_record_removed": last_removed, "history_rows_removed": rows_removed, "backups_may_still_hold_them": leftovers}


def main() -> None:
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("init")
    p = sub.add_parser("add-jd"); p.add_argument("--file", required=True)
    p = sub.add_parser("get-jd"); p.add_argument("--id", required=True)
    p = sub.add_parser("synonyms"); p.add_argument("action", choices=["add", "list"]); p.add_argument("role"); p.add_argument("syn", nargs="*")
    p = sub.add_parser("plan-queries"); p.add_argument("--jd", required=True); p.add_argument("--max", type=int, default=6)
    p.add_argument("--open-to-work", action="store_true", help="append the optional open-to-work query as an extra")
    p = sub.add_parser("run-start"); p.add_argument("--jd", required=True); p.add_argument("--queries", default="[]")
    p = sub.add_parser("run-finish"); p.add_argument("--run", type=int, required=True); p.add_argument("--notes", default="")
    p = sub.add_parser("seen"); p.add_argument("--jd", required=True)
    p = sub.add_parser("feedback"); p.add_argument("--jd", required=True); p.add_argument("--key", required=True)
    p.add_argument("--status", choices=STATUSES); p.add_argument("--score", type=int); p.add_argument("--note")
    p = sub.add_parser("sync-feedback"); p.add_argument("--jd", required=True)
    p = sub.add_parser("feedback-list"); p.add_argument("--jd", required=True)
    p = sub.add_parser("history"); p.add_argument("--jd")
    p = sub.add_parser("where"); p.add_argument("key")
    p = sub.add_parser("forget"); p.add_argument("key"); p.add_argument("--yes", action="store_true"); p.add_argument("--reason", default="")
    a = ap.parse_args()
    con = connect()
    out: object
    if a.cmd == "init":
        out = {"db": str(DB_PATH), "ok": True}
    elif a.cmd == "add-jd":
        jd = json.loads(open(a.file).read())
        out = add_jd(con, jd)
        cr = jd["criteria"]
        progress.result(f"JD '{jd['id']}' registered · sheet '{out['sheet_name']}' · "
                        f"{len(cr.get('mandatory_skills', []))} mandatory, {len(cr.get('preferred_skills', []))} preferred skills")
    elif a.cmd == "get-jd":
        out = get_jd(con, a.id)
    elif a.cmd == "synonyms":
        if a.action == "add":
            con.executemany("INSERT OR IGNORE INTO role_synonyms VALUES (?,?)", [(a.role, s) for s in a.syn]); con.commit()
        out = [r[0] for r in con.execute("SELECT synonym FROM role_synonyms WHERE role=? ORDER BY synonym", (a.role,))]
    elif a.cmd == "plan-queries":
        out = queries_for(con, a.jd, a.max)
        if a.open_to_work:
            q = open_to_work_query(get_jd(con, a.jd)["criteria"])
            if q not in out:
                out.append(q)
        for i, q in enumerate(out, 1):
            progress.sub(f"planned query {i}/{len(out)} · {q}" + (" (optional: open to work)" if q.startswith('"open to work"') else ""))
    elif a.cmd == "run-start":
        get_jd(con, a.jd)
        cur = con.execute("INSERT INTO runs (jd_id, started_at, queries_json) VALUES (?,?,?)", (a.jd, now(), a.queries)); con.commit()
        out = {"run_id": cur.lastrowid}
        progress.sub(f"run {cur.lastrowid} started for {a.jd} with {len(json.loads(a.queries))} planned queries")
    elif a.cmd == "run-finish":
        con.execute("UPDATE runs SET finished_at=?, notes=? WHERE id=?", (now(), a.notes, a.run)); con.commit()
        out = dict(con.execute("SELECT * FROM runs WHERE id=?", (a.run,)).fetchone())
        jd = get_jd(con, out["jd_id"])
        rows = con.execute("SELECT count(*) FROM candidate_scores WHERE jd_id=?", (out["jd_id"],)).fetchone()[0]
        progress.result(f"run {a.run} closed · {out['found']} processed, {out['inserted']} inserted, {out['updated']} updated")
        otw = open_to_work_counts(jd["sheet_name"])
        out["open_to_work"] = otw
        progress.summary(f"Run {a.run} · {jd['title']}\n"
                         f"rows inserted {out['inserted']}, updated {out['updated']} · {rows} candidates held for this JD\n"
                         f"open to work: {otw['Yes']} Yes, {otw['Likely']} Likely of {otw['rows']} rows in the sheet\n"
                         f"sheet '{jd['sheet_name']}' in {WORKBOOK_PATH}\n"
                         f"gaps / notes: {a.notes or 'none recorded'}")
    elif a.cmd == "seen":
        out = [dict(r) for r in con.execute(
            "SELECT c.dedupe_key, c.name, c.company, s.match_score, s.confidence, s.updated_at FROM candidate_scores s "
            "JOIN candidates c USING(dedupe_key) WHERE s.jd_id=? ORDER BY s.match_score DESC", (a.jd,))]
    elif a.cmd == "feedback":
        con.execute("INSERT INTO feedback (dedupe_key, jd_id, recruiter_status, score_override, note, created_at) VALUES (?,?,?,?,?,?)",
                    (a.key, a.jd, a.status, a.score, a.note, now())); con.commit()
        out = {"ok": True}
    elif a.cmd == "feedback-list":
        out = [dict(r) for r in con.execute(
            "SELECT f.*, c.name FROM feedback f LEFT JOIN candidates c USING(dedupe_key) WHERE f.jd_id=? ORDER BY f.id", (a.jd,))]
    elif a.cmd == "sync-feedback":
        out = {"synced": sync_feedback(con, a.jd)}
    elif a.cmd == "where":
        out = where(con, a.key)
        progress.result(f"{out['dedupe_key']}: {len(out['scores'])} scores, {len(out['feedback'])} feedback, "
                        f"{len(out['sources'])} sources, {len(out['workbook_rows'])} workbook rows"
                        + (", SUPPRESSED" if out["suppressed"] else ""))
    elif a.cmd == "forget":
        if not a.yes:
            progress.fail_exit("forget deletes the person from memory.db and the workbook for good; add --yes")
        out = forget(con, a.key, reason=a.reason)
        progress.result(f"forgot {out['dedupe_key']}: {out['workbook_rows_deleted']} workbook rows, "
                        f"{sum(out['deleted'].values())} memory rows deleted; suppressed for future runs")
        if out["backups_may_still_hold_them"]:
            progress.sub("WARNING: backup copies are not cleaned and may still hold this person: "
                         + ", ".join(out["backups_may_still_hold_them"]) + ". Delete the ones you no longer need.")
    else:
        q = "SELECT * FROM runs" + (" WHERE jd_id=?" if a.jd else "") + " ORDER BY id DESC LIMIT 20"
        out = [dict(r) for r in con.execute(q, (a.jd,) if a.jd else ())]
    print(json.dumps(out, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
