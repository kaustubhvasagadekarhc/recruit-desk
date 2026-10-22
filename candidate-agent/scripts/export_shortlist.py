"""Export one JD's shortlist as a small file to hand to a hiring manager (plan item 9, 2026-10-18).

uv run scripts/export_shortlist.py --jd <jd_id> [--statuses Shortlisted,Contacted] [--format xlsx|csv] [--out exports/<file>]

Reads the JD's sheet of Candidate_Research.xlsx (read only) and writes exports/<jd>-shortlist-<date>.xlsx (or .csv) with only
the rows whose Recruiter Status is in --statuses (default Shortlisted), best Match Score first, and the minimum columns:
name, title, company, location, experience, score, confidence, open to work, profile link, up to 3 cited facts from the
Evidence sheet, status and date found. No email or phone number is copied: contacts leave through export_crm.py, which
applies the self-published rule. Nothing is written when no row qualifies.
"""
from __future__ import annotations

import argparse
import csv
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).parent))
import memory  # noqa: E402
import progress  # noqa: E402
from common import EVIDENCE_SHEET, HEADERS, ROOT, STATUSES, WORKBOOK_PATH, header_len, row_key, structured_fields  # noqa: E402

COLUMNS = ["Name", "Title", "Company", "Location", "Experience (years)", "Match Score", "Evidence Confidence", "Open to Work",
           "Profile URL", "Cited facts", "Recruiter Status", "Date Found"]
EXPORT_DIR = ROOT / "exports"


def shown(path: Path) -> str:
    """The path as the user should see it: relative to this folder when it is inside it."""
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def safe_cell(v):
    """A text cell that starts like a formula would be run by Excel or a CSV reader; keep it as text."""
    return "'" + v if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v


def cited_facts(ev_rows: list[tuple], key: str, jd_id: str, limit: int = 3) -> str:
    """Up to `limit` 'fact: “quote”' strings from this JD's Evidence rows of one person (rows with a quote first). "Open to work"
    rows are left out, as in the outreach drafts: it is not something to hand a hiring manager as a reason."""
    mine = [r for r in ev_rows if r[0] == key and r[1] and not str(r[1]).lower().startswith("open to work") and len(r) > 7 and r[7] in (None, jd_id)]
    mine.sort(key=lambda r: (not r[2], str(r[1])))
    return " | ".join(f"{r[1]}: “{r[2]}”" if r[2] else str(r[1]) for r in mine[:limit])


def read_rows(wb_path: Path, sheet: str, statuses: set[str], jd_id: str = "") -> list[dict]:
    wb = load_workbook(wb_path, read_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        return []
    ev_rows = []
    if EVIDENCE_SHEET in wb.sheetnames:
        ev_rows = [r for r in wb[EVIDENCE_SHEET].iter_rows(min_row=2, values_only=True) if r and r[0]]
    it = wb[sheet].iter_rows(values_only=True)
    hdr = list(next(it, []) or [])
    n = header_len(hdr)
    out = []
    if n is not None:
        for vals in it:
            cells = {HEADERS[i]: vals[i] for i in range(min(n, len(vals)))}
            if not cells.get("Candidate Name") or (cells.get("Recruiter Status") or "New") not in statuses:
                continue
            derived = structured_fields({k: v for k, v in zip(["candidate_name", "current_role_company", "experience", "score_breakdown", "match_score", "location"],
                                                               [cells.get("Candidate Name"), cells.get("Current Role / Company"), cells.get("Experience"),
                                                                cells.get("Score Breakdown"), cells.get("Match Score (0-100)"), cells.get("Location")])})
            key = row_key(cells.get("Profile URL") or "", cells["Candidate Name"], cells.get("Current Role / Company") or "")
            out.append({"Name": cells["Candidate Name"], "Title": cells.get("Title") or derived["current_title"],
                        "Company": cells.get("Company") or derived["current_company"], "Location": cells.get("Location"),
                        "Experience (years)": cells.get("Experience Years") if cells.get("Experience Years") not in (None, "") else derived["experience_years"],
                        "Match Score": cells.get("Match Score (0-100)"), "Evidence Confidence": cells.get("Evidence Confidence"),
                        "Open to Work": str(cells.get("Open to Work") or "").split(":")[0].strip() or None,
                        "Profile URL": cells.get("Profile URL"), "Cited facts": cited_facts(ev_rows, key, jd_id) or None,
                        "Recruiter Status": cells.get("Recruiter Status") or "New", "Date Found": cells.get("Date Found")})
    wb.close()
    return sorted(out, key=lambda r: -(r["Match Score"] if isinstance(r["Match Score"], (int, float)) else -1))


def write(rows: list[dict], out: Path, fmt: str) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "csv":
        with out.open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(COLUMNS)
            w.writerows([[safe_cell(r[c]) if r[c] is not None else "" for c in COLUMNS] for r in rows])
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Shortlist"
    ws.append(COLUMNS)
    widths = [24, 28, 22, 24, 12, 10, 12, 12, 44, 70, 14, 12]
    for i, c in enumerate(ws[1], start=1):
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = widths[i - 1]
    for r in rows:
        ws.append([safe_cell(r[c]) for c in COLUMNS])
        link = ws.cell(ws.max_row, COLUMNS.index("Profile URL") + 1)
        if link.value and str(link.value).startswith("http"):
            link.hyperlink = link.value
            link.font = Font(color="0563C1", underline="single")
        ws.cell(ws.max_row, COLUMNS.index("Cited facts") + 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(out)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--jd", required=True)
    ap.add_argument("--statuses", default="Shortlisted")
    ap.add_argument("--format", choices=["xlsx", "csv"], default="xlsx")
    ap.add_argument("--out")
    a = ap.parse_args()
    statuses = {s.strip() for s in a.statuses.split(",") if s.strip()}
    bad = statuses - set(STATUSES)
    if bad:
        progress.fail_exit(f"unknown status {', '.join(sorted(bad))}; use {', '.join(STATUSES)}")
    jd = memory.get_jd(memory.connect(), a.jd)
    if not WORKBOOK_PATH.exists():
        progress.fail_exit(f"no workbook at {WORKBOOK_PATH}")
    rows = read_rows(WORKBOOK_PATH, jd["sheet_name"], statuses, a.jd)
    if not rows:
        progress.result(f"no rows with status {', '.join(sorted(statuses))} in '{jd['sheet_name']}'; nothing written")
        return
    out = Path(a.out) if a.out else EXPORT_DIR / f"{a.jd}-shortlist-{date.today().isoformat()}.{a.format}"
    out = out if out.is_absolute() else ROOT / out
    try:
        out.resolve().relative_to(EXPORT_DIR.resolve())
    except ValueError:
        progress.fail_exit("the shortlist is written under exports/ only")
    write(rows, out, a.format)
    progress.result(f"{len(rows)} people ({', '.join(sorted(statuses))}) from '{jd['sheet_name']}' → {shown(out)}")


if __name__ == "__main__":
    main()
