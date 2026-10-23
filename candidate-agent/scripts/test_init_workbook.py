"""Offline test: the Excel file is created when it does not exist, once, with its sheets and headers, and an existing file is never touched.

Run: uv run python scripts/test_init_workbook.py
"""
from __future__ import annotations

import os
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
results: list[tuple[str, bool]] = []


def check(name: str, ok: bool) -> None:
    results.append((name, bool(ok)))
    print(("PASS " if ok else "FAIL ") + name)


def run_init(folder: Path) -> str:
    out = subprocess.run([sys.executable, str(ROOT / "scripts" / "init_workbook.py")], cwd=ROOT, capture_output=True, text=True,
                         env={**os.environ, "RECRUIT_DESK_WORKBOOK_DIR": str(folder), "PYTHONDONTWRITEBYTECODE": "1"})
    return out.stdout + out.stderr


tmp = Path(tempfile.mkdtemp(prefix="initwb-"))
folder = tmp / "not-yet-there" / "sub"          # the folder itself does not exist either
wb_path = folder / "Candidate_Research.xlsx"
out = run_init(folder)
check("created in a folder that did not exist", wb_path.exists() and "created" in out)
wb = load_workbook(wb_path)
check("has the starting sheets", all(n in wb.sheetnames for n in ["Evidence"]) and "Sheet" not in wb.sheetnames)
check("every sheet has a header row", all(ws.cell(1, 1).value for ws in wb))
wb.close()
before = wb_path.read_bytes()
out2 = run_init(folder)
check("a second run leaves the file alone", wb_path.read_bytes() == before and "already there" in out2)
bad = [n for n, ok in results if not ok]
print(f"\n{len(results) - len(bad)}/{len(results)} passed")
sys.exit(1 if bad else 0)
