"""Weight tuning, shortlist export and outreach drafts (2026-10-18, plan item 9). Temp copies only."""
from __future__ import annotations

import contextlib
import csv
import io
import json
import sqlite3
import sys
from unittest import mock

from openpyxl import load_workbook

from helpers import JD, TempTestCase

import common
import entities
import excel_upsert
import export_shortlist
import memory
import outreach
import progress
import tune_weights
from test_structured_crm import scored

W = {"core_skills": 35, "experience": 20, "domain_projects": 15, "seniority": 10, "location": 10, "preferred_skills": 10}


def share(**over):
    base = {d: 0.5 for d in W}
    base.update(over)
    return base


class Analyse(TempTestCase):
    def test_moves_points_from_the_dimension_that_tells_nothing_to_the_one_that_does(self):
        pos = [share(core_skills=0.9, location=0.5)] * 6
        neg = [share(core_skills=0.4, location=0.5)] * 6
        basis, new = tune_weights.analyse(pos, neg, W)
        self.assertEqual(basis["core_skills"]["gap"], 0.5)
        self.assertEqual(sum(new.values()), 100)
        self.assertEqual(new["core_skills"] - W["core_skills"], 5)       # at most 5 points move
        self.assertEqual(sum(W[d] - new[d] for d in W if new[d] < W[d]), 5)
        self.assertTrue(all(new[d] >= tune_weights.FLOOR for d in new))

    def test_no_change_when_nothing_separates_the_groups(self):
        _, new = tune_weights.analyse([share()] * 6, [share()] * 6, W)
        self.assertEqual(new, W)

    def test_no_change_without_a_receiver_or_without_a_donor(self):
        pos = [share(core_skills=0.9)] * 6                                  # one receiver, but every other gap is 0.. donors exist
        _, new = tune_weights.analyse(pos, [share(core_skills=0.5)] * 6, W)
        self.assertGreater(new["core_skills"], W["core_skills"])
        everyone_tells = [share(**{d: 0.9 for d in W})] * 6                 # every gap >= 0.10: nobody can donate
        _, same = tune_weights.analyse(everyone_tells, [share(**{d: 0.4 for d in W})] * 6, W)
        self.assertEqual(same, W)

    def test_floor_is_kept(self):
        tight = {"core_skills": 35, "experience": 20, "domain_projects": 15, "seniority": 6, "location": 6, "preferred_skills": 18}
        pos = [share(core_skills=0.9)] * 6
        _, new = tune_weights.analyse(pos, [share(core_skills=0.4)] * 6, tight)
        self.assertTrue(all(v >= tune_weights.FLOOR for v in new.values()))
        self.assertEqual(sum(new.values()), 100)

    def test_max_shift_is_capped_at_five(self):
        pos = [share(core_skills=0.9)] * 6
        _, new = tune_weights.analyse(pos, [share(core_skills=0.4)] * 6, W, max_shift=50)
        self.assertEqual(new["core_skills"] - W["core_skills"], 5)


class DbCase(TempTestCase):
    def setUp(self):
        super().setUp()
        self.db, self.wb = self.tmp / "memory.db", self.tmp / "wb.xlsx"
        self.exports = self.tmp / "exports"
        self._p = [mock.patch.object(m, n, v) for m, n, v in (
            (excel_upsert, "DB_PATH", self.db), (excel_upsert, "WORKBOOK_PATH", self.wb), (memory, "DB_PATH", self.db),
            (memory, "WORKBOOK_PATH", self.wb), (common, "DB_PATH", self.db), (entities, "DB_PATH", self.db),
            (entities, "WORKBOOK_PATH", self.wb), (export_shortlist, "WORKBOOK_PATH", self.wb), (outreach, "WORKBOOK_PATH", self.wb),
            (export_shortlist, "EXPORT_DIR", self.exports), (outreach, "EXPORT_DIR", self.exports), (tune_weights, "ROOT", self.tmp))]
        for p in self._p:
            p.start()
        self.con = memory.connect()
        memory.add_jd(self.con, JD)
        self.con.commit()

    def tearDown(self):
        self.con.close()
        for p in self._p:
            p.stop()
        super().tearDown()

    def run_main(self, mod, *argv):
        with mock.patch.object(sys, "argv", [mod.__name__ + ".py", *argv]):
            mod.main()


class Tuning(DbCase):
    def seed(self, n_pos=6, n_neg=6):
        def person(i, status, core, loc):
            key = f"url:https://x.test/{status}{i}"
            self.con.execute("INSERT INTO candidates VALUES (?,?,?,?,?,?)", (key, f"P{status}{i}", "Acme", key[4:], "2026-10-15", "2026-10-15"))
            br = f"Core skills {core}/35 · Experience 10/20 · Domain 8/15 · Seniority 5/10 · Location {loc}/10 · Preferred 5/10"
            self.con.execute("INSERT INTO candidate_scores VALUES (?,?,?,?,?,?,?)", (key, JD["id"], 60, br, "Medium", None, "2026-10-15"))
            self.con.execute("INSERT INTO feedback (dedupe_key, jd_id, recruiter_status, note, created_at) VALUES (?,?,?,?,?)",
                             (key, JD["id"], status, "", "2026-10-16"))
        for i in range(n_pos):
            person(i, "Shortlisted", 32, 6)
        for i in range(n_neg):
            person(i, "Rejected", 14, 6)
        self.con.commit()

    def test_collect_groups_by_latest_status(self):
        self.seed(2, 2)
        key = "url:https://x.test/Rejected0"
        self.con.execute("INSERT INTO feedback (dedupe_key, jd_id, recruiter_status, note, created_at) VALUES (?,?,?,?,?)",
                         (key, JD["id"], "Shortlisted", "", "2026-10-17"))   # changed their mind: now a positive
        self.con.commit()
        pos, neg = tune_weights.collect(self.con, JD["id"], JD["weights"])
        self.assertEqual((len(pos), len(neg)), (3, 1))
        self.assertAlmostEqual(pos[0]["core_skills"], 32 / 35)

    def test_too_little_feedback_changes_nothing(self):
        self.seed(4, 6)
        self.run_main(tune_weights, "--jd", JD["id"], "--apply")
        self.assertEqual(memory.get_jd(self.con, JD["id"])["weights"], JD["weights"])
        self.assertEqual(self.con.execute("SELECT count(*) FROM weight_changes").fetchone()[0], 0)

    def test_dry_run_prints_and_saves_nothing(self):
        self.seed()
        self.run_main(tune_weights, "--jd", JD["id"])
        self.assertEqual(memory.get_jd(self.con, JD["id"])["weights"], JD["weights"])
        self.assertEqual(self.con.execute("SELECT count(*) FROM weight_changes").fetchone()[0], 0)

    def test_apply_saves_weights_the_log_and_the_jd_file(self):
        self.seed()
        (self.tmp / "jds").mkdir()
        jd_file = self.tmp / "jds" / f"{JD['id']}.json"
        jd_file.write_text(json.dumps({"id": JD["id"], "weights": JD["weights"]}))
        self.run_main(tune_weights, "--jd", JD["id"], "--apply")
        w = memory.get_jd(self.con, JD["id"])["weights"]
        self.assertEqual(sum(w.values()), 100)
        self.assertGreater(w["core_skills"], JD["weights"]["core_skills"])
        self.assertEqual(json.loads(jd_file.read_text())["weights"], w)   # a later add-jd must not undo it
        row = self.con.execute("SELECT before_json, after_json, n_positive, n_negative FROM weight_changes").fetchone()
        self.assertEqual((json.loads(row[0]), json.loads(row[1]), row[2], row[3]), (JD["weights"], w, 6, 6))

    def test_forgetting_a_person_removes_their_feedback_from_the_next_suggestion(self):
        self.seed()
        memory.forget(self.con, "url:https://x.test/Shortlisted0")
        pos, _ = tune_weights.collect(self.con, JD["id"], JD["weights"])
        self.assertEqual(len(pos), 5)


class Workbook_(DbCase):
    URL = "https://www.linkedin.com/in/test-person-1"

    def person(self, **over):
        base = dict(scores={"core_skills": 30, "domain_projects": 0, "seniority": 0, "preferred_skills": 0},
                    quotes={"core_skills": ["React | TypeScript"]}, facts=[{"fact": "7+ YOE", "source": "LinkedIn headline"}],
                    evidence_notes="Headline: 'Senior Frontend Engineer | React | TypeScript | 7+ YOE'. #OpenToWork",
                    email_mobile="a@b.com", contact_source_url="https://x.test")
        base.update(over)
        return scored(**base)

    def upsert(self, cands, *extra):
        p = self.write_json("b.json", {"jd_id": JD["id"], "candidates": cands})
        with mock.patch.object(sys, "argv", ["excel_upsert.py", "--input", str(p), *extra]):
            excel_upsert.main()

    def set_status(self, name, status):
        wb = load_workbook(self.wb)
        ws = wb[JD["title"]]
        h = [c.value for c in ws[1]]
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, h.index("Candidate Name") + 1).value == name:
                ws.cell(r, h.index("Recruiter Status") + 1, status)
        wb.save(self.wb)

    def seed(self):
        self.upsert([self.person(), self.person(candidate_name="Other One", profile_url="https://www.linkedin.com/in/other-one")], "--contacts")
        self.set_status("Test Person", "Shortlisted")


class Shortlist(Workbook_):
    def test_only_shortlisted_rows_and_no_contact_data(self):
        self.seed()
        self.run_main(export_shortlist, "--jd", JD["id"])
        f = next(self.exports.glob("*-shortlist-*.xlsx"))
        vals = list(load_workbook(f)["Shortlist"].iter_rows(values_only=True))
        self.assertEqual(list(vals[0]), export_shortlist.COLUMNS)
        self.assertEqual(len(vals) - 1, 1)
        row = dict(zip(vals[0], vals[1]))
        self.assertEqual((row["Name"], row["Title"], row["Company"], row["Recruiter Status"]), ("Test Person", "Senior Frontend Engineer", "Acme", "Shortlisted"))
        self.assertIn("Core skills: “React | TypeScript”", row["Cited facts"])
        self.assertFalse(any("a@b.com" in str(v) for v in vals[1]))

    def test_statuses_csv_and_best_score_first(self):
        self.seed()
        self.set_status("Other One", "Contacted")
        self.run_main(export_shortlist, "--jd", JD["id"], "--statuses", "Shortlisted,Contacted", "--format", "csv")
        rows = list(csv.DictReader(next(self.exports.glob("*.csv")).open(encoding="utf-8")))
        self.assertEqual(len(rows), 2)
        self.assertGreaterEqual(float(rows[0]["Match Score"]), float(rows[1]["Match Score"]))

    def test_nothing_written_when_no_row_qualifies(self):
        self.seed()
        self.run_main(export_shortlist, "--jd", JD["id"], "--statuses", "Rejected")
        self.assertFalse(self.exports.exists() and list(self.exports.iterdir()))

    def test_out_must_be_under_exports(self):
        self.seed()
        with self.assertRaises(SystemExit):
            self.run_main(export_shortlist, "--jd", JD["id"], "--out", str(self.tmp / "elsewhere.xlsx"))

    def test_unknown_status_refused(self):
        self.seed()
        with self.assertRaises(SystemExit):
            self.run_main(export_shortlist, "--jd", JD["id"], "--statuses", "Hired")


class Outreach(Workbook_):
    def plan(self):
        people = outreach.plan(self.wb, JD["title"], JD["id"], {"Shortlisted"})
        self.assertEqual(len(people), 1)
        return people[0]

    def write(self, drafts):
        f = self.write_json("d.json", {"drafts": drafts})
        self.run_main(outreach, "write", "--jd", JD["id"], "--company", "Acme Corp", "--input", str(f))

    def test_plan_lists_citable_facts_without_open_to_work(self):
        self.seed()
        p = self.plan()
        self.assertEqual(p["name"], "Test Person")
        facts = {f["fact"] for f in p["facts"]}
        self.assertIn("Core skills", facts)
        self.assertFalse(any(x.lower().startswith("open to work") for x in facts))
        self.assertEqual(len({f["id"] for f in p["facts"]}), len(p["facts"]))
        self.assertEqual([f["id"] for f in p["facts"]], [f["id"] for f in self.plan()["facts"]])   # ids are stable

    def test_valid_draft_is_written_with_the_notice_appended(self):
        self.seed()
        p = self.plan()
        fact = next(f for f in p["facts"] if f["fact"] == "Core skills")
        self.write([{"key": p["key"], "fact_id": fact["id"],
                     "message": "Hi Test, your work with React | TypeScript caught our eye. We are hiring a Senior React Frontend Developer in Pune. Open to a quick chat?"}])
        text = next(self.exports.glob("*-outreach-*.md")).read_text(encoding="utf-8")
        self.assertIn("## Test Person", text)
        self.assertIn("Nothing has been sent", text)
        self.assertIn(common.art14_notice(p["source"], p["date_found"], JD["title"], "Acme Corp"), text)

    def refused(self, drafts, needle):
        err = io.StringIO()
        with contextlib.redirect_stderr(err), self.assertRaises(SystemExit):
            self.write(drafts)
        self.assertIn(needle, err.getvalue())
        self.assertFalse(self.exports.exists() and list(self.exports.iterdir()), "nothing may be written")

    def test_refusals(self):
        self.seed()
        p = self.plan()
        fact = next(f for f in p["facts"] if f["fact"] == "Core skills")
        ok = "Your React | TypeScript work caught our eye."
        self.refused([{"key": "url:https://nobody", "fact_id": fact["id"], "message": ok}], "is not on the plan")
        self.refused([{"key": p["key"], "fact_id": "deadbeef", "message": ok}], "is not one of this person")
        self.refused([{"key": p["key"], "fact_id": fact["id"], "message": "Hello, we have a role for you."}], "must repeat the cited quote")
        self.refused([{"key": p["key"], "fact_id": fact["id"], "message": ok + " Write to me at a@b.com"}], "email address")
        self.refused([{"key": p["key"], "fact_id": fact["id"], "message": ok + " Call +91 98765 43210"}], "phone number")
        self.refused([{"key": p["key"], "fact_id": fact["id"], "message": ok + " x" * 600}], "at most 900")
        self.refused([], "no drafts")

    def test_people_who_are_not_shortlisted_get_no_plan(self):
        self.seed()
        names = {p["name"] for p in outreach.plan(self.wb, JD["title"], JD["id"], {"Shortlisted"})}
        self.assertEqual(names, {"Test Person"})


class ReviewFixes(Workbook_):
    """Findings of the independent review (2026-10-18), each pinned by a test."""

    def test_phone_and_email_detection(self):
        for text in ("Call +91 98765 43210", "Ring 415.555.0100 any time", "mail jane@example.com", "write jane at example dot com"):
            self.assertTrue(outreach.has_contact(text), text)
        for text in ("You worked there 2019-2023 as a lead", "From 2019 - 2023 at Acme", "Series B, $75M round"):
            self.assertFalse(outreach.has_contact(text), text)

    def test_short_quote_duplicates_bad_input_and_no_overwrite(self):
        self.seed()
        p = Outreach.plan(self)
        fact = next(f for f in p["facts"] if f["fact"] == "Core skills")
        good = {"key": p["key"], "fact_id": fact["id"], "message": "Your React | TypeScript work caught our eye. Open to a chat?"}
        short = dict(good, fact_id=next(f for f in p["facts"] if f["fact"] == "7+ YOE")["id"], message="Hello 7+ YOE")
        for drafts, needle in (([short], "proves nothing"), ([good, good], "more than one draft"), (["not an object"], "must be an object")):
            err = io.StringIO()
            with contextlib.redirect_stderr(err), self.assertRaises(SystemExit):
                Outreach.write(self, drafts)
            self.assertIn(needle, err.getvalue())
        Outreach.write(self, [good])
        Outreach.write(self, [good])   # same day again: the first file must survive
        self.assertEqual(len(list(self.exports.glob("*-outreach-*.md"))), 2)

    def test_plan_marks_short_facts_unusable(self):
        self.seed()
        facts = Outreach.plan(self)["facts"]
        self.assertTrue(all("usable" in f for f in facts))
        self.assertFalse(next(f for f in facts if f["fact"] == "7+ YOE")["usable"])

    def test_shortlist_leaves_out_open_to_work_and_neutralises_formulas(self):
        self.upsert([self.person(candidate_name="=HYPERLINK(1)", profile_url="https://www.linkedin.com/in/formula-guy")], "--contacts")
        self.set_status("=HYPERLINK(1)", "Shortlisted")
        self.run_main(export_shortlist, "--jd", JD["id"])
        vals = list(load_workbook(next(self.exports.glob("*.xlsx")))["Shortlist"].iter_rows(values_only=True))
        row = dict(zip(vals[0], vals[1]))
        self.assertTrue(row["Name"].startswith("'="))
        self.assertNotIn("Open to work", row["Cited facts"] or "")

    def test_shortlist_ignores_evidence_of_another_jd(self):
        self.seed()
        key = f"url:{self.URL}"
        wb = load_workbook(self.wb)
        wb["Evidence"].append([key, "Core skills", "A quote from another JD", "https://x.test/other", "Web", "2026-10-16", "Test Person", "another-jd"])
        wb.save(self.wb)
        self.run_main(export_shortlist, "--jd", JD["id"])
        vals = list(load_workbook(next(self.exports.glob("*.xlsx")))["Shortlist"].iter_rows(values_only=True))
        self.assertNotIn("another JD", str(vals[1]))

    def test_entity_sheets_are_marked_read_only(self):
        self.seed()
        for name in ("Candidates", "Matches"):
            ws = load_workbook(self.wb)[name]
            self.assertTrue(ws.protection.sheet)
            self.assertFalse(ws.protection.autoFilter or ws.protection.sort)

    def test_a_jd_sheet_named_like_an_entity_sheet_is_refused_not_deleted(self):
        self.seed()
        wb = load_workbook(self.wb)
        del wb["Matches"]
        ws = wb.create_sheet("Matches")
        ws.append(["Role", "Pay"])
        ws.append(["Chef", 10])
        wb.save(self.wb)
        with self.assertRaises(SystemExit):
            entities.migrate(self.wb, self.db, stamp="x")
        self.assertEqual(load_workbook(self.wb)["Matches"]["A2"].value, "Chef")

    def test_entities_dry_run_creates_no_database(self):
        self.seed()
        gone = self.tmp / "other.db"
        entities.migrate(self.wb, gone, dry_run=True)
        self.assertFalse(gone.exists())

    def test_forget_warns_about_backups(self):
        self.seed()
        (self.tmp / "wb.backup-2026-10-18-before-entities.xlsx").write_bytes(b"x")
        out = memory.forget(self.con, self.URL)
        self.assertEqual(out["backups_may_still_hold_them"], ["wb.backup-2026-10-18-before-entities.xlsx"])


class TuningFixes(Tuning):
    def test_a_score_correction_without_status_does_not_hide_a_shortlisted_person(self):
        self.seed(2, 2)
        key = "url:https://x.test/Shortlisted0"
        self.con.execute("INSERT INTO feedback (dedupe_key, jd_id, recruiter_status, score_override, note, created_at) VALUES (?,?,?,?,?,?)",
                         (key, JD["id"], None, 70, "score fix", "2026-10-17"))
        self.con.commit()
        pos, _ = tune_weights.collect(self.con, JD["id"], JD["weights"])
        self.assertEqual(len(pos), 2)

    def test_min_each_zero_is_refused(self):
        with self.assertRaises(SystemExit):
            self.run_main(tune_weights, "--jd", JD["id"], "--min-each", "0")

    def test_reverting_a_status_to_new_stops_counting_as_shortlisted(self):
        self.seed(1, 0)
        self.con.execute("INSERT INTO candidates VALUES (?,?,?,?,?,?)", ("url:https://x.test/rev", "Rev", "Acme", "https://x.test/rev", "d", "d"))
        self.con.commit()
        (self.tmp / "x").mkdir()
        # Recruiter Status edits in the workbook: Shortlisted, then back to New
        from test_structured_crm import scored as _scored
        cand = _scored(candidate_name="Rev", profile_url="https://x.test/rev")
        p = self.write_json("b.json", {"jd_id": JD["id"], "candidates": [cand]})
        with mock.patch.object(sys, "argv", ["excel_upsert.py", "--input", str(p)]):
            excel_upsert.main()

        def set_status(v):
            wb = load_workbook(self.wb)
            ws = wb[JD["title"]]
            h = [c.value for c in ws[1]]
            ws.cell(2, h.index("Recruiter Status") + 1, v)
            wb.save(self.wb)
        for status in ("Shortlisted", "New"):
            set_status(status)
            memory.sync_feedback(self.con, JD["id"])
        key = common.row_key("https://x.test/rev", "Rev", "")
        last = self.con.execute("SELECT recruiter_status FROM feedback WHERE dedupe_key=? ORDER BY id DESC LIMIT 1", (key,)).fetchone()[0]
        self.assertEqual(last, "New")
