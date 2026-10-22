"""memory.py (JD registry, synonyms, plan-queries) and excel_upsert idempotence on temp copies."""
from __future__ import annotations

import json
import sqlite3
import sys
from unittest import mock

from openpyxl import load_workbook

from helpers import JD, TempTestCase, candidate

import common
import excel_upsert
import memory


class MemoryDb(TempTestCase):
    def setUp(self):
        super().setUp()
        self.db = self.tmp / "memory.db"
        self._p = mock.patch.object(memory, "DB_PATH", self.db)
        self._p.start()
        self.con = memory.connect()

    def tearDown(self):
        self.con.close()
        self._p.stop()
        super().tearDown()

    def test_add_jd_is_an_upsert_and_stores_synonyms(self):
        a = memory.add_jd(self.con, JD)
        b = memory.add_jd(self.con, dict(JD, title=JD["title"]))
        self.assertEqual(a["sheet_name"], b["sheet_name"])
        self.assertEqual(self.con.execute("SELECT count(*) FROM jds").fetchone()[0], 1)
        syns = [r[0] for r in self.con.execute("SELECT synonym FROM role_synonyms WHERE role=?", (JD["title"],))]
        self.assertEqual(sorted(syns), sorted(JD["criteria"]["role_synonyms"]))

    def test_bad_weights_rejected(self):
        with self.assertRaises(SystemExit):
            memory.add_jd(self.con, dict(JD, weights={"core_skills": 50}))

    def test_queries_use_stored_synonyms(self):
        memory.add_jd(self.con, JD)
        self.con.execute("INSERT INTO role_synonyms VALUES (?,?)", (JD["title"], "Frontend Lead"))
        qs = memory.queries_for(self.con, JD["id"])
        self.assertIn("site:linkedin.com/in Frontend Lead Pune React", qs)

    def test_plan_queries_cli_prints_each_query(self):
        memory.add_jd(self.con, JD)
        self.con.commit()
        with mock.patch.object(sys, "argv", ["memory.py", "plan-queries", "--jd", JD["id"]]):
            memory.main()
        out = self._out._new_target.getvalue()
        self.assertIn("planned query 1/", out)
        self.assertIn("category:people Senior React Frontend Developer Pune", out)


class Upsert(TempTestCase):
    def setUp(self):
        super().setUp()
        self.db, self.wb = self.tmp / "memory.db", self.tmp / "book.xlsx"
        self._ps = [mock.patch.object(memory, "DB_PATH", self.db), mock.patch.object(excel_upsert, "DB_PATH", self.db),
                    mock.patch.object(excel_upsert, "WORKBOOK_PATH", self.wb)]
        for p in self._ps:
            p.start()
        con = memory.connect()
        memory.add_jd(con, JD)
        con.close()

    def tearDown(self):
        for p in self._ps:
            p.stop()
        super().tearDown()

    def batch(self, *cands):
        rows = []
        for c in cands:
            c = dict(c, match_score=80, score_breakdown="Core skills 30/35")
            rows.append(c)
        return self.write_json("batch.json", {"jd_id": JD["id"], "candidates": rows})

    def upsert(self, path, *extra):
        buf = self._out._new_target
        buf.seek(0)
        buf.truncate()
        with mock.patch.object(sys, "argv", ["excel_upsert.py", "--input", str(path), *extra]):
            excel_upsert.main()
        text = buf.getvalue()
        return json.loads(text[text.index("\n{") + 1:])

    def sheet(self):
        ws = load_workbook(self.wb)[JD["title"]]
        return ws, {h: i for i, h in enumerate(common.HEADERS)}

    def test_rerun_updates_without_duplicating_and_keeps_recruiter_status(self):
        a = candidate(profile_url="https://www.linkedin.com/in/a-1", candidate_name="A One")
        b = candidate(profile_url="https://github.com/b-two", candidate_name="B Two")
        p = self.batch(a, b)
        r1 = self.upsert(p)
        self.assertEqual((r1["inserted"], r1["updated"], r1["rows_in_sheet"]), (2, 0, 2))
        wb = load_workbook(self.wb)
        ws = wb[JD["title"]]
        col = common.HEADERS.index("Recruiter Status") + 1
        ws.cell(2, col, "Shortlisted")
        wb.save(self.wb)
        a2 = dict(a, profile_url="https://in.linkedin.com/in/A-1/?trk=x", key_skills="React, Redux")
        r2 = self.upsert(self.batch(a2, b))
        self.assertEqual((r2["inserted"], r2["updated"], r2["rows_in_sheet"]), (0, 2, 2))
        ws, h = self.sheet()
        self.assertEqual(ws.cell(2, h["Recruiter Status"] + 1).value, "Shortlisted")
        self.assertEqual(ws.cell(2, h["Key Skills"] + 1).value, "React, Redux")
        n = sqlite3.connect(self.db).execute("SELECT count(*) FROM candidate_scores WHERE jd_id=?", (JD["id"],)).fetchone()[0]
        self.assertEqual(n, 2)

    def test_dry_run_saves_nothing(self):
        r = self.upsert(self.batch(candidate()), "--dry-run")
        self.assertTrue(r["dry_run"])
        self.assertFalse(self.wb.exists())

    def test_invalid_batch_writes_nothing(self):
        bad = self.write_json("bad.json", {"jd_id": JD["id"], "candidates": [candidate(profile_url="", match_score=140)]})
        with mock.patch.object(sys, "argv", ["excel_upsert.py", "--input", str(bad)]), \
             mock.patch.object(sys, "stderr"):
            with self.assertRaises(SystemExit):
                excel_upsert.main()
        self.assertFalse(self.wb.exists())
