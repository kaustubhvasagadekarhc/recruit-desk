"""Plan items 6 (evidence depth, honest confidence) and 10 (privacy, audit trail), on temp copies only."""
from __future__ import annotations

import io
import json
import sqlite3
import sys
from contextlib import redirect_stdout
from datetime import date, timedelta
from unittest import mock

from openpyxl import Workbook, load_workbook

from helpers import JD, TempTestCase, candidate

import common
import dedupe
import evidence
import excel_upsert
import memory
import retention
import score

W, CRIT = JD["weights"], JD["criteria"]
GH = "https://github.com/test-person"


def corroboration(**over) -> dict:
    c = {"url": GH, "source_type": "GitHub", "fetched_at": "2026-10-15T08:00:00+00:00",
         "snippet": "test-person · 40 repositories · Most used: TypeScript, React. Pinned: react-table-kit",
         "confirms": [{"fact": "40 public repositories", "quote": "40 repositories"},
                      {"fact": "React and TypeScript work", "quote": "TypeScript, React"}]}
    c.update(over)
    return c


class Evidence(TempTestCase):
    def test_single_source_high_becomes_medium_and_caps_at_85(self):
        c = candidate(evidence_confidence="High", scores={"core_skills": 35, "domain_projects": 15, "seniority": 10, "preferred_skills": 10})
        score.score(c, W, CRIT)
        self.assertEqual(c["evidence_confidence"], "Medium")
        self.assertEqual(c["confidence_claimed"], "High")
        self.assertEqual(c["evidence_sources"], 1)
        self.assertEqual(c["match_score"], 85)
        self.assertIn("Evidence Sources 1", c["score_breakdown"])
        self.assertIn("Medium confidence (capped at 85)", c["score_breakdown"])

    def test_corroborated_medium_becomes_high_and_is_not_capped(self):
        c = candidate(scores={"core_skills": 35, "domain_projects": 15, "seniority": 10, "preferred_skills": 10},
                      facts=[{"fact": "7+ YOE", "source": "LinkedIn headline"}], corroboration=corroboration())
        score.score(c, W, CRIT)
        self.assertEqual((c["evidence_confidence"], c["evidence_sources"]), ("High", 2))
        self.assertGreater(c["match_score"], 85)
        st = {f["fact"]: (f["status"], f["source"]) for f in c["facts"]}
        self.assertEqual(st["7+ YOE"], ("claimed", "LinkedIn headline"))
        self.assertEqual(st["40 public repositories"], ("corroborated", f"GitHub: {GH}"))
        self.assertIn("corroborated (GitHub: https://github.com/test-person): 40 public repositories", c["evidence_notes"])

    def test_rescoring_is_idempotent(self):
        c = candidate(corroboration=corroboration(), facts=["7+ YOE"])
        score.score(c, W, CRIT)
        first = (c["match_score"], c["score_breakdown"], c["evidence_notes"])
        score.score(c, W, CRIT)
        self.assertEqual(first, (c["match_score"], c["score_breakdown"], c["evidence_notes"]))

    def test_quote_missing_from_snippet_stops(self):
        c = candidate(corroboration=corroboration(confirms=[{"fact": "x", "quote": "500 stars"}]))
        with self.assertRaises(SystemExit):
            score.score(c, W, CRIT)

    def test_linkedin_or_same_host_corroboration_rejected(self):
        for url in ("https://in.linkedin.com/in/test-person-1/details", "https://www.indeed.com/r/x"):
            with self.assertRaises(SystemExit):
                score.score(candidate(corroboration=corroboration(url=url)), W, CRIT)
        with self.assertRaises(SystemExit):
            score.score(candidate(profile_url=GH, corroboration=corroboration(url=GH + "?tab=repositories")), W, CRIT)

    def test_only_one_fetched_source_per_person(self):
        with self.assertRaises(SystemExit):
            score.score(candidate(corroboration=[corroboration(), corroboration(url="https://me.dev")]), W, CRIT)

    def test_low_stays_low_and_caps_at_60(self):
        c = candidate(evidence_confidence="Low", corroboration=corroboration())
        score.score(c, W, CRIT)
        self.assertEqual((c["evidence_confidence"], c["match_score"]), ("Low", 60))

    def test_plan_picks_top_n_and_one_public_url_each(self):
        cands = [candidate(candidate_name=f"P{i}", match_score=50 + i,
                           evidence_notes=f"Snippet. Site https://p{i}.dev and https://www.linkedin.com/in/p{i}",
                           public_urls=[f"https://github.com/p{i}"] if i % 2 else [])
                 for i in range(12)]
        cands.append(candidate(candidate_name="NoUrl", match_score=99))
        rows = evidence.plan(cands, top=10)
        self.assertEqual(len(rows), 10)
        self.assertEqual(rows[0]["candidate"], "NoUrl")
        self.assertIsNone(rows[0]["fetch"])
        by = {r["candidate"]: r for r in rows}
        self.assertEqual(by["P11"]["fetch"], "https://github.com/p11")   # GitHub ranked above a personal site
        self.assertEqual(by["P10"]["fetch"], "https://p10.dev")
        self.assertNotIn("P0", by)
        for r in rows:
            self.assertNotIn("linkedin", r["fetch"] or "")


class Privacy(TempTestCase):
    def setUp(self):
        super().setUp()
        self.db, self.wb = self.tmp / "memory.db", self.tmp / "book.xlsx"
        self._ps = [mock.patch.object(m, "DB_PATH", self.db) for m in (memory, excel_upsert, dedupe, common, score)]
        self._ps += [mock.patch.object(m, "WORKBOOK_PATH", self.wb) for m in (memory, excel_upsert)]
        for p in self._ps:
            p.start()
        self.con = memory.connect()
        memory.add_jd(self.con, JD)

    def tearDown(self):
        self.con.close()
        for p in self._ps:
            p.stop()
        super().tearDown()

    def upsert(self, cands, *extra):
        rows = [dict(c, match_score=c.get("match_score", 80), score_breakdown="Core skills 30/35") for c in cands]
        path = self.write_json("batch.json", {"jd_id": JD["id"], "candidates": rows})
        with mock.patch.object(sys, "argv", ["excel_upsert.py", "--input", str(path), *extra]):
            excel_upsert.main()

    def sheet(self):
        ws = load_workbook(self.wb)[JD["title"]]
        return ws, {h: i + 1 for i, h in enumerate(common.HEADERS)}

    def test_sources_table_and_where_trail(self):
        a = candidate(query_id="run1-q2", corroboration=corroboration(), evidence_sources=2)
        self.upsert([a])
        self.upsert([a])   # rerun: no duplicate trail rows
        key = common.row_key(a["profile_url"], a["candidate_name"], a["current_role_company"])
        out = memory.where(self.con, a["profile_url"])
        self.assertEqual(out["dedupe_key"], key)
        self.assertEqual(len(out["sources"]), 2)
        prim = next(s for s in out["sources"] if s["url"] == a["profile_url"])
        self.assertEqual((prim["query_id"], prim["source_type"]), ("run1-q2", "fixture"))
        self.assertEqual(prim["snippet_hash"], common.snippet_hash(a["evidence_notes"]))
        self.assertEqual(out["workbook_rows"], [{"sheet": JD["title"], "row": 2}])
        ws, h = self.sheet()
        self.assertEqual(ws.cell(1, h["Evidence Sources"]).value, "Evidence Sources")
        self.assertEqual(ws.cell(2, h["Evidence Sources"]).value, 2)

    def test_forget_erases_everywhere_and_suppresses(self):
        a = candidate(candidate_name="A One", profile_url="https://www.linkedin.com/in/a-1")
        b = candidate(candidate_name="B Two", profile_url="https://github.com/b-two")
        self.upsert([a, b])
        self.con.execute("INSERT INTO feedback (dedupe_key, jd_id, recruiter_status, created_at) VALUES (?,?,?,?)",
                         (common.row_key(a["profile_url"], "A One", a["current_role_company"]), JD["id"], "Reviewed", "x"))
        self.con.commit()
        out = memory.forget(self.con, "https://in.linkedin.com/in/A-1/?trk=x")
        self.assertEqual(out["workbook_rows_deleted"], 1)
        self.assertEqual(out["deleted"]["candidates"], 1)
        self.assertEqual(out["deleted"]["feedback"], 1)
        ws, h = self.sheet()
        self.assertEqual([ws.cell(r, h["Candidate Name"]).value for r in range(2, ws.max_row + 1)], ["B Two"])
        self.assertEqual(ws.cell(2, h["Profile URL"]).hyperlink.target, "https://github.com/b-two")   # link moved with its row
        trail = memory.where(self.con, a["profile_url"])
        self.assertIsNone(trail["candidate"])
        self.assertEqual((trail["scores"], trail["sources"]), ([], []))
        self.assertTrue(trail["suppressed"])
        names = [r[0] for r in self.con.execute("SELECT name FROM candidates")]
        self.assertEqual(names, ["B Two"])
        self.assertNotIn("A One", json.dumps([list(r) for r in self.con.execute("SELECT * FROM suppressed")]))
        # a later run never re-adds the person: not by URL variant, and not by name+company without a URL
        self.upsert([dict(a, profile_url="https://www.linkedin.com/in/a-1/"), b])
        ws, h = self.sheet()
        self.assertEqual(ws.max_row, 2)

    def test_dedupe_and_score_drop_suppressed_before_scoring(self):
        a = candidate(candidate_name="A One", profile_url="https://github.com/a-one")
        self.upsert([a])
        memory.forget(self.con, "url:https://github.com/a-one")
        b = candidate(candidate_name="B Two", profile_url="https://github.com/b-two")
        same_person_other_url = dict(a, profile_url="https://a-one.dev")   # name|company still matches
        path = self.write_json("d.json", {"jd_id": JD["id"], "candidates": [a, b, same_person_other_url]})
        with mock.patch.object(sys, "argv", ["dedupe.py", "--input", str(path)]):
            dedupe.main()
        kept = json.loads(path.read_text())["candidates"]
        self.assertEqual([c["candidate_name"] for c in kept], ["B Two"])
        path = self.write_json("s.json", {"jd_id": JD["id"], "candidates": [a, b]})
        with mock.patch.object(sys, "argv", ["score.py", "--input", str(path)]):
            score.main()
        self.assertEqual([c["candidate_name"] for c in json.loads(path.read_text())["candidates"]], ["B Two"])

    def test_forget_cli_needs_yes(self):
        with mock.patch.object(sys, "argv", ["memory.py", "forget", "url:https://github.com/x"]):
            with self.assertRaises(SystemExit):
                memory.main()

    def test_contacts_blank_unless_flag(self):
        c = candidate(email_mobile="a.one@acme-mail.in", contact_source_url="https://a-one.dev/contact")
        self.upsert([c])
        ws, h = self.sheet()
        self.assertIsNone(ws.cell(2, h["Email / Mobile"]).value)
        self.upsert([c], "--contacts")
        ws, h = self.sheet()
        self.assertEqual(ws.cell(2, h["Email / Mobile"]).value, "a.one@acme-mail.in")

    def test_old_sheet_gets_evidence_sources_column(self):
        wb = Workbook()
        ws = wb.active
        ws.title = JD["title"]
        ws.append(common.BASE_HEADERS)
        old = candidate(candidate_name="Old Row")
        ws.append([old.get(f, "") or None for f, _, _ in common.COLUMNS[:14]])
        wb.save(self.wb)
        self.upsert([candidate(candidate_name="New Row", profile_url="https://github.com/new", evidence_sources=1)])
        ws, h = self.sheet()
        self.assertEqual([c.value for c in ws[1]], common.HEADERS)
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.cell(3, h["Evidence Sources"]).value, 1)

    def test_retention_lists_unreviewed_new_rows_after_90_days(self):
        old = (date.today() - timedelta(days=120)).isoformat()
        self.upsert([candidate(candidate_name="Old New", profile_url="https://github.com/old", date_found=old),
                     candidate(candidate_name="Fresh New", profile_url="https://github.com/fresh")])
        before = self.wb.read_bytes()
        buf = io.StringIO()
        with mock.patch.object(sys, "argv", ["retention.py", "--workbook", str(self.wb), "--db", str(self.db)]), redirect_stdout(buf):
            retention.main()
        text = buf.getvalue()
        self.assertIn("not reviewed: 1 rows still New", text)
        self.assertIn("Old New · New · found " + old, text)
        self.assertNotIn("Fresh New", text)
        self.assertEqual(before, self.wb.read_bytes())

    def test_art14_notice_names_source_and_date(self):
        t = common.art14_notice("GitHub", "2026-10-15", "Senior React Developer", "Humancloud")
        for part in ("public profile on GitHub on 2026-10-15", "Senior React Developer", "Humancloud", "delete it"):
            self.assertIn(part, t)
