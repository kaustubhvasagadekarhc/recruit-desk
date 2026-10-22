"""Structured fields, Evidence sheet and CRM export (2026-10-16, data-structure plan items 1, 4 and 8). Temp copies only."""
from __future__ import annotations

import csv
import sqlite3
import sys
from unittest import mock

from openpyxl import Workbook, load_workbook

from helpers import JD, TempTestCase, candidate

import backfill_structured
import common
import excel_upsert
import export_crm
import memory
import score

W, CRIT = JD["weights"], JD["criteria"]
PARTS = ("Core Skills pts", "Experience pts", "Domain pts", "Seniority pts", "Location pts", "Preferred pts", "OTW pts",
         "Cap Adjustment pts")


def scored(**over) -> dict:
    c = candidate(**over)
    for f in ("open_to_work", "open_to_work_quote", *common.FACTOR_FIELDS):
        c.pop(f, None)
    score.score(c, W, CRIT)
    return c


class Splits(TempTestCase):
    def test_names(self):
        self.assertEqual(common.split_name("Rahul Sethi"), ("Rahul", "Sethi"))
        self.assertEqual(common.split_name("Dr. A. B. Rao, PMP (He/Him)"), ("A.", "B. Rao"))
        self.assertEqual(common.split_name("Madonna"), ("Madonna", ""))
        self.assertEqual(common.split_name(""), ("", ""))

    def test_role_company(self):
        self.assertEqual(common.split_role_company("Senior Frontend Engineer @ Acme"), ("Senior Frontend Engineer", "Acme"))
        self.assertEqual(common.split_role_company("Lead UI Developer at Infosys Ltd"), ("Lead UI Developer", "Infosys Ltd"))
        self.assertEqual(common.split_role_company("Senior React.js Engineer (company not shown in snippet)"),
                         ("Senior React.js Engineer", ""))
        self.assertEqual(common.split_role_company("Frontend Dev @ (not stated)"), ("Frontend Dev", ""))

    def test_experience_and_work_mode(self):
        self.assertEqual(common.experience_basis("9.5 yrs (self-stated headline)"), "self-stated")
        self.assertEqual(common.experience_basis("~6 yrs, estimated from timeline"), "computed")
        self.assertEqual(common.experience_basis("~4 yrs (estimated)"), "estimate")
        self.assertEqual(common.experience_basis("senior"), "")
        self.assertEqual(common.work_mode_of("Pune (Remote)"), "Remote")
        self.assertEqual(common.work_mode_of("Pune, Maharashtra, India"), "Unknown")   # silence is never Onsite
        self.assertEqual(common.work_mode_of("Remote or hybrid"), "Unknown")

    def test_rahul_example_from_the_report(self):
        s = common.structured_fields({"candidate_name": "Rahul Sethi", "experience": "9.5 yrs (self-stated headline)",
                                      "current_role_company": "Senior React.js Engineer (company not shown in snippet)"})
        self.assertEqual((s["first_name"], s["last_name"], s["current_title"], s["current_company"]),
                         ("Rahul", "Sethi", "Senior React.js Engineer", ""))
        self.assertEqual((s["experience_years"], s["experience_basis"]), (9.5, "self-stated"))

    def test_batch_values_win(self):
        s = common.structured_fields({"candidate_name": "A B", "first_name": "Anil", "work_mode": "Hybrid"})
        self.assertEqual((s["first_name"], s["last_name"], s["work_mode"]), ("Anil", "B", "Hybrid"))


class ScoreParts(TempTestCase):
    def test_parts_add_up_to_match_score(self):
        for over in ({}, {"evidence_notes": "React dev #OpenToWork"}, {"mandatory_missing": ["TypeScript"]},
                     {"experience": "2 yrs"}, {"evidence_confidence": "Low"}):
            c = scored(**over)
            s = common.structured_fields(c)
            total = sum(s[f] for f in common.SCORE_PART_FIELDS) + s["pts_cap"]
            self.assertEqual(total, c["match_score"], over)
            self.assertEqual(s["pts_core"], c["scores"]["core_skills"])

    def test_gate_shows_as_negative_cap(self):
        c = scored(mandatory_missing=["TypeScript"])
        self.assertLess(common.structured_fields(c)["pts_cap"], 0)

    def test_older_llm_breakdown_and_partial(self):
        p = common.score_parts("Core skills 32/35 · Experience 15/20 · Domain 12/15 · Seniority 8/10 · Location 10/10 · Preferred 6/10")
        self.assertEqual(p, {"pts_core": 32, "pts_experience": 15, "pts_domain": 12, "pts_seniority": 8,
                             "pts_location": 10, "pts_preferred": 6, "pts_otw": 0})   # pre-bonus breakdown: OTW 0
        s = common.structured_fields({"match_score": 80, "score_breakdown": "Core skills 30/35"})
        self.assertEqual((s["pts_core"], s["pts_cap"], s["pts_otw"]), (30, "", ""))   # incomplete: no cap, no guess


class Workbook_(TempTestCase):
    def setUp(self):
        super().setUp()
        self.db, self.wb = self.tmp / "memory.db", self.tmp / "wb.xlsx"
        self._p = [mock.patch.object(excel_upsert, "DB_PATH", self.db), mock.patch.object(excel_upsert, "WORKBOOK_PATH", self.wb),
                   mock.patch.object(memory, "DB_PATH", self.db), mock.patch.object(memory, "WORKBOOK_PATH", self.wb),
                   mock.patch.object(common, "DB_PATH", self.db)]
        for p in self._p:
            p.start()
        self.con = memory.connect()
        memory.add_jd(self.con, JD)
        self.con.commit()

    def tearDown(self):
        self.con.close()
        for p in self._p:
            p.stop()
        super().tearDown()

    def upsert(self, cands, *extra, jd=JD["id"]):
        path = self.write_json("b.json", {"jd_id": jd, "candidates": cands})
        with mock.patch.object(sys, "argv", ["excel_upsert.py", "--input", str(path), *extra]):
            excel_upsert.main()

    def rows(self, sheet):
        ws = load_workbook(self.wb)[sheet]
        vals = list(ws.iter_rows(values_only=True))
        return vals[0], [dict(zip(vals[0], v)) for v in vals[1:]]

    def rich(self, **over):
        base = dict(scores={"core_skills": 30, "domain_projects": 0, "seniority": 0, "preferred_skills": 0},
                    quotes={"core_skills": ["React | TypeScript"]}, facts=[{"fact": "7+ YOE", "source": "LinkedIn headline"}],
                    evidence_notes="Headline: 'Senior Frontend Engineer | React | TypeScript | 7+ YOE'. #OpenToWork",
                    corroboration={"url": "https://github.com/test-person", "source_type": "GitHub", "fetched_at": "2026-10-16T05:00:00Z",
                                   "snippet": "Test Person · Pune · builds React apps in TypeScript",
                                   "confirms": [{"fact": "Works in TypeScript", "quote": "React apps in TypeScript"}]})
        base.update(over)
        if base["corroboration"] is None:
            del base["corroboration"]
        return scored(**base)

    def test_new_columns_filled_and_old_ones_unchanged(self):
        self.upsert([self.rich()])
        hdr, rows = self.rows(JD["title"])
        self.assertEqual(list(hdr[:19]), common.HEADERS[:19])
        self.assertEqual(list(hdr[19:35]), ["First Name", "Last Name", "Title", "Company", "Experience Years", "Experience Basis",
                                          "Work Mode", *PARTS, "Evidence Count"])
        r = rows[0]
        self.assertEqual((r["First Name"], r["Last Name"], r["Title"], r["Company"]), ("Test", "Person", "Senior Frontend Engineer", "Acme"))
        self.assertEqual((r["Experience Years"], r["Experience Basis"], r["Work Mode"]), (7, "self-stated", "Unknown"))
        self.assertEqual(sum(r[p] for p in PARTS), r["Match Score (0-100)"])

    def test_evidence_sheet_and_count_idempotent(self):
        self.upsert([self.rich()])
        self.upsert([self.rich()])   # rerun: no duplicate evidence rows
        hdr, ev = self.rows(common.EVIDENCE_SHEET)
        self.assertEqual(list(hdr[:6]), ["Candidate Key", "Fact", "Quote", "Source URL", "Source Type", "Fetched On"])
        facts = {(e["Fact"], e["Quote"], e["Source URL"]) for e in ev}
        url = "https://www.linkedin.com/in/test-person-1"
        self.assertIn(("7+ YOE", None, url), facts)
        self.assertIn(("Core skills", "React | TypeScript", url), facts)
        self.assertIn(("Open to work: Yes", "#OpenToWork", url), facts)
        self.assertIn(("Works in TypeScript", "React apps in TypeScript", "https://github.com/test-person"), facts)
        self.assertEqual(len(ev), 4)
        self.assertTrue(all(e["Candidate Key"] == "url:https://www.linkedin.com/in/test-person-1" for e in ev))
        self.assertEqual({e["Fetched On"] for e in ev if "github" in e["Source URL"]}, {"2026-10-16"})
        _, rows = self.rows(JD["title"])
        self.assertEqual(rows[0]["Evidence Count"], 4)

    def test_source_without_facts_gets_a_blank_row(self):
        self.upsert([scored(profile_url="https://github.com/plain")])
        _, ev = self.rows(common.EVIDENCE_SHEET)
        self.assertEqual([(e["Fact"], e["Quote"], e["Source URL"]) for e in ev], [(None, None, "https://github.com/plain")])

    def test_19_column_sheet_migrates_in_place(self):
        wb = Workbook()
        ws = wb.active
        ws.title = JD["title"]
        ws.append(common.HEADERS[:19])
        ws.append([candidate(candidate_name="Old Row").get(f) or None for f, _, _ in common.COLUMNS[:19]])
        wb.save(self.wb)
        self.upsert([self.rich(profile_url="https://github.com/new", corroboration=None)])
        hdr, rows = self.rows(JD["title"])
        self.assertEqual(list(hdr), common.HEADERS)
        self.assertEqual(rows[0]["Candidate Name"], "Old Row")
        self.assertIsNone(rows[0]["First Name"])   # untouched until backfilled or upserted
        self.assertEqual(rows[0]["Evidence Count"], 0)

    def test_backfill_fills_existing_rows(self):
        wb = Workbook()
        ws = wb.active
        ws.title = JD["title"]
        ws.append(common.HEADERS[:15])
        old = candidate(candidate_name="Rahul Sethi", current_role_company="Senior React.js Engineer (company not shown in snippet)",
                        experience="9.5 yrs (self-stated headline)", match_score=70,
                        score_breakdown="Core skills 24/35 · Experience 19/20 · Domain/projects 4/15 · Seniority 6/10 · "
                                        "Location 10/10 · Preferred skills 2/10 · Open to work +5 (Yes: \"x\")")
        ws.append([old.get(f) or None for f, _, _ in common.COLUMNS[:15]])
        wb.save(self.wb)
        before = list(load_workbook(self.wb)[JD["title"]].iter_rows(values_only=True))
        r = backfill_structured.backfill(self.wb, self.db)
        self.assertEqual(r["rows_filled"], 1)
        hdr, rows = self.rows(JD["title"])
        self.assertEqual(list(hdr), common.HEADERS)
        self.assertEqual(tuple(list(self.rows(JD["title"])[0])[:15]), before[0])
        x = rows[0]
        self.assertEqual((x["First Name"], x["Last Name"], x["Company"], x["Experience Years"]), ("Rahul", "Sethi", None, 9.5))
        self.assertEqual((x["Core Skills pts"], x["OTW pts"], x["Cap Adjustment pts"]), (24, 5, 0))
        self.assertIsNone(x["Open to Work"])   # key-factor columns stay blank: not rescored
        self.assertIn(common.EVIDENCE_SHEET, load_workbook(self.wb).sheetnames)

    def test_forget_removes_evidence(self):
        self.upsert([self.rich()])
        memory.forget(self.con, "https://www.linkedin.com/in/test-person-1")
        _, ev = self.rows(common.EVIDENCE_SHEET)
        self.assertEqual(ev, [])
        self.assertEqual(sqlite3.connect(self.db).execute("SELECT count(*) FROM evidence").fetchone()[0], 0)

    def test_where_lists_evidence(self):
        self.upsert([self.rich()])
        out = memory.where(self.con, "https://www.linkedin.com/in/test-person-1")
        self.assertEqual(len(out["evidence"]), 4)

    def test_evidence_is_a_reserved_sheet_name(self):
        r = memory.add_jd(self.con, dict(JD, id="evidence-jd", title="Evidence"))
        self.assertNotEqual(r["sheet_name"], common.EVIDENCE_SHEET)

    # ---------- CRM export (item 8) ----------
    def export(self, fmt):
        out = self.tmp / f"{fmt}.csv"
        export_crm.export(fmt, out, self.wb)
        with out.open() as f:
            return list(csv.DictReader(f))

    def test_hubspot_keys_by_email_only_when_self_published(self):
        a = self.rich(candidate_name="Asha Rao", profile_url="https://github.com/asha", corroboration=None,
                      email_mobile="asha@example.dev", contact_source_url="https://asha.dev/contact")
        b = self.rich(candidate_name="Bo Li", profile_url="https://github.com/boli", corroboration=None,
                      evidence_notes="Headline: 'Senior Frontend Engineer | React | TypeScript | 7+ YOE'. mail bo@leak.dev #OpenToWork")
        self.upsert([a, b], "--contacts")
        rows = {r["Last Name"]: r for r in self.export("hubspot")}
        self.assertEqual(list(rows["Rao"].keys()), export_crm.FORMATS["hubspot"])
        self.assertEqual((rows["Rao"]["Record Key"], rows["Rao"]["Record Key Type"], rows["Rao"]["Email"]),
                         ("asha@example.dev", "email", "asha@example.dev"))
        self.assertEqual((rows["Li"]["Record Key"], rows["Li"]["Record Key Type"], rows["Li"]["Email"]),
                         ("https://github.com/boli", "profile_url", ""))   # never harvested from notes
        self.assertEqual((rows["Rao"]["Job Title"], rows["Rao"]["Company Name"], rows["Rao"]["City"], rows["Rao"]["Country/Region"]),
                         ("Senior Frontend Engineer", "Acme", "Pune", "India"))

    def test_without_contacts_flag_no_email_reaches_the_crm(self):
        a = self.rich(profile_url="https://github.com/asha", corroboration=None,
                      email_mobile="asha@example.dev", contact_source_url="https://asha.dev/contact")
        self.upsert([a])
        (row,) = self.export("hubspot")
        self.assertEqual((row["Email"], row["Record Key Type"]), ("", "profile_url"))

    def test_salesforce_merges_jds_and_one_word_name_goes_to_last_name(self):
        jd2 = dict(JD, id="test-react-mumbai", title="React Developer Mumbai")
        memory.add_jd(self.con, jd2)
        self.con.commit()
        p = self.rich(candidate_name="Madonna", corroboration=None)
        self.upsert([p])
        self.upsert([dict(p, match_score=50)], jd=jd2["id"])
        (row,) = self.export("salesforce")
        self.assertEqual(list(row.keys()), export_crm.FORMATS["salesforce"])
        self.assertEqual((row["First Name"], row["Last Name"]), ("", "Madonna"))
        self.assertEqual(row["Matched JDs"], f"{JD['title']}; {jd2['title']}")
        self.assertEqual(row["Account Name"], "Acme")
        self.assertEqual(row["Record Key"], "https://www.linkedin.com/in/test-person-1")

    def test_status_filter(self):
        self.upsert([self.rich(corroboration=None)])
        out = self.tmp / "x.csv"
        r = export_crm.export("hubspot", out, self.wb, statuses={"Shortlisted"})
        self.assertEqual(r["contacts"], 0)
