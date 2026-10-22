"""GitHub and Stack Exchange skill evidence (2026-10-16, data-structure plan item 9). Fixtures only, no live calls."""
from __future__ import annotations

import json
import sys
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from helpers import JD, TempTestCase, candidate

import common
import dev_evidence
import excel_upsert
import memory
import score

FIXTURES = json.loads((Path(__file__).parent / "fixtures" / "dev_api.json").read_text())
W, CRIT = JD["weights"], JD["criteria"]
GH = "https://github.com/test-person"
SO = "https://stackoverflow.com/users/4242/test-person"


class FakeAPI:
    """Answers from the fixture file; scripted failures first. Records every URL asked for."""

    def __init__(self, failures=None, clock=1_000_000.0):
        self.urls, self.failures, self.now, self.slept = [], list(failures or []), clock, []

    def transport(self, url, headers):
        assert "Authorization" not in headers or headers["Authorization"].startswith("Bearer ")
        self.urls.append(url)
        if self.failures:
            return self.failures.pop(0)
        if url in FIXTURES:
            return 200, {}, json.loads(json.dumps(FIXTURES[url]))
        return 404, {}, {"message": "Not Found"}

    def sleep(self, s):
        self.slept.append(s)
        self.now += s

    def clock(self):
        return self.now

    def client(self, **kw):
        return dev_evidence.Client(transport=self.transport, sleep=self.sleep, clock=self.clock, **kw)


def linked(**over):
    return candidate(**{"match_score": 80, "public_urls": [GH, SO], **over})


class Links(TempTestCase):
    def test_only_bare_profile_links(self):
        self.assertEqual(dev_evidence.github_login("https://www.github.com/Test-Person/"), "Test-Person")
        self.assertIsNone(dev_evidence.github_login("https://github.com/someone/their-repo"))   # may be someone else's
        self.assertIsNone(dev_evidence.github_login("https://github.com/orgs"))
        self.assertIsNone(dev_evidence.github_login("https://gist.github.com/test-person"))
        self.assertEqual(dev_evidence.se_profile(SO), ("stackoverflow", 4242))
        self.assertEqual(dev_evidence.se_profile("https://unix.stackexchange.com/users/77/x"), ("unix", 77))
        self.assertIsNone(dev_evidence.se_profile("https://meta.stackexchange.com/users/77/x"))
        self.assertIsNone(dev_evidence.se_profile("https://stackoverflow.com/questions/1/x"))

    def test_links_come_from_the_candidates_own_entry(self):
        c = candidate(evidence_notes=f"Bio links {GH} and {SO}.")
        self.assertEqual(set(dev_evidence.linked_profiles(c)), {"github", "stackexchange"})
        self.assertEqual(dev_evidence.linked_profiles(candidate()), {})   # LinkedIn only: nothing to call


class Fetch(TempTestCase):
    def test_signals_and_facts(self):
        api = FakeAPI()
        cands = [linked(), candidate(candidate_name="No Links", match_score=90)]
        stats = dev_evidence.fetch(cands, api.client(), now="2026-10-16T08:00:00Z")
        self.assertEqual((stats["github"], stats["stackexchange"], stats["no_link"]), (1, 1, 1))
        gh = cands[0]["dev_signals"]["github"]
        self.assertEqual(gh["languages"], {"TypeScript": 307200, "Go": 102400, "CSS": 51200, "Shell": 2048})
        self.assertEqual(gh["last_active"], "2026-10-01")
        self.assertEqual([r["name"] for r in gh["repos"]], ["ui-kit", "go-tools", "dotfiles"])   # the fork is left out
        self.assertIn({"fact": "GitHub repo ui-kit (TypeScript), 120 stars",
                       "quote": "ui-kit: stargazers_count 120, language TypeScript"}, gh["facts"])
        se = cands[0]["dev_signals"]["stackexchange"]
        self.assertEqual(se["top_tags"][0], {"tag": "reactjs", "answer_score": 120, "answer_count": 45})
        self.assertEqual(se["reputation"], 5310)
        self.assertNotIn("dev_signals", cands[1])
        # never a name search: only the linked profiles' own endpoints were called
        self.assertFalse(any("search" in u or "No%20Links" in u or "Test%20Person" in u for u in api.urls))
        self.assertEqual(len(api.urls), 7)

    def test_never_stores_an_email(self):
        cands = [linked()]
        dev_evidence.fetch(cands, FakeAPI().client())
        blob = json.dumps(cands[0]["dev_signals"])
        self.assertNotIn("@", blob)
        self.assertNotIn("email", blob.lower())
        self.assertNotIn("Acme", blob)   # profile name / company are not copied either
        self.assertEqual(common.validate_dev_signals(cands[0]), [])

    def test_validation_rejects_an_email(self):
        c = candidate(dev_signals={"github": {"url": GH, "facts": [{"fact": "x", "quote": "mail test@mail.invalid"}]}})
        self.assertTrue(common.validate_dev_signals(c))
        self.assertTrue(common.validate_dev_signals(candidate(dev_signals={"github": {"email": ""}})))

    def test_github_rate_limit_waits_for_the_reset(self):
        api = FakeAPI()
        api.failures = [(403, {"x-ratelimit-remaining": "0", "x-ratelimit-reset": str(int(api.now) + 30)}, {"message": "rate limit"})]
        cands = [candidate(match_score=80, public_urls=[GH])]
        stats = dev_evidence.fetch(cands, api.client())
        self.assertEqual(api.slept, [31.0])
        self.assertEqual(stats["github"], 1)

    def test_far_reset_stops_cleanly_and_keeps_earlier_candidates(self):
        api = FakeAPI()
        cands = [linked(match_score=90), candidate(candidate_name="Second", match_score=80, public_urls=[GH])]
        client = api.client(max_wait=60)
        real = api.transport
        def transport(url, headers):
            if len(api.urls) >= 7:   # the second person's first call hits an exhausted hourly window
                api.urls.append(url)
                return 403, {"x-ratelimit-remaining": "0", "x-ratelimit-reset": str(int(api.now) + 3000)}, {}
            return real(url, headers)
        client.transport = transport
        stats = dev_evidence.fetch(cands, client)
        self.assertIn("more than --max-wait", stats["stopped"])
        self.assertIn("github", cands[0]["dev_signals"])
        self.assertNotIn("dev_signals", cands[1])
        self.assertEqual(api.slept, [])

    def test_server_errors_back_off_exponentially(self):
        api = FakeAPI(failures=[(502, {}, None), (503, {}, None)])
        dev_evidence.fetch([candidate(match_score=80, public_urls=[GH])], api.client())
        self.assertEqual(api.slept, [2, 4])

    def test_stack_exchange_backoff_is_honoured(self):
        api = FakeAPI()
        url = "https://api.stackexchange.com/2.3/users/4242?site=stackoverflow"
        FIX = dict(FIXTURES[url], backoff=5)
        real = api.transport
        api_t = lambda u, h: (api.urls.append(u), (200, {}, json.loads(json.dumps(FIX))))[1] if u == url else real(u, h)
        client = api.client()
        client.transport = api_t
        dev_evidence.fetch([candidate(match_score=80, public_urls=[SO])], client)
        self.assertEqual(api.slept, [5.0])   # waited before the next SE call, as the API asked

    def test_missing_profile_is_not_found(self):
        stats = dev_evidence.fetch([candidate(match_score=80, public_urls=["https://github.com/nobody-here"])], FakeAPI().client())
        self.assertEqual((stats["github"], stats["not_found"]), (0, 1))


def with_signals(**over) -> dict:
    c = linked(**over)
    dev_evidence.fetch([c], FakeAPI().client(), now="2026-10-16T08:00:00Z")
    for f in ("open_to_work", "open_to_work_quote", *common.FACTOR_FIELDS):
        c.pop(f, None)
    return c


class Scoring(TempTestCase):
    def test_api_sources_lift_medium_to_high(self):
        c = with_signals(scores={"core_skills": 30, "domain_projects": 0, "seniority": 0, "preferred_skills": 0},
                         quotes={"core_skills": ["React | TypeScript"]})
        score.score(c, W, CRIT)
        self.assertEqual(c["evidence_sources"], 3)
        self.assertEqual(c["evidence_confidence"], "High")
        self.assertIn("GitHub API", c["score_breakdown"])
        n = len(c["facts"])
        score.score(c, W, CRIT)   # rerun: same facts, no duplicates
        self.assertEqual(len(c["facts"]), n)
        self.assertEqual(c["evidence_sources"], 3)

    def test_same_site_never_counts_twice(self):
        c = with_signals(profile_url=GH, public_urls=[GH], scores={"core_skills": 30, "domain_projects": 0, "seniority": 0,
                                                                   "preferred_skills": 0}, quotes={"core_skills": ["React | TypeScript"]})
        score.score(c, W, CRIT)
        self.assertEqual(c["evidence_sources"], 1)
        self.assertEqual(c["evidence_confidence"], "Medium")

    def test_columns(self):
        cols = common.dev_columns(with_signals())
        self.assertEqual(cols["github_languages"], "TypeScript 300 KB (66%), Go 100 KB (22%), CSS 50 KB (11%), Shell 2 KB (0%)")
        self.assertEqual(cols["github_last_active"], "2026-10-01")
        self.assertEqual(cols["so_top_tags"], "reactjs (score 120, 45 answers), typescript (score 61, 20 answers), "
                                              "javascript (score 55, 30 answers)")
        self.assertEqual(common.dev_columns(candidate()), {"github_languages": "", "github_last_active": "", "so_top_tags": ""})
        self.assertEqual(common.HEADERS[-3:], ["GitHub Languages", "GitHub Last Active", "SO Top Tags"])


class Upsert(TempTestCase):
    def setUp(self):
        super().setUp()
        self.db, self.wb = self.tmp / "memory.db", self.tmp / "wb.xlsx"
        self._p = [mock.patch.object(excel_upsert, "DB_PATH", self.db), mock.patch.object(excel_upsert, "WORKBOOK_PATH", self.wb),
                   mock.patch.object(memory, "DB_PATH", self.db), mock.patch.object(memory, "WORKBOOK_PATH", self.wb),
                   mock.patch.object(common, "DB_PATH", self.db)]
        for p in self._p:
            p.start()
        self.con = memory.connect()
        memory.add_jd(self.con, JD)
        self.con.commit()

    def tearDown(self):
        self.con.close()
        for p in self._p:
            p.stop()
        super().tearDown()

    def upsert(self, cands):
        path = self.write_json("b.json", {"jd_id": JD["id"], "candidates": cands})
        with mock.patch.object(sys, "argv", ["excel_upsert.py", "--input", str(path)]):
            excel_upsert.main()

    def rows(self, sheet):
        vals = list(load_workbook(self.wb)[sheet].iter_rows(values_only=True))
        return vals[0], [dict(zip(vals[0], v)) for v in vals[1:]]

    def scored(self, **over):
        c = with_signals(scores={"core_skills": 30, "domain_projects": 0, "seniority": 0, "preferred_skills": 0},
                         quotes={"core_skills": ["React | TypeScript"]}, **over)
        score.score(c, W, CRIT)
        return c

    def test_columns_and_evidence_rows(self):
        self.upsert([self.scored()])
        hdr, rows = self.rows(JD["title"])
        self.assertEqual(list(hdr[:35]), common.HEADERS[:35])   # existing headers unchanged
        r = rows[0]
        self.assertTrue(r["GitHub Languages"].startswith("TypeScript 300 KB (66%)"))
        self.assertEqual(r["GitHub Last Active"], "2026-10-01")
        self.assertTrue(r["SO Top Tags"].startswith("reactjs (score 120"))
        self.assertEqual(r["Evidence Confidence"], "High")
        _, ev = self.rows(common.EVIDENCE_SHEET)
        kinds = {e["Source Type"] for e in ev}
        self.assertTrue({"GitHub API", "Stack Exchange API"} <= kinds)
        self.assertIn(("GitHub language: TypeScript (66% of code in 3 recent repos)", "TypeScript: 307200 bytes", GH),
                      {(e["Fact"], e["Quote"], e["Source URL"]) for e in ev})
        self.assertFalse(any("@" in str(v) for e in ev for v in e.values()))
        srcs = {u for (u,) in self.con.execute("SELECT url FROM sources")}
        self.assertTrue({GH, SO} <= srcs)

    def test_later_batch_without_signals_keeps_the_cells(self):
        self.upsert([self.scored()])
        c = candidate(match_score=60, score_breakdown="Core skills 30/35")
        self.upsert([c])
        _, rows = self.rows(JD["title"])
        self.assertEqual(rows[0]["GitHub Last Active"], "2026-10-01")

    def test_35_column_sheet_migrates_in_place(self):
        wb = Workbook()
        ws = wb.active
        ws.title = JD["title"]
        ws.append(common.HEADERS[:35])
        ws.append([candidate(candidate_name="Old Row").get(f) or None for f, _, _ in common.COLUMNS[:35]])
        wb.save(self.wb)
        self.upsert([self.scored(profile_url="https://www.linkedin.com/in/new-one")])
        hdr, rows = self.rows(JD["title"])
        self.assertEqual(list(hdr), common.HEADERS)
        self.assertEqual(rows[0]["Candidate Name"], "Old Row")
        self.assertIsNone(rows[0]["GitHub Languages"])
        self.assertEqual(rows[1]["GitHub Last Active"], "2026-10-01")
