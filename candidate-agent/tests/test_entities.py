"""Candidates and Matches entity sheets (2026-10-18, data-structure plan item 7). Temp copies only."""
from __future__ import annotations

import sqlite3
import sys
from unittest import mock

from openpyxl import load_workbook

from helpers import JD, TempTestCase

import common
import entities
import excel_upsert
import memory
import progress
from test_structured_crm import scored

JD2 = dict(JD, id="test-vue-remote", title="Vue Frontend Developer")
URL = "https://www.linkedin.com/in/test-person-1"


class Entities(TempTestCase):
    def setUp(self):
        super().setUp()
        self.db, self.wb = self.tmp / "memory.db", self.tmp / "wb.xlsx"
        self._p = [mock.patch.object(m, n, v) for m, n, v in (
            (excel_upsert, "DB_PATH", self.db), (excel_upsert, "WORKBOOK_PATH", self.wb), (memory, "DB_PATH", self.db),
            (memory, "WORKBOOK_PATH", self.wb), (common, "DB_PATH", self.db), (entities, "DB_PATH", self.db),
            (entities, "WORKBOOK_PATH", self.wb))]
        for p in self._p:
            p.start()
        self.con = memory.connect()
        memory.add_jd(self.con, JD)
        memory.add_jd(self.con, JD2)
        self.con.commit()

    def tearDown(self):
        self.con.close()
        for p in self._p:
            p.stop()
        super().tearDown()

    def upsert(self, cands, jd):
        path = self.write_json("b.json", {"jd_id": jd["id"], "candidates": cands})
        with mock.patch.object(sys, "argv", ["excel_upsert.py", "--input", str(path)]):
            excel_upsert.main()

    def rows(self, sheet):
        vals = list(load_workbook(self.wb)[sheet].iter_rows(values_only=True))
        return vals[0], [dict(zip(vals[0], v)) for v in vals[1:]]

    def two_jds(self):
        self.upsert([scored(), scored(candidate_name="Other One", profile_url="https://www.linkedin.com/in/other-one")], JD)
        self.upsert([scored(experience="2 yrs")], JD2)

    def test_one_person_two_jds_is_one_candidate_and_two_matches(self):
        self.two_jds()
        hdr, cands = self.rows("Candidates")
        self.assertEqual(list(hdr), [h for h, _ in entities.CANDIDATE_COLUMNS])
        self.assertEqual(len(cands), 2)
        me = next(c for c in cands if c["Candidate Key"] == f"url:{URL}")
        self.assertEqual((me["JDs Matched"], me["First Name"], me["Last Name"], me["Company"]), (2, "Test", "Person", "Acme"))
        _, matches = self.rows("Matches")
        mine = [m for m in matches if m["Candidate Key"] == f"url:{URL}"]
        self.assertEqual({m["JD Id"] for m in mine}, {JD["id"], JD2["id"]})
        self.assertEqual(len(matches), 3)
        self.assertEqual(me["Best Match"], max(m["Match Score (0-100)"] for m in mine))

    def test_match_score_parts_add_up(self):
        self.two_jds()
        _, matches = self.rows("Matches")
        for m in matches:
            self.assertEqual(sum(m[h] for h in entities.PTS), m["Match Score (0-100)"])

    def test_recruiter_status_comes_from_the_jd_sheet(self):
        self.two_jds()
        wb = load_workbook(self.wb)
        ws = wb[JD["title"]]
        col = common.HEADERS.index("Recruiter Status") + 1
        ws.cell(2, col, "Shortlisted")
        wb.save(self.wb)
        self.upsert([scored()], JD)   # any upsert rebuilds the entity sheets
        _, matches = self.rows("Matches")
        self.assertIn("Shortlisted", {m["Recruiter Status"] for m in matches if m["JD Id"] == JD["id"]})

    def test_idempotent_and_jd_sheets_untouched(self):
        self.two_jds()
        before = self.rows(JD["title"])
        self.upsert([scored()], JD)
        self.upsert([scored()], JD)
        self.assertEqual(self.rows(JD["title"])[1], before[1])
        self.assertEqual(len(self.rows("Candidates")[1]), 2)
        self.assertEqual(len(self.rows("Matches")[1]), 3)

    def test_no_email_copied(self):
        self.upsert([scored(email_mobile="a@b.com", contact_source_url="https://x.test")], JD)
        for sheet in ("Candidates", "Matches"):
            hdr, rows = self.rows(sheet)
            self.assertNotIn("Email / Mobile", hdr)
            self.assertFalse(any("a@b.com" in str(v) for r in rows for v in r.values()))

    def test_entity_sheets_are_reserved_names(self):
        for t in ("Candidates", "Matches"):
            self.assertNotEqual(memory.add_jd(self.con, dict(JD, id=f"{t}-jd", title=t))["sheet_name"], t)

    def test_forget_removes_the_person_from_both_sheets(self):
        self.two_jds()
        memory.forget(self.con, URL)
        self.assertEqual([c["Candidate Name"] for c in self.rows("Candidates")[1]], ["Other One"])
        self.assertEqual({m["Candidate Name"] for m in self.rows("Matches")[1]}, {"Other One"})

    def test_migrate_backs_up_and_dry_run_changes_nothing(self):
        self.two_jds()
        wb = load_workbook(self.wb)
        del wb["Candidates"], wb["Matches"]   # a workbook written before this change
        wb.save(self.wb)
        r = entities.migrate(self.wb, self.db, dry_run=True)
        self.assertEqual((r["candidates"], r["matches"], r["backups"]), (2, 3, []))
        self.assertNotIn("Candidates", load_workbook(self.wb).sheetnames)
        r = entities.migrate(self.wb, self.db, stamp="2026-10-18")
        self.assertEqual(len(r["backups"]), 2)
        self.assertTrue((self.tmp / "wb.backup-2026-10-18-before-entities.xlsx").exists())
        self.assertEqual(len(self.rows("Matches")[1]), 3)
        with self.assertRaises(SystemExit):
            entities.migrate(self.wb, self.db, stamp="2026-10-18")   # never overwrite a backup

    def test_sheet_order_and_old_sheets_still_read(self):
        self.two_jds()
        names = load_workbook(self.wb).sheetnames
        self.assertEqual(names[-3:], ["Candidates", "Matches", "Evidence"])
        con = sqlite3.connect(self.db)
        self.assertEqual(con.execute("SELECT count(*) FROM candidates").fetchone()[0], 2)
