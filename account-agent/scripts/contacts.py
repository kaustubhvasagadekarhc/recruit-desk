"""Leaders and contacts: the rules for the draft's "_contacts" and "_channels", and their two workbook sheets.

A draft (runs/<slug>.draft.json) row may carry two working keys next to its columns:

  "_contacts": up to 6 senior people (CEO, founders, CTO, VP / Head of Engineering, COO, CHRO or talent head,
    any US-based leader), each
      {"name", "title", "location", "source": <https URL naming the person and title>, "verified": "YYYY-MM-DD",
       "email":  {"value": "a@b.com", "quote": "<words on the page that contain a@b.com>", "source": <URL>}
                 or "Not published (checked: ...)",
       "phone":  {"value": "+1 ...", "kind": "personal" | "company line", "quote": "<words containing it>",
                  "source": <URL>} or "Not published (checked: ...)",
       "linkedin", "x", "github", "site", "podcast", "talks", "booking": a URL (talks may add a short label
                 before it) or "Not published (checked: ...)"; a channel left out is written as
                 "Not published (checked: <channels_checked>)",
       "channels_checked": "<what was searched for this person's channels>"}
  "_channels": the company's own published channels,
      {"press_email", "sales", "partnerships", "investor_relations", "contact_form": {"value", "source", "quote"?}
       or "Not published (checked: ...)",
       "offices": [{"label", "address", "phone"?, "source", "quote"?}],
       "checked": "<what was searched>"}

The hard rules (validate_contacts, also run by the hook on every draft or rows Write):
  - An email is accepted only with the source URL and a quote from that page that contains the exact
    address. A phone likewise (digits compared). Nothing else may hold an email or a phone number.
  - Wording that admits a guess ("likely", "pattern", "firstname@", "inferred", ...) is refused anywhere.
  - LinkedIn is a linkedin.com/in/ profile URL taken from search results; the page itself is never fetched.
  - Sources are http(s) URLs; paid enrichment and data-broker hosts (Apollo pages, Hunter, RocketReach, ZoomInfo,
    Lusha, ContactOut, SignalHire, ...) are refused as sources.
  - The one exception (approved by the human, 2026-10-16): the Apollo.io API through scripts/apollo.py. A person,
    email or phone from it carries "source": "apollo" and the Apollo person id ("apollo_id"); an Apollo email also
    needs "email_status": "verified" (guessed, likely or unavailable emails are refused). No quote is needed, and
    the sheet and the app show "Source: Apollo (verified)" apart from self-published facts.

Storage (data minimisation, like candidate-agent): a person's own email and phone are written to the
Contacts sheet only when excel_upsert.py runs with --contacts; otherwise those cells say so. A company line
and the company channels are business details and are always written.

CLI:
  uv run python scripts/contacts.py check --draft runs/<slug>.draft.json   # rules only, offline
  uv run python scripts/contacts.py verify --draft runs/<slug>.draft.json  # re-reads each email / phone source page
                                                                           # and confirms the value is on it
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date
from pathlib import Path
from urllib.parse import urlsplit

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import is_not_found, today  # noqa: E402

MAX_PEOPLE = 6
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
EMAIL_FULL = re.compile(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$")
PHONE_RE = re.compile(r"(?<![\w/])\+?\d[\d\s().-]{6,}\d(?![\w/])")
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
URL_FULL = re.compile(r"^https?://[^\s<>\"']+$")
URL_IN = re.compile(r"https?://[^\s<>\"']+")
GUESS = re.compile(
    r"(?i)\b(likely|probabl[ey]|presumabl[ey]|guess\w*|inferr?\w*|construct\w*|pattern|typical(ly)? format|"
    r"format is|assumed?|predicted|derived|best guess|unverified)\b|"
    r"\b(first(name)?|last(name)?|f?name|first\.last|flast|firstinitial)\s*[.@]"
)
PAID = re.compile(r"(?i)(^|\.)(apollo\.io|hunter\.io|rocketreach\.co|zoominfo\.com|lusha\.com|contactout\.com|"
                  r"signalhire\.com|seamless\.ai|clearbit\.com|peopledatalabs\.com|snov\.io|uplead\.com|"
                  r"kaspr\.io|leadiq\.com|adapt\.io|getprospect\.com|voilanorbert\.com|anymailfinder\.com|"
                  r"emailhippo\.com|spokeo\.com|whitepages\.com|truepeoplesearch\.com|beenverified\.com|"
                  r"fastpeoplesearch\.com|radaris\.com|zabasearch\.com|intelius\.com|signal\.nfx\.com|"
                  r"crunchbase\.com/person|rocketreach|datanyze\.com|lead411\.com)$")
CHANNEL_KEYS = ("linkedin", "x", "github", "site", "podcast", "talks", "booking")
COMPANY_KEYS = ("press_email", "sales", "partnerships", "investor_relations", "contact_form")
CHANNEL_HOSTS = {
    "linkedin": re.compile(r"(?i)^https://([a-z]{2,3}\.)?linkedin\.com/in/[^/?#\s]+/?$"),
    "x": re.compile(r"(?i)^https://(www\.)?(x|twitter)\.com/[A-Za-z0-9_]{1,15}/?$"),
    "github": re.compile(r"(?i)^https://(www\.)?github\.com/[A-Za-z0-9-]+/?$"),
}
PERSON_KEYS = {"name", "title", "location", "source", "verified", "email", "phone", "channels_checked", "apollo_id",
               *CHANNEL_KEYS}
APOLLO = "apollo"
APOLLO_LABEL = "Source: Apollo (verified)"
APOLLO_ID = re.compile(r"^[A-Za-z0-9]{8,40}$")

CONTACT_COLUMNS = [
    "Domain", "Company", "Name", "Title", "Location", "Email", "Email Source", "Phone", "Phone Source",
    "LinkedIn (from search results)", "X / Twitter", "GitHub", "Website / Blog", "Podcast / Newsletter",
    "Upcoming Talks", "Booking Link", "Source", "Last Verified",
]
CHANNEL_COLUMNS = ["Domain", "Company", "Channel", "Value", "Source", "Quote", "Last Verified"]
CHANNEL_LABELS = {"press_email": "Press / media", "sales": "Sales", "partnerships": "Partnerships",
                  "investor_relations": "Investor relations", "contact_form": "Contact form"}
CONTACTS_SHEET, CHANNELS_SHEET = "Contacts", "Company channels"
WITHHELD = "Not stored (this run did not use --contacts)"


def digits(s: str) -> str:
    return re.sub(r"\D", "", s or "")


def has_phone(text: str) -> bool:
    return any(len(digits(m.group(0))) >= 10 for m in PHONE_RE.finditer(text or ""))


def _src_ok(u, where: str, e) -> bool:
    u = str(u or "").strip()
    if not URL_FULL.match(u):
        e(f"{where}: source must be the https URL the fact was read on (got {u[:60]!r})")
        return False
    host = (urlsplit(u).hostname or "").lower()
    if PAID.search(host) or PAID.search(host + urlsplit(u).path.lower()):
        e(f"{where}: {host} is a paid enrichment or data-broker source; free public sources only")
        return False
    return True


def _check_apollo(v: dict, kind: str, where: str, e, allow_kind: bool) -> None:
    """An email or phone from the Apollo API: person id required; an email must be email_status 'verified'."""
    val = str(v.get("value") or "").strip()
    if kind not in ("email", "phone"):
        e(f"{where}: source 'apollo' is accepted only for a person's email or phone")
        return
    if not APOLLO_ID.match(str(v.get("apollo_id") or "")):
        e(f"{where}: an Apollo value needs the Apollo person id (apollo_id) it came from")
    if kind == "email":
        if not EMAIL_FULL.match(val):
            e(f"{where}: {val!r} is not one email address")
        if str(v.get("email_status") or "").strip().lower() != "verified":
            e(f"{where}: only emails Apollo marks email_status 'verified' are kept; guessed, likely or unavailable ones are refused")
    else:
        if len(digits(val)) < 7:
            e(f"{where}: {val!r} is not a phone number")
        if allow_kind and v.get("kind") not in ("personal", "company line"):
            e(f"{where}: kind must be 'personal' or 'company line'")
    for part in (val, str(v.get("kind") or "")):
        if GUESS.search(part):
            e(f"{where}: refused wording {GUESS.search(part).group(0)!r}; write only what Apollo verified")


def _check_value(v, kind: str, where: str, e, allow_kind: bool = False) -> None:
    """kind: email | phone | text. v is a dict {value, source, quote} or a 'Not published (checked: ...)' line."""
    if isinstance(v, str):
        s = v.strip()
        if not s:
            e(f"{where} is empty; write 'Not published (checked: <what was searched>)'")
        elif not is_not_published(s):
            if EMAIL_RE.search(s) or has_phone(s):
                e(f"{where}: an email or phone needs its source URL and a quote containing it "
                  f"({{\"value\", \"quote\", \"source\"}}); unsourced or constructed values are refused")
            else:
                e(f"{where} must be {{\"value\", \"source\"}} or 'Not published (checked: ...)'")
        return
    if not isinstance(v, dict):
        e(f"{where} must be an object or 'Not published (checked: ...)'")
        return
    val = str(v.get("value") or "").strip()
    quote = str(v.get("quote") or "").strip()
    if not val:
        e(f"{where}: value is empty")
        return
    if str(v.get("source") or "").strip().lower() == APOLLO:
        _check_apollo(v, kind, where, e, allow_kind)
        return
    _src_ok(v.get("source"), where, e)
    if kind == "email":
        if not EMAIL_FULL.match(val):
            e(f"{where}: {val!r} is not one email address")
        elif val.lower() not in quote.lower():
            e(f"{where}: the quote must contain {val} exactly as published; a constructed or unsourced email is refused")
    elif kind == "phone":
        if len(digits(val)) < 7:
            e(f"{where}: {val!r} is not a phone number")
        elif digits(val) not in digits(quote):
            e(f"{where}: the quote must contain {val} as published; a guessed phone is refused")
        if allow_kind and v.get("kind") not in ("personal", "company line"):
            e(f"{where}: kind must be 'personal' (the person published it) or 'company line' (the company's main number)")
    else:
        for m in EMAIL_RE.finditer(val):
            if m.group(0).lower() not in quote.lower():
                e(f"{where}: the email {m.group(0)} needs a quote from the source page that contains it")
        if has_phone(val) and not any(len(digits(p.group(0))) >= 7 and digits(p.group(0)) in digits(quote)
                                      for p in PHONE_RE.finditer(val)):
            e(f"{where}: the phone number needs a quote from the source page that contains it")
    for part in (val, str(v.get("kind") or "")):
        if GUESS.search(part):
            e(f"{where}: refused wording {GUESS.search(part).group(0)!r}; write only what was published")


def is_not_published(s: str) -> bool:
    s = (s or "").strip()
    return bool(re.match(r"^Not published \(checked: [^)]{6,}\)\.?$", s)) or is_not_found(s)


def _check_channel(key: str, v, where: str, e) -> None:
    if isinstance(v, dict):
        v = str(v.get("value") or "")
    s = str(v or "").strip()
    if not s:
        return
    if is_not_published(s):
        return
    if EMAIL_RE.search(s) or has_phone(s):
        e(f"{where}: emails and phones go only in the email / phone fields, with their quote")
        return
    if GUESS.search(s):
        e(f"{where}: refused wording {GUESS.search(s).group(0)!r}")
    url = s if key != "talks" else (URL_IN.findall(s) or [""])[-1]
    if not _src_ok(url, where, e):
        return
    rx = CHANNEL_HOSTS.get(key)
    if rx and not rx.match(url):
        want = {"linkedin": "https://www.linkedin.com/in/<id> (a profile URL seen in search results)",
                "x": "https://x.com/<handle>", "github": "https://github.com/<login>"}[key]
        e(f"{where}: must be {want}")


def validate_contacts(row: dict, today_: date | None = None) -> list[str]:
    """Errors for a row's "_contacts" and "_channels" (empty list when absent or valid)."""
    errs: list[str] = []
    tag = f"contacts ({row.get('Company') or '?'})"
    people = row.get("_contacts")
    chans = row.get("_channels")
    if people is None and chans is None:
        return errs

    def e(msg: str) -> None:
        errs.append(f"{tag}: {msg}")

    t = (today_ or date.today()).isoformat()
    if people is not None:
        if not isinstance(people, list):
            e("_contacts must be a list of people")
            people = []
        if len(people) > MAX_PEOPLE:
            e(f"_contacts holds {len(people)} people; keep the {MAX_PEOPLE} most senior")
        seen = set()
        for i, p in enumerate(people, 1):
            if not isinstance(p, dict):
                e(f"person {i} must be an object")
                continue
            who = f"person {i} ({p.get('name') or '?'})"
            for k in p:
                if k not in PERSON_KEYS:
                    e(f"{who}: unknown key {k!r} (allowed: {', '.join(sorted(PERSON_KEYS))})")
            for k in ("name", "title"):
                if not str(p.get(k) or "").strip():
                    e(f"{who}: {k} is required")
            n = str(p.get("name") or "").strip().lower()
            if n and n in seen:
                e(f"{who} appears twice")
            seen.add(n)
            for k in ("name", "title", "location", "channels_checked"):
                s = str(p.get(k) or "")
                if EMAIL_RE.search(s) or has_phone(s):
                    e(f"{who}: {k} holds an email or phone; those go only in email / phone, with their quote")
                if k in ("name", "location") and GUESS.search(s):
                    e(f"{who}: {k}: refused wording {GUESS.search(s).group(0)!r}")
            if not str(p.get("location") or "").strip():
                e(f"{who}: location is required (a place, or 'Not published (checked: ...)')")
            if str(p.get("source") or "").strip().lower() == APOLLO:
                if not APOLLO_ID.match(str(p.get("apollo_id") or "")):
                    e(f"{who}: a person found through Apollo needs the Apollo person id (apollo_id)")
            else:
                _src_ok(p.get("source"), f"{who} source", e)
            for k in ("email", "phone"):
                pv = p.get(k)
                if isinstance(pv, dict) and str(pv.get("source") or "").lower() == APOLLO and p.get("apollo_id") \
                        and pv.get("apollo_id") and pv.get("apollo_id") != p.get("apollo_id"):
                    e(f"{who} {k}: apollo_id differs from the person's apollo_id")
            v = str(p.get("verified") or "")
            if not DATE_RE.match(v):
                e(f"{who}: verified must be the YYYY-MM-DD the source was read")
            elif v > t:
                e(f"{who}: verified date {v} is in the future")
            _check_value(p.get("email", ""), "email", f"{who} email", e)
            _check_value(p.get("phone", ""), "phone", f"{who} phone", e, allow_kind=True)
            for k in CHANNEL_KEYS:
                _check_channel(k, p.get(k), f"{who} {k}", e)
            if any(not p.get(k) for k in CHANNEL_KEYS) and not str(p.get("channels_checked") or "").strip():
                e(f"{who}: channels_checked is required when a channel is left out (what was searched)")
    if chans is not None:
        if not isinstance(chans, dict):
            e("_channels must be an object")
            return errs
        for k in chans:
            if k not in (*COMPANY_KEYS, "offices", "checked"):
                e(f"_channels: unknown key {k!r}")
        for k in COMPANY_KEYS:
            if k in chans:
                kind = "email" if isinstance(chans[k], dict) and EMAIL_FULL.match(str(chans[k].get("value") or "").strip()) else "text"
                _check_value(chans[k], kind, f"channel {k}", e)
        if any(k not in chans for k in COMPANY_KEYS) and not str(chans.get("checked") or "").strip():
            e("_channels: 'checked' is required when a channel is left out (what was searched)")
        offices = chans.get("offices") or []
        if not isinstance(offices, list):
            e("_channels.offices must be a list")
            offices = []
        for i, o in enumerate(offices, 1):
            if not isinstance(o, dict) or not str(o.get("address") or "").strip():
                e(f"office {i}: needs label, address and source")
                continue
            _check_value({"value": o.get("address"), "source": o.get("source"), "quote": o.get("quote") or o.get("address")},
                         "text", f"office {i} address", e)
            if o.get("phone"):
                _check_value({"value": o.get("phone"), "source": o.get("source"), "quote": o.get("quote", "")},
                             "phone", f"office {i} phone", e)
    return errs


# ---- workbook rows --------------------------------------------------------------------------------
def _cell(v, missing: str) -> tuple[str, str]:
    """(value text, source text) for an email / phone / channel value."""
    if isinstance(v, dict):
        val = str(v.get("value") or "").strip()
        if v.get("kind") == "company line":
            val = f"{val} (company line)"
        src = str(v.get("source") or "")
        if src.lower() == APOLLO:
            label = APOLLO_LABEL if v.get("email_status") == "verified" else "Source: Apollo"
            return val, f"{label} · Apollo person id {v.get('apollo_id') or '?'}"
        q = str(v.get("quote") or "").strip()
        return val, (f'{src} — "{q[:160]}"' if q else src)
    s = str(v or "").strip()
    return (s or missing), ""


def contact_rows(row: dict, keep_personal: bool) -> list[list[str]]:
    """Contacts sheet rows for one company, in CONTACT_COLUMNS order."""
    out = []
    dom, comp = str(row.get("Domain") or ""), str(row.get("Company") or "")
    for p in row.get("_contacts") or []:
        missing = f"Not published (checked: {str(p.get('channels_checked') or 'not searched').strip()})"
        email, esrc = _cell(p.get("email"), missing)
        phone, psrc = _cell(p.get("phone"), missing)
        personal_phone = isinstance(p.get("phone"), dict) and p["phone"].get("kind") != "company line"
        if not keep_personal:
            if isinstance(p.get("email"), dict):
                email, esrc = WITHHELD, ""
            if personal_phone:
                phone, psrc = WITHHELD, ""
        chans = [_cell(p.get(k), missing)[0] for k in CHANNEL_KEYS]
        psource = str(p.get("source") or "")
        if psource.lower() == APOLLO:
            psource = f"Source: Apollo · Apollo person id {p.get('apollo_id') or '?'}"
        out.append([dom, comp, str(p.get("name") or ""), str(p.get("title") or ""), str(p.get("location") or ""),
                    email, esrc, phone, psrc, *chans, psource, str(p.get("verified") or today())])
    return out


def channel_rows(row: dict) -> list[list[str]]:
    ch = row.get("_channels")
    if not isinstance(ch, dict):
        return []
    dom, comp = str(row.get("Domain") or ""), str(row.get("Company") or "")
    missing = f"Not published (checked: {str(ch.get('checked') or 'not searched').strip()})"
    stamp = str(row.get("Last Checked") or today())
    out = []
    for k in COMPANY_KEYS:
        v = ch.get(k, missing)
        if isinstance(v, dict):
            out.append([dom, comp, CHANNEL_LABELS[k], str(v.get("value") or ""), str(v.get("source") or ""),
                        str(v.get("quote") or "")[:200], stamp])
        else:
            out.append([dom, comp, CHANNEL_LABELS[k], str(v or missing), "", "", stamp])
    for o in ch.get("offices") or []:
        val = str(o.get("address") or "") + (f" · {o['phone']}" if o.get("phone") else "")
        out.append([dom, comp, f"Office: {o.get('label') or 'office'}", val, str(o.get("source") or ""),
                    str(o.get("quote") or "")[:200], stamp])
    return out


def write_sheets(wb, rows: list[dict], keep_personal: bool) -> tuple[int, int]:
    """Replace each written company's rows in the Contacts and Company channels sheets. Returns (people, channels)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def sheet(name: str, cols: list[str]):
        ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        old = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)] if ws.max_row > 1 else []
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)
        for i, c in enumerate(cols, 1):
            h = ws.cell(row=1, column=i, value=c)
            h.font, h.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
            h.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = 16 if c in ("Domain", "Last Verified", "Channel") else 30
        ws.freeze_panes = "C2"
        return ws, old

    done = {str(r.get("Domain") or "").lower() for r in rows if r.get("_contacts") is not None}
    ws, old = sheet(CONTACTS_SHEET, CONTACT_COLUMNS)
    keep = [r for r in old if str(r[0] or "").lower() not in done]
    new = [cr for r in rows for cr in contact_rows(r, keep_personal)]
    for vals in keep + new:
        ws.append([("" if v is None else str(v)) or None for v in vals[: len(CONTACT_COLUMNS)]])
    done_c = {str(r.get("Domain") or "").lower() for r in rows if r.get("_channels") is not None}
    wc, oldc = sheet(CHANNELS_SHEET, CHANNEL_COLUMNS)
    keepc = [r for r in oldc if str(r[0] or "").lower() not in done_c]
    newc = [cr for r in rows for cr in channel_rows(r)]
    for vals in keepc + newc:
        wc.append([("" if v is None else str(v)) or None for v in vals[: len(CHANNEL_COLUMNS)]])
    for w in (ws, wc):
        for r in w.iter_rows(min_row=2):
            for c in r:
                c.alignment = Alignment(vertical="top", wrap_text=True)
    # The list sheets stay first, so a viewer that opens the first sheet shows the accounts.
    for name in (CONTACTS_SHEET, CHANNELS_SHEET):
        wb.move_sheet(name, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index(name))
    return len(new), len(newc)


# ---- live re-check of the quotes ------------------------------------------------------------------
def on_page(text: str, value: str, kind: str) -> bool:
    if kind == "phone":
        return bool(digits(value)) and digits(value) in digits(text)
    return value.lower() in (text or "").lower()


def _targets(row: dict):
    for p in row.get("_contacts") or []:
        for k in ("email", "phone"):
            v = p.get(k)
            if isinstance(v, dict) and str(v.get("source") or "").lower() != APOLLO:  # Apollo values are not on a page
                yield f"{p.get('name')} {k}", v, k
    ch = row.get("_channels") or {}
    for k in COMPANY_KEYS:
        v = ch.get(k)
        if isinstance(v, dict) and (EMAIL_RE.search(str(v.get("value"))) or has_phone(str(v.get("value")))):
            yield f"channel {k}", v, "email" if EMAIL_RE.search(str(v.get("value"))) else "phone"


def verify(row: dict) -> list[str]:
    """Fetch each email / phone source page (robots-aware, LinkedIn refused) and confirm the value is on it."""
    import lookup  # noqa: PLC0415 - network only for this command

    bad = []
    for label, v, kind in _targets(row):
        url = str(v.get("source") or "")
        host = urlsplit(url).hostname or ""
        if lookup.BLOCKED_HOSTS.search(host):
            bad.append(f"{label}: source {url} cannot be re-read (LinkedIn / Indeed); use a public page that shows it")
            continue
        try:
            if not lookup.allowed(url):
                bad.append(f"{label}: robots.txt disallows {url}")
                continue
            r = lookup.get(url)
            text = r.text if r.status_code == 200 else ""
        except Exception as ex:  # noqa: BLE001
            bad.append(f"{label}: {url} unavailable ({type(ex).__name__})")
            continue
        if not on_page(text, str(v.get("value")), kind):
            bad.append(f"{label}: {v.get('value')} is not on {url}; replace it with 'Not published (checked: ...)'")
        else:
            print(f"  ok  {label}: {v.get('value')} is on {url}")
    return bad


def _load(path: str) -> dict:
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    rows = payload.get("rows") if isinstance(payload, dict) else payload
    if not rows:
        raise SystemExit("no rows in " + path)
    return rows[0]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("cmd", choices=["check", "verify"])
    ap.add_argument("--draft", required=True)
    a = ap.parse_args()
    row = _load(a.draft)
    errs = validate_contacts(row)
    people = row.get("_contacts") or []
    n_apollo = sum(isinstance(p.get(k), dict) and str(p[k].get("source") or "").lower() == APOLLO
                   for p in people for k in ("email", "phone"))
    if n_apollo:
        print(f"Apollo values {n_apollo} (person id and email_status checked; not re-read from a page)")
    print(f"people {len(people)}; published emails {sum(isinstance(p.get('email'), dict) for p in people)}; "
          f"published phones {sum(isinstance(p.get('phone'), dict) for p in people)}")
    if a.cmd == "verify" and not errs:
        errs = verify(row)
    if errs:
        print("PROBLEMS:\n" + "\n".join(errs[:30]))
        sys.exit(1)
    print("contacts ok")


if __name__ == "__main__":
    main()
