"""Builder self-test for the account agent: field validators, India Connection evidence, starter
citations, team matching, CSV batch, roles summary, assemble and the idempotent Excel upsert.

  uv run python scripts/test_account.py

Uses temporary files only; the real workbook, memory.db, team.csv and company-agent are not touched.
"""
from __future__ import annotations

import json
import sys
import tempfile
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import assemble  # noqa: E402
import batch  # noqa: E402
import excel_upsert  # noqa: E402
import match_team  # noqa: E402
import memory  # noqa: E402
import roles_from_company_agent as roles  # noqa: E402
import starter  # noqa: E402
from common import TEAM_NOTE, add_source, citations, format_sources, parse_sources  # noqa: E402
from validate_rows import validate_payload, validate_row  # noqa: E402

TMP = Path(tempfile.mkdtemp(prefix="account-test-"))
import h1b  # noqa: E402
h1b.INDEX_PATH = TMP / "no-h1b-index.sqlite"  # the real LCA index (data/) is never read by the tests
RESULTS: list[tuple[bool, str]] = []
IC = "India Connection (US-based leader)"
ST = "Conversation Starter (cold email / conference)"
WARM = "Warm Connection / Common Ground"
BEST = "Best Person to Reach Out (HumanCloud side)"


def check(cond: bool, name: str) -> None:
    RESULTS.append((bool(cond), name))
    print(("PASS " if cond else "FAIL ") + name)


def row(**kw) -> dict:
    base = {
        "Company": "Acme", "Domain": "acme.com", "Product Type": "SaaS; Fintech", "Product?": "Yes",
        "What They Build (brief)": "Payment APIs for online shops [1]",
        "Company Brief (founded, size, customers, milestones)": "Founded 2014 [2]; approx. 300 employees (About page) [1]; raised $75M Series D in 2026 [3]",
        "India Office": "Not found (checked: site offices page, Wikidata, 2 Exa searches)",
        IC: "No evidence found (checked: team page, 2 leader bios, 3 Exa searches)",
        "Leadership Education & Background": "Jane Doe: ex-Stripe [1]",
        "HQ / US Office": "San Francisco, CA [1]", "Funding": "$75M Series D, June 2026 [3]",
        "Who to Contact": "Jane Doe, CTO [1]", "Sources": "[1] https://acme.com/about\n[2] https://www.wikidata.org/wiki/Q1\n[3] https://acme.com/press/series-d",
        "Confidence": "Medium", "Last Checked": "2026-10-15",
        WARM: TEAM_NOTE, BEST: TEAM_NOTE,
    }
    base.update(kw)
    return base


def errs(r: dict, team=None, draft=False) -> list[str]:
    return validate_row(r, 0, draft=draft, team=team)


# ---- search depth and freshness (0.6.2) ----------------------------------------------------------
from datetime import date as _d  # noqa: E402
from validate_rows import stale_flags  # noqa: E402

BRIEF = "Company Brief (founded, size, customers, milestones)"
check(any("India Office is empty" in e for e in errs(row(**{"India Office": ""}), team=[])), "depth: an empty reason field is refused")
check(errs(row(Funding="Not found (checked: site press page, EDGAR Form D, 3 Exa searches)"), team=[]) == [],
      "depth: 'Not found (checked: ...)' needs no citation")
check(any("is empty" in e for e in errs(row(Funding="Not found (checked: x)"), team=[])) or
      any("no citation" in e for e in errs(row(Funding="Not found (checked: x)"), team=[])),
      "depth: a 'checked' list too short to name a source is not a reason")
check(errs(row(**{"Product Type": "Not found (checked: home page, about page)", "Product?": "Not found (checked: pricing page)"}), team=[]) == [],
      "depth: closed fields accept a not-found reason")
check(any("employee size" in e for e in errs(row(**{BRIEF: "Founded 2014 [2]"}), team=[])), "depth: the brief must state size")
check(errs(row(**{BRIEF: "Founded 2014 [2]; size not found (checked: Wikidata P1128, Wikipedia, site, 3 Exa searches)"}), team=[]) == [],
      "depth: 'size not found (checked: ...)' satisfies the size rule")
check(errs(row(**{BRIEF: "Founded 2014 [2]; approx. 3,000 employees per 2025 press release [1]"}), team=[]) == [],
      "depth: a cited approx. size passes")
T = _d(2026, 10, 16)
check(any("2 or more years old" in f for f in stale_flags(row(Funding="Series F, $375M, December 2021 [3]"), T)),
      "stale: a 2021 round in 2026 is flagged")
check(stale_flags(row(Funding="Series F, $375M, December 2021 [3]; No newer round found (checked: 3 Exa searches, GDELT news, EDGAR Form D)"), T) == [],
      "stale: an old round with 'No newer round found (checked: ...)' passes")
check(stale_flags(row(Funding="Series C, $50M, March 2025 [3]"), T) == [], "stale: a 2025 round is fresh")
check(stale_flags(row(Funding="Public (NASDAQ: FRSH) since 2021 [3]"), T) == [], "stale: a listed company is not flagged by year")
check(not any("2026" in f for f in stale_flags(row(Funding="Series B in 2019 [3]", Sources="[3] https://x.com/2019/a"), T)) and
      any("2019" in f for f in stale_flags(row(Funding="Series B in 2019 [3]"), T)), "stale: the newest year names the flag")
check(any("conflicting CTO" in f for f in stale_flags(row(**{"Who to Contact": "Jane Doe, CTO [1]; Sam Roe, CTO and co-founder [2]"}), T)),
      "stale: two CTOs are a conflict")
check(any("conflicting CTO" in f for f in stale_flags(row(**{"Who to Contact": "Jane Doe, CTO [1] (a 2026 article calls Sam Roe CTO, so confirm the current holder [2])"}), T)),
      "stale: 'confirm the current' wording is a conflict")
check(stale_flags(row(**{"Who to Contact": "Jane Doe, CTO [1]; Sam Roe, CTO [2] (conflict: Doe per [1] 2022 vs Roe per [2] Aug 2026; newer source used)"}), T) == [],
      "stale: a noted conflict passes")
check(stale_flags(row(**{"Who to Contact": "Jane Doe, CTO [1]; Sam Roe, CEO [2]"}), T) == [], "stale: one CTO and a CEO are no conflict")
check(any("2 or more years old" in e for e in validate_row(row(Funding="Seed, 2020 [3]"), 0, team=[], today=T)),
      "stale: validate_row refuses an unflagged stale round")

# ---- field validators ---------------------------------------------------------------------------
check(errs(row(), team=[]) == [], "a complete cited row is valid")
check(any("Sources is required" in e for e in errs(row(Sources=""), team=[])), "Sources is required")
check(any("Last Checked is required" in e for e in errs(row(**{"Last Checked": ""}), team=[])), "Last Checked required on a final row")
check(not any("Last Checked" in e for e in errs(row(**{"Last Checked": ""}), team=[], draft=True)), "a draft may leave Last Checked to assemble.py")
check(any("Domain must be" in e for e in errs(row(Domain="https://acme.com/x"), team=[])), "Domain must be a bare host")
check(any("Product Type" in e for e in errs(row(**{"Product Type": "Crypto magic"}), team=[])), "Product Type must come from the list")
check(any("Product?" in e for e in errs(row(**{"Product?": "Maybe"}), team=[])), "Product? is Yes / No / Both")
check(any("Confidence" in e for e in errs(row(Confidence="Sure"), team=[])), "Confidence is High / Medium / Low")
check(any("no citation" in e for e in errs(row(Funding="$75M Series D"), team=[])), "a cited field without a citation is refused")
check(any("cites [9]" in e for e in errs(row(Funding="$75M [9]"), team=[])), "a citation must exist in Sources")
check(any("not in Sources" in e for e in errs(row(Funding="$75M https://news.example.com/x"), team=[])), "a URL in a cell must be in Sources")
check(any("not `[n] https://url`" in e for e in errs(row(Sources="acme.com/about"), team=[])), "Sources lines must be `[n] url`")
check(any("unknown column" in e for e in errs(row(Notes="x"), team=[])), "unknown columns are refused")
check(errs(row(_leaders=[{"name": "Jane"}]), team=[]) == [], "underscore working keys are allowed")
check(any("email or phone" in e for e in errs(row(Funding="ask cfo@acme.com [3]"), team=[])), "no email outside Who to Contact")
check(any("email or phone" in e for e in errs(row(**{"HQ / US Office": "SF, +1 415 555 0100 [1]"}), team=[])), "no phone outside Who to Contact")
check(any("self-published" in e for e in errs(row(**{"Who to Contact": "Jane Doe, CTO, jane@acme.com [1]"}), team=[])), "an email needs 'self-published [n]'")
check(errs(row(**{"Who to Contact": "Jane Doe, CTO, jane@acme.com (self-published [1])"}), team=[]) == [], "a self-published email with a citation passes")
check(errs(row(**{"Company Brief (founded, size, customers, milestones)": "2014-2020 growth phase [2]; approx. 300 employees [1]"}), team=[]) == [], "a year range is not a phone number")
check(any("appears twice" in e for e in validate_payload({"rows": [row(), row()]})), "a domain appears once per file")

# ---- India Connection evidence rule ---------------------------------------------------------------
check(any("needs evidence" in e for e in errs(row(**{IC: "Yes: CTO studied at IIT Delhi"}), team=[])), "India Connection 'Yes' without an evidence URL is rejected")
check(any("needs evidence" in e or "cites [7]" in e for e in errs(row(**{IC: "Yes: CTO studied at IIT Delhi [7]"}), team=[])), "India Connection citing a missing source is rejected")
check(errs(row(**{IC: "Yes: CTO Jane Doe (San Francisco) studied at IIT Delhi [1]"}), team=[]) == [], "India Connection with cited evidence passes")
check(errs(row(**{IC: "No evidence found (checked: team page, 4 Exa searches)"}), team=[]) == [], "'No evidence found (checked: ...)' passes")
check(any("say what was searched" in e for e in errs(row(**{IC: "No evidence found"}), team=[])), "a bare 'No evidence found' is refused")
check(any("India Connection must be" in e for e in errs(row(**{IC: "No"}), team=[])), "a bare 'No' is refused (use 'No evidence found')")
for w in ("Yes: surname suggests an Indian origin [1]", "Yes: name sounds Indian [1]", "Yes: appears to be of Indian descent [1]",
          "Yes: ethnicity Indian [1]", "Yes: Indian-American CTO [1]"):
    check(any("self-published facts only" in e for e in errs(row(**{IC: w}), team=[])), f"inference wording refused: {w[:34]}")
check(errs(row(**{IC: "Yes: CTO led the Indian Institute of Technology Delhi alumni chapter [1]"}), team=[]) == [], "'Indian Institute' is a fact, not an inference")

# ---- Conversation Starter citation check ----------------------------------------------------------
good_starter = "Email: Congratulations on the $75M Series D [3]. Conference: How is the payments API roadmap shaping up? Facts used: [3]"
check(errs(row(**{ST: good_starter}), team=[]) == [], "a starter built from cited row facts passes")
check(any("must cite" in e for e in errs(row(**{ST: "Congrats on the round!"}), team=[])), "a starter without a citation is refused")
check(any("cites [8]" in e for e in errs(row(**{ST: "Congrats on the round [8]"}), team=[])), "a starter citing a source not in Sources is refused")
check(any("not in Sources" in e for e in errs(row(**{ST: "Congrats https://techcrunch.com/x"}), team=[])), "a starter URL must be in Sources")
check(any("no other cell" in e for e in errs(row(**{ST: "Congrats on the $90M Series E [3]"}), team=[])), "a starter number not in the row is refused")
check(any("business facts only" in e for e in errs(row(**{ST: "As a fellow Indian, congrats [3]"}), team=[])), "starter refuses identity wording")
check(any("business facts only" in e for e in errs(row(**{ST: "Saw your India Connection [3]"}), team=[])), "starter never mentions the India Connection field")

rows_file = TMP / "acme.rows.json"
rows_file.write_text(json.dumps({"meta": {"list": "Accounts"}, "rows": [row()]}))
sys.argv = ["starter.py", "set", "--rows", str(rows_file), "--text", good_starter]
starter.main()
check(json.loads(rows_file.read_text())["rows"][0].get(ST) == good_starter, "starter.py set writes a valid starter")
sys.argv = ["starter.py", "set", "--rows", str(rows_file), "--text", "Congrats on the $90M round [3]"]
starter.main()
check(str(json.loads(rows_file.read_text())["rows"][0].get(ST)).startswith("Not found (checked:"),
      "starter.py set writes the not-found reason when the check fails")

# ---- team matching --------------------------------------------------------------------------------
team_csv = TMP / "team.csv"
team_csv.write_text(
    "name,role,email,seniority,regions,past_employers,colleges,cities_lived,communities,conferences,industries_known,linkedin_url,consent,default_for\n"
    "Priya Sharma,\"Director, Sales (US)\",p@x.com,Director,US-East; India,Infosys; Persistent Systems,COEP Pune; Carnegie Mellon,Pune; Boston,TiE Boston,SaaStr 2026,fintech,,yes,US\n"
    "Rahul Mehta,\"VP, Delivery\",r@x.com,VP,India,TCS; Wipro,IIT-B,Mumbai,NASSCOM,NASSCOM Technology & Leadership Forum 2025,enterprise software,,yes,India\n"
    "No Consent,Engineer,n@x.com,IC,US,Infosys,IIT Bombay,,,,,,no,\n"
    ",Missing name,,IC,US,,,,,,,,yes,\n", encoding="utf-8")
team, problems = match_team.load_team(team_csv)
check([m["name"] for m in team] == ["Priya Sharma", "Rahul Mehta"], "team.csv: only consented, complete rows are used")
check(len(problems) == 2, "team.csv: skipped rows are reported")
leaders = [
    {"name": "Jane Doe", "title": "CTO", "location": "San Francisco, US", "employers": ["Tata Consultancy Services"],
     "colleges": ["Indian Institute of Technology Bombay"], "conferences": ["NASSCOM Technology & Leadership Forum 2026"], "cite": "[1]"},
    {"name": "Sam Roe", "title": "VP Engineering", "location": "Boston, USA", "employers": ["Persistent"], "cities": ["Boston"],
     "conferences": ["SaaStr 2026"], "cite": "[2]"},
]
pairs = match_team.match_pairs(team, leaders)
rj = next(p for p in pairs if p["member"]["name"] == "Rahul Mehta" and p["leader"]["name"] == "Jane Doe")
check(rj["score"] == 3 + 3 + 1, "alias match: TCS = Tata Consultancy Services (3), IIT-B = IIT Bombay (3), same event other year (1)")
ps = next(p for p in pairs if p["member"]["name"] == "Priya Sharma" and p["leader"]["name"] == "Sam Roe")
check(ps["score"] == 3 + 1 + 2, "shared employer (3) + city (1) + same conference and year (2)")
warm, best = match_team.warm_and_best(team, leaders, ["Fintech"])
check("Rahul Mehta and CTO Jane Doe both worked at Tata Consultancy Services [1]" in warm and warm.count("common ground") == 2, "warm cell names the two best matches with their citations")
check("know each other" not in warm and "common ground" in warm, "warm cell claims common ground only")
check(best.startswith("Rahul Mehta") or best.startswith("Priya Sharma"), "best person is a team member")
tie_team = [dict(team[0], name="A Manager", seniority="Manager"), dict(team[0], name="B Director", seniority="Director")]
_, best_tie = match_team.warm_and_best(tie_team, [leaders[1]])
check(best_tie.startswith("B Director"), "a score tie goes to the better seniority fit (VP contact -> Director)")
_, best_default = match_team.warm_and_best(team, [{"name": "X", "title": "CEO", "location": "New York, US"}])
check(best_default.startswith("Priya Sharma") and "default owner for US" in best_default, "no match -> default owner for the region")
check(match_team.warm_and_best([], leaders) == ("", ""), "no team -> both cells empty")
check(errs(row(**{WARM: "x", BEST: "y"}), team=[]) != [], "no team rows -> the two cells must be blank or 'add team.csv'")
check(errs(row(**{BEST: "Priya Sharma (Director): score 6"}), team=["Priya Sharma"]) == [], "with team rows, Best Person names a team member")
check(any("must name a person" in e for e in errs(row(**{BEST: "Somebody Else"}), team=["Priya Sharma"])), "Best Person outside team.csv is refused")
check(any("common ground only" in e for e in errs(row(**{WARM: "Priya knows Jane well [1]"}), team=["Priya Sharma"])), "Warm Connection never claims they know each other")

# ---- CSV batch -----------------------------------------------------------------------------------
csv_file = TMP / "companies.csv"
csv_file.write_text("Company,Domain,List\nRazorpay,https://www.razorpay.com/,Q4 India\nAcme,acme.com,\nacme again,acme.com,\n"
                    "Bad Domain,not a domain,\n,,\nNo Domain Co,,\n", encoding="utf-8")
items, probs = batch.read_csv(csv_file)
check([i["company"] for i in items] == ["Razorpay", "Acme", "Bad Domain", "No Domain Co"], "batch: duplicates dropped, blank lines ignored")
check(items[0]["domain"] == "razorpay.com" and items[0]["list"] == "Q4 India", "batch: domain normalised, list kept")
check(items[1]["list"] == "Accounts", "batch: default sheet is Accounts")
check(items[2]["domain"] == "" and any("not a website domain" in p for p in probs), "batch: a bad domain is reported and verified instead")
bad_csv = TMP / "bad.csv"
bad_csv.write_text("name,site\nAcme,acme.com\n")
check(batch.read_csv(bad_csv)[0] == [] and "company" in batch.read_csv(bad_csv)[1][0], "batch: a CSV without a company column is refused")
batch.STATE = TMP / "batch.json"
batch.recent_domains = lambda days: {"acme.com": date.today().isoformat()}
st = batch.plan(csv_file, 14, False, "Accounts")
check([i["state"] for i in st["items"]] == ["todo", "skipped", "todo", "todo"], "batch: a company checked recently is skipped (resumable)")
st = batch.plan(csv_file, 14, True, "Accounts")
check(all(i["state"] == "todo" for i in st["items"]), "batch: --refresh rescans everything")

# ---- roles summary -------------------------------------------------------------------------------
sample = [
    {"Job Title": "Senior Backend Engineer (Go)", "Location / Work Mode": "Bengaluru, India", "Status": "Open", "Last Checked": "2026-10-15", "Career Page": "https://acme.com/careers"},
    {"Job Title": "Android Developer", "Location / Work Mode": "Pune", "Status": "Open", "Last Checked": "2026-10-15"},
    {"Job Title": "Account Executive", "Location / Work Mode": "New York", "Status": "Open", "Last Checked": "2026-10-15"},
    {"Job Title": "Data Engineer", "Location / Work Mode": "Remote", "Status": "Closed", "Last Checked": "2026-10-01"},
]
s = roles.summarise(sample)
check((s["open"], s["engineering"], s["india"]) == (3, 2, True), "roles: open, engineering and India counts")
check(s["source"] == "https://acme.com/careers", "roles: cites the careers page")
check(roles.cell_text(s).startswith("2 open engineering roles of 3 open roles"), "roles: cell text")

# ---- assemble + Excel upsert ---------------------------------------------------------------------
draft = row(**{"Last Checked": "", WARM: "", BEST: ""})
draft["_leaders"] = leaders
draft_file = TMP / "acme.draft.json"
draft_file.write_text(json.dumps({"meta": {"list": "Accounts"}, "rows": [draft]}))
roles_file = TMP / "acme.roles.json"
roles_file.write_text(json.dumps({**s, "cell": roles.cell_text(s)}))
assemble.RUNS_DIR = TMP
assemble.load_team = lambda: ([], ["team.csv not found"])
sys.argv = ["assemble.py", "--draft", str(draft_file), "--roles", str(roles_file)]
assemble.main()
out = json.loads((TMP / "acme.rows.json").read_text())["rows"][0]
srcs, _ = parse_sources(out["Sources"])
check(out["Development Requirement (current)"].endswith(f"[{max(srcs)}]") and srcs[max(srcs)] == "https://acme.com/careers", "assemble: roles cell cites the careers page as a new source")
check(out[WARM] == TEAM_NOTE and out[BEST] == TEAM_NOTE, "assemble: empty team.csv -> 'add team.csv'")
sys.argv = ["assemble.py", "--draft", str(draft_file), "--roles", str(TMP / "missing.roles.json")]
assemble.main()
nr = json.loads((TMP / "acme.rows.json").read_text())
check(nr["rows"][0]["Development Requirement (current)"].startswith("Not found (checked:") and
      "Development Requirement (current)" in nr["meta"]["not_found"], "assemble: no roles -> a not-found reason, listed in meta")
sys.argv = ["assemble.py", "--draft", str(draft_file), "--roles", str(roles_file)]
assemble.main()
check(out["Last Checked"] == date.today().isoformat(), "assemble: Last Checked is today")
check(citations("a [1, 3] b [2]") == [1, 3, 2] and add_source({1: "u"}, "v") == 2 and format_sources({2: "b", 1: "a"}) == "[1] a\n[2] b", "citation helpers")


# ---- leaders and contacts (0.6.2) -----------------------------------------------------------------
import contacts as C  # noqa: E402

NP = "Not published (checked: site team page, contact page, 2 Exa searches)"
JANE = {"name": "Jane Doe", "title": "CTO", "location": "San Francisco, CA", "source": "https://acme.com/about",
        "verified": "2026-10-15", "email": NP, "phone": NP, "linkedin": "https://www.linkedin.com/in/janedoe",
        "github": "https://github.com/janedoe", "channels_checked": "site team page, 2 Exa searches"}
CH = {"press_email": {"value": "press@acme.com", "source": "https://acme.com/press", "quote": "Media: press@acme.com"},
      "contact_form": {"value": "https://acme.com/contact", "source": "https://acme.com/contact"},
      "offices": [{"label": "HQ", "address": "1 Main St, San Francisco, CA 94105", "phone": "+1 415 555 0100",
                   "source": "https://acme.com/contact", "quote": "1 Main St, San Francisco, CA 94105 +1 415 555 0100"}],
      "checked": "contact, press and investors pages"}


def cerrs(people=None, chans=None, **kw):
    r = row(**kw)
    if people is not None:
        r["_contacts"] = people
    if chans is not None:
        r["_channels"] = chans
    return errs(r, team=[])


check(cerrs([JANE], CH) == [], "contacts: a sourced person and company channels pass")
pub = {**JANE, "email": {"value": "jane@acme.com", "quote": "Write to me at jane@acme.com", "source": "https://janedoe.dev"}}
check(cerrs([pub]) == [], "contacts: a self-published email with its quote passes")
made = {**JANE, "email": {"value": "jane.doe@acme.com", "quote": "Jane Doe, CTO", "source": "https://acme.com/about"}}
check(any("constructed or unsourced" in e for e in cerrs([made])), "contacts: a constructed email (not in the quote) is refused")
check(any("constructed" in e for e in cerrs([{**JANE, "email": "jane.doe@acme.com"}])), "contacts: a bare email with no source is refused")
check(any("source must be" in e for e in cerrs([{**JANE, "email": {"value": "jane@acme.com", "quote": "jane@acme.com"}}])),
      "contacts: an email without a source URL is refused")
check(any("refused wording" in e for e in cerrs([{**JANE, "email": {"value": "jane@acme.com", "quote": "jane@acme.com",
      "source": "https://acme.com", "kind": "likely pattern"}}])), "contacts: 'likely' / 'pattern' wording is refused")
check(any("guessed phone" in e for e in cerrs([{**JANE, "phone": {"value": "+1 415 555 0199", "kind": "personal",
      "quote": "Call our office", "source": "https://acme.com"}}])), "contacts: a phone not in its quote is refused")
check(any("an email or phone needs" in e for e in cerrs([{**JANE, "phone": "+1 415 555 0199"}])), "contacts: an unsourced phone is refused")
check(cerrs([{**JANE, "phone": {"value": "+1 415 555 0100", "kind": "company line", "quote": "HQ +1 415 555 0100",
      "source": "https://acme.com/contact"}}]) == [], "contacts: the company's published line, labelled, passes")
check(any("kind must be" in e for e in cerrs([{**JANE, "phone": {"value": "+1 415 555 0100", "quote": "+1 415 555 0100",
      "source": "https://acme.com/contact"}}])), "contacts: a phone needs kind personal or company line")
check(any("paid enrichment" in e for e in cerrs([{**JANE, "source": "https://rocketreach.co/jane-doe"}])),
      "contacts: a data-broker source is refused")
check(any("linkedin.com/in" in e for e in cerrs([{**JANE, "linkedin": "https://www.linkedin.com/company/acme"}])),
      "contacts: LinkedIn must be a /in/ profile URL")
check(any("keep the 6" in e for e in cerrs([{**JANE, "name": f"P{i}"} for i in range(7)])), "contacts: at most 6 people")
check(any("email or phone" in e for e in cerrs([{**JANE, "location": "SF, jane@acme.com"}])), "contacts: no email outside the email field")
check(any("channels_checked" in e for e in cerrs([{k: v for k, v in JANE.items() if k != "channels_checked"}])),
      "contacts: a missing channel needs channels_checked")
check(any("future" in e for e in cerrs([{**JANE, "verified": "2999-01-01"}])), "contacts: verified date cannot be in the future")
check(any("constructed or unsourced" in e for e in cerrs(chans={**CH, "sales": {"value": "sales@acme.com",
      "source": "https://acme.com", "quote": "Talk to sales"}})), "contacts: a company email needs its quote too")
check(any("phone" in e for e in cerrs(chans={**CH, "offices": [{**CH["offices"][0], "quote": "1 Main St, San Francisco, CA 94105"}]})),
      "contacts: an office phone not in its quote is refused")
check(C.on_page("<a href='mailto:jane@acme.com'>", "jane@acme.com", "email") and C.on_page("Tel +1 (415) 555-0100", "+14155550100", "phone")
      and not C.on_page("Jane Doe", "jane@acme.com", "email"), "contacts: verify finds the value on the page")
r1 = row(**{"Last Checked": "2026-10-15"})
r1["_contacts"], r1["_channels"] = [pub, {**JANE, "name": "Sam Roe", "title": "CEO", "phone": {"value": "+1 415 555 0100",
    "kind": "company line", "quote": "+1 415 555 0100", "source": "https://acme.com/contact"}}], CH
cr = C.contact_rows(r1, keep_personal=False)
check(cr[0][5] == C.WITHHELD and cr[1][7].endswith("(company line)"), "contacts: without --contacts the personal email is withheld, the company line kept")
check(C.contact_rows(r1, keep_personal=True)[0][5] == "jane@acme.com", "contacts: with --contacts the published email is kept")
check(C.contact_rows(r1, True)[0][13].startswith("Not published (checked:"), "contacts: a channel left out says what was checked")
check([c[2] for c in C.channel_rows(r1)][:5] == ["Press / media", "Sales", "Partnerships", "Investor relations", "Contact form"]
      and C.channel_rows(r1)[1][3].startswith("Not published"), "contacts: company channels list every channel")
from common import safe_sheet_name as _ssn  # noqa: E402
check(_ssn("Contacts") != "Contacts" and _ssn("company channels") != "company channels", "contacts: a list cannot be named Contacts")

memory.DB_PATH = TMP / "memory.db"
memory_connect = memory.connect
memory.connect = lambda path=None: memory_connect(TMP / "memory.db")
wb = TMP / "Account_Research.xlsx"
for _ in range(2):
    sys.argv = ["excel_upsert.py", "--rows", str(TMP / "acme.rows.json"), "--workbook", str(wb)]
    excel_upsert.main()
from openpyxl import load_workbook  # noqa: E402

ws = load_workbook(wb)["Accounts"]
check(ws.max_row == 2, "excel: two upserts of one company leave one row (idempotent)")
check(ws.cell(1, 1).value == "Company" and ws.cell(1, 19).value == "Last Checked", "excel: 19 columns in order")
check(memory.last_checked(memory.connect(), "acme.com") == date.today().isoformat(), "memory: the row is recorded")
cp = TMP / "acme.contacts.rows.json"
cpay = json.loads((TMP / "acme.rows.json").read_text())
cpay["rows"][0]["_contacts"], cpay["rows"][0]["_channels"] = r1["_contacts"], CH
cp.write_text(json.dumps(cpay))
for _ in range(2):
    sys.argv = ["excel_upsert.py", "--rows", str(cp), "--workbook", str(wb)]
    excel_upsert.main()
wbx = load_workbook(wb)
check(wbx.sheetnames[0] == "Accounts" and wbx["Contacts"].max_row == 3 and wbx["Contacts"].cell(1, 1).value == "Domain",
      "excel: Contacts sheet, one row per person, linked by Domain, idempotent")
check(wbx["Contacts"].cell(2, 6).value == C.WITHHELD, "excel: personal email not stored without --contacts")
check(wbx["Company channels"].max_row == 1 + 6, "excel: Company channels sheet, 5 channels + 1 office")
sys.argv = ["excel_upsert.py", "--rows", str(cp), "--workbook", str(wb), "--contacts"]
excel_upsert.main()
check(load_workbook(wb)["Contacts"].cell(2, 6).value == "jane@acme.com", "excel: --contacts keeps the published email")

# ---- Apollo.io client (0.6.3): mocked HTTP only, no real key or call ------------------------------
import io  # noqa: E402
import contextlib  # noqa: E402
import apollo as A  # noqa: E402

FAKE_KEY = "TESTKEY0123456789abcdef"
PID1, PID2 = "64a1b2c3d4e5f60718293a4b", "64a1b2c3d4e5f60718293a4c"


class Resp:
    def __init__(self, status: int, body=None, headers=None):
        self.status_code, self._body, self.headers = status, body or {}, headers or {}

    def json(self):
        return self._body


class FakeApollo:
    """Answers like api.apollo.io; records every call; can be told to fail one endpoint with a status."""

    def __init__(self, fail: dict | None = None):
        self.calls, self.fail = [], fail or {}

    def request(self, method, url, headers=None, json=None, params=None, timeout=None, allow_redirects=True):
        path = url.split("/api/v1", 1)[1]
        self.calls.append((method, path, dict(headers or {}), allow_redirects))
        if path in self.fail:
            st = self.fail[path]
            return Resp(st, {"error": "nope"}, {"retry-after": "0"} if st == 429 else {})
        if path == "/organizations/enrich":
            return Resp(200, {"organization": {"name": "Acme", "founded_year": 2014, "estimated_num_employees": 300}})
        if path == "/mixed_people/api_search":
            return Resp(200, {"people": [{"id": PID1, "first_name": "Jane", "title": "CTO"},
                                         {"id": PID2, "first_name": "Raj", "title": "VP Engineering"}]})
        if path == "/people/match":
            if json["id"] == PID1:
                return Resp(200, {"person": {"id": PID1, "name": "Jane Doe", "title": "CTO", "city": "San Francisco",
                                             "country": "United States", "email": "jane@acme.com", "email_status": "verified",
                                             "linkedin_url": "http://www.linkedin.com/in/janedoe",
                                             "phone_numbers": [{"sanitized_number": "+14155550100", "type": "mobile"}]}})
            return Resp(200, {"person": {"id": PID2, "name": "Raj Rao", "title": "VP Engineering", "country": "India",
                                         "email": "raj@acme.com", "email_status": "guessed"}})
        return Resp(404)


def apollo_run(fake, db, budget=6, now=None):
    con = A.cache_db(db)
    cl = A.Client(FAKE_KEY, budget, con, session=fake, sleep=lambda s: None, **({"now": now} if now else {}))
    return A.enrich("acme.com", cl), cl


adb = TMP / "apollo.db"
fake = FakeApollo()
res, cl = apollo_run(fake, adb)
jane = next(p for p in res["people"] if p["apollo_id"] == PID1)
raj = next(p for p in res["people"] if p["apollo_id"] == PID2)
check(isinstance(jane["email"], dict) and jane["email"]["value"] == "jane@acme.com" and jane["email"]["email_status"] == "verified"
      and jane["email"]["source"] == "apollo", "apollo: a verified email is kept, with source apollo and the person id")
check(isinstance(raj["email"], str) and raj["email"].startswith("Not published") and "raj@acme.com" not in json.dumps(raj),
      "apollo: a guessed email is dropped (Not published, the address not kept)")
check(jane["phone"]["value"] == "+14155550100" and jane["phone"]["kind"] == "personal", "apollo: a phone Apollo returned directly is kept")
check(all(c[2].get("X-Api-Key") == FAKE_KEY and c[3] is False for c in fake.calls) and len(fake.calls) == 5,
      "apollo: key sent only as the X-Api-Key header, redirects not followed; 1 org + 2 searches + 2 matches")
check(cl.credits == 2 and cl.org_credits == 1,
      "apollo: credits counted (2 people, 1 organization)")
check(errs({**row(), "_contacts": [jane]}) == [], "apollo: the Apollo person passes the contact rules")
fake2 = FakeApollo()
res2, cl2 = apollo_run(fake2, adb)
check(fake2.calls == [] and cl2.credits == 0 and cl2.cached == 5 and len(res2["people"]) == 2,
      "apollo: the 30-day cache stops a second spend on a rescan")
from datetime import datetime as _dt, timedelta as _td  # noqa: E402
fake3 = FakeApollo()
apollo_run(fake3, adb, now=lambda: _dt.now() + _td(days=31))
check(len(fake3.calls) == 5, "apollo: a cached answer older than 30 days is fetched again")
check(A.forget("acme.com", A.cache_db(adb)) == 5, "apollo: forget <domain> clears that company's cached answers")
fake4 = FakeApollo()
res4, cl4 = apollo_run(fake4, TMP / "apollo-budget.db", budget=1)
check(cl4.credits == 1 and len(res4["people"]) == 1 and sum(c[1] == "/people/match" for c in fake4.calls) == 1,
      "apollo: APOLLO_MAX_CREDITS caps people-match calls per company")
for st, kind in ((402, "plan"), (403, "plan"), (401, "key"), (429, "rate")):
    f = FakeApollo(fail={"/people/match": st})
    r_, c_ = apollo_run(f, TMP / f"apollo-{st}.db")
    check(r_["error"] and str(st) in r_["error"] and r_["people"] == [] and c_.credits == 0 and FAKE_KEY not in r_["error"]
          and (st != 429 or sum(c[1] == "/people/match" for c in f.calls) == 2),
          f"apollo: HTTP {st} stops cleanly, no credits counted" + (" (one retry)" if st == 429 else ""))
envf = TMP / "apollo.env"
envf.write_text("EXA_API_KEY=x\nAPOLLO_API_KEY=" + FAKE_KEY + "\nAPOLLO_MAX_CREDITS=3\n")
check(A.config(envf) == (FAKE_KEY, 3), "apollo: key and budget read from .env")
envf.write_text("APOLLO_API_KEY=\n")
check(A.config(envf)[0] == "" and A.config(TMP / "missing.env") == ("", 6), "apollo: an empty or missing .env means not configured, budget 6")
_old = A.ENV_PATH
A.ENV_PATH = TMP / "missing.env"
buf = io.StringIO()
with contextlib.redirect_stdout(buf):
    rc = A.main(["enrich", "acme.com"])
A.ENV_PATH = _old
check(rc == 0 and "Apollo not configured" in buf.getvalue(), "apollo: no key prints 'Apollo not configured' and exits 0")
buf = io.StringIO()
with contextlib.redirect_stdout(buf):
    A.print_result(res, cl, False)
    A.print_result(res, cl, True)
    A.print_result(apollo_run(FakeApollo(fail={"/organizations/enrich": 401}), TMP / "apollo-k.db")[0], cl, False)
out = buf.getvalue()
check(FAKE_KEY not in out and "Apollo credits used: 3" in out and "Source: Apollo (verified)" in out,
      "apollo: the key never appears in output; credits used printed")
# contact rules for Apollo values
AP = {"name": "Jane Doe", "title": "CTO", "location": "San Francisco", "source": "apollo", "apollo_id": PID1,
      "verified": "2026-10-16", "email": {"value": "jane@acme.com", "source": "apollo", "apollo_id": PID1, "email_status": "verified"},
      "phone": "Not published (checked: Apollo people match)", "channels_checked": "Apollo people match"}
def aerrs(**kw):
    return C.validate_contacts({"Company": "Acme", "_contacts": [{**AP, **kw}]}, date(2026, 10, 16))
check(aerrs() == [], "contacts: an Apollo-verified email with person id passes")
check(any("person id" in e for e in aerrs(email={**AP["email"], "apollo_id": ""})), "contacts: an Apollo email without the person id is refused")
check(any("verified" in e for e in aerrs(email={**AP["email"], "email_status": "likely to engage"})), "contacts: an Apollo email not 'verified' is refused")
check(any("person id" in e for e in aerrs(apollo_id="")), "contacts: an Apollo person without the person id is refused")
check(any("paid enrichment" in e for e in aerrs(source="https://app.apollo.io/people/1")), "contacts: Apollo web pages stay refused as sources")
check(any("constructed" in e for e in aerrs(email="jane.doe@acme.com")), "contacts: a bare (guessed) email stays refused")
ar = C.contact_rows({"Domain": "acme.com", "Company": "Acme", "_contacts": [AP]}, keep_personal=True)[0]
check(ar[5] == "jane@acme.com" and ar[6].startswith("Source: Apollo (verified)") and PID1 in ar[6],
      "excel: Contacts sheet shows 'Source: Apollo (verified)' with the person id")
check(C.contact_rows({"Domain": "acme.com", "Company": "Acme", "_contacts": [AP]}, keep_personal=False)[0][5] == C.WITHHELD,
      "excel: an Apollo email still needs --contacts to be stored")

# ---- Round 1 data improvements: Evidence sheet, freshness, Form D, GLEIF, DNS, CRM export ----------
import csv as _csv  # noqa: E402
import struct as _struct  # noqa: E402
import time as _time  # noqa: E402
import evidence as EV  # noqa: E402
import export_crm as X  # noqa: E402
import lookup as L  # noqa: E402

check(EV.split_facts("Founded 2014 [3]; approx. 5,200 employees (March 2026) [16]; serves 8M [4] [5]") ==
      [("Founded 2014", [3]), ("approx. 5,200 employees (March 2026)", [16]), ("serves 8M", [4]), ("serves 8M", [5])],
      "evidence: a cell splits into one fact per [n] clause; adjacent markers share the clause")
check(EV.split_facts("Payments [1, 2]") == [("Payments", [1, 2])], "evidence: [1, 2] gives both sources to one fact")
# Real Razorpay cells (runs/razorpay.rows.json): a citation inside a parenthetical belongs to it; no cut inside brackets.
_rz = ("Yes: Country Head US Volker Steinle (San Francisco per a profile snippet [19]) wrote in a June 2026 post that he "
       "first interned in Bangalore during his MBA, at Infosys, and spent his first Razorpay week in Bangalore [20]; Infosys "
       "is also listed among his past employers [21].")
_rzf = EV.split_facts(_rz)
check(_rzf == [("Yes: Country Head US Volker Steinle (San Francisco per a profile snippet)", [19]),
               ("Yes: Country Head US Volker Steinle (San Francisco per a profile snippet) wrote in a June 2026 post that he "
                "first interned in Bangalore during his MBA, at Infosys, and spent his first Razorpay week in Bangalore", [20]),
               ("Infosys is also listed among his past employers", [21])],
      "evidence: Razorpay India Connection keeps the parenthetical whole and never starts a fact with ')'")
check(all(f.count("(") == f.count(")") for f, _ in _rzf), "evidence: every Razorpay fact has balanced brackets")
_rzw = EV.split_facts("Shashank Kumar, co-founder, quoted as CTO in August 2026 press [13] (conflict: CTO per [13] 2026-08-04 vs MD "
                      "and Co-founder per [8] 2025-11-10 and [14] 2026-03-12; former CTO Murali Brahmadesam announced his exit "
                      "on 2025-07-07 [22]); Prabu Rambadran, Senior Vice President, Engineering since November 2025 [8].")
_conf = ("Shashank Kumar, co-founder, quoted as CTO in August 2026 press (conflict: CTO per 2026-08-04 vs MD and Co-founder "
         "per 2025-11-10 and 2026-03-12)")
check(_rzw == [("Shashank Kumar, co-founder, quoted as CTO in August 2026 press", [13]), (_conf, [13]), (_conf, [8]),
               (_conf, [14]),
               ("Shashank Kumar, co-founder, quoted as CTO in August 2026 press (former CTO Murali Brahmadesam announced his "
                "exit on 2025-07-07)", [22]),
               ("Prabu Rambadran, Senior Vice President, Engineering since November 2025", [8])],
      "evidence: markers mid-parenthetical cite the parenthetical clause, ';' inside brackets splits only within them")
check(EV.split_facts("education not found (checked: Razorpay appointment release, Techcircle article [13]).") ==
      [("education not found (checked: Razorpay appointment release, Techcircle article)", [13])],
      "evidence: a marker just before ')' cites the clause with its parenthetical")
check(EV.split_facts("A (unclosed; x [1]; y [2]") == [("A (x)", [1]), ("A (y)", [2])],
      "evidence: an unclosed bracket is never cut at top level")
check(EV.split_facts("Office in U.S. Seattle [1]. Founded 2014 [2]") == [("Office in U.S. Seattle", [1]), ("Founded 2014", [2])],
      "evidence: sentence ends split facts, abbreviations like U.S. do not")
check(EV.split_facts("Harshil Mathur: IIT Roorkee, graduated 2013 [9]. Shashank Kumar: IIT Roorkee [10]")[1] ==
      ("Shashank Kumar: IIT Roorkee", [10]), "evidence: a new sentence starts a new fact")
check([EV.source_type(u, "acme.com") for u in (
    "https://www.sec.gov/Archives/x", "https://acme.com/about", "https://blog.acme.com/p", "https://www.wikidata.org/wiki/Q1",
    "https://en.wikipedia.org/wiki/Acme", "https://www.reuters.com/x", "https://search.gleif.org/#/record/X",
    "https://dns.google/resolve?name=acme.com&type=MX", "https://boards.greenhouse.io/acme", "https://github.com/acme",
    "https://www.reveliolabs.com/x", "https://someblog.example/x")] ==
      ["Filing", "Company site", "Company site", "Wikidata", "Wikipedia", "News", "Registry", "DNS", "Careers / ATS",
       "Code host", "Data provider", "Web"], "evidence: source types (filing, company site, wikidata, news, registry, DNS, ATS, ...)")
check(EV.EVIDENCE_COLUMNS == ["Domain", "Field", "Fact", "Source URL", "Source Type", "Checked On", "Age (days)", "Confidence"]
      and EV.DERIVED_COLUMNS == ["Evidence Count", "Freshness", "Key Field Confidence"], "evidence: exact header names")
evr = EV.evidence_for_row(row(), on=date(2026, 10, 20))
check(len(evr) == 8 and all(r[0] == "acme.com" and r[3].startswith("https://") for r in evr),
      "evidence: one row per cited fact, every row has a source URL")
check(not any(r[1] in ("India Office", IC, WARM, BEST, ST, "Confidence") for r in evr),
      "evidence: Not found cells, HumanCloud side, starter and bookkeeping fields are not facts")
check(all(r[5] == "2026-10-15" and r[6] == 5 for r in evr), "evidence: Checked On is the row's Last Checked; Age (days) computed")
fd = next(r for r in evr if r[1] == "Funding")
check(fd[2] == "$75M Series D, June 2026" and fd[4] == "Company site" and fd[7] == "High", "evidence: funding fact, type and confidence")
check(EV.derived(evr) == {"Evidence Count": 8, "Freshness": "Fresh", "Key Field Confidence": "High"},
      "evidence: derived columns for a fresh, company-sourced row")
weak = row(**{"Who to Contact": "Jane Doe, CTO [4]",
              "Sources": "[1] https://acme.com/about\n[2] https://www.wikidata.org/wiki/Q1\n[3] https://acme.com/press/series-d\n[4] https://someblog.example/x"})
check(EV.derived(EV.evidence_for_row(weak))["Key Field Confidence"] == "Low",
      "evidence: row confidence = weakest key field (a contact only on an unknown blog -> Low)")
check([EV.freshness(a) for a in (0, 30, 31, 180, 181, None)] == ["Fresh", "Fresh", "Ageing", "Ageing", "Stale", ""],
      "evidence: Freshness bands 30 / 180 days")
check(EV.derived(EV.evidence_for_row(row(**{"Last Checked": "2025-01-01"})))["Freshness"] == "Stale", "evidence: a 2025 check is Stale")
ff = EV.evidence_from_facts("acme.com", [
    {"field": "Funding", "value": "$75M Series D [3]; total $90M [1]", "checked": "2026-10-15",
     "sources": json.dumps(["https://acme.com/press/series-d", "https://acme.com/about"])},
    {"field": "HQ / US Office", "value": "San Francisco [1] [9]", "checked": "2026-10-15",
     "sources": json.dumps(["https://acme.com/about"])},
    {"field": "India Office", "value": "Not found (checked: site, Wikidata, 2 Exa searches)", "checked": "2026-10-15", "sources": "[]"},
    {"field": "Who to Contact", "value": "Jane [1]", "checked": "2026-10-15", "sources": "[]"}])
check([(r[1], r[2], r[3]) for r in ff] == [("Funding", "$75M Series D", "https://acme.com/press/series-d"),
                                           ("Funding", "total $90M", "https://acme.com/about"),
                                           ("HQ / US Office", "San Francisco", "https://acme.com/about")],
      "evidence: memory.db facts rebuild pairs markers with URLs; a mismatch falls back; no source = no row")

ws = load_workbook(wb)["Accounts"]
check([ws.cell(1, i).value for i in range(1, 20)] == list(excel_upsert.COLUMNS) and
      [ws.cell(1, i).value for i in range(20, 23)] == EV.DERIVED_COLUMNS, "excel: 19 headers unchanged, 3 derived columns appended")
check(isinstance(ws.cell(2, 20).value, int) and ws.cell(2, 20).value > 0 and ws.cell(2, 21).value == "Fresh"
      and ws.cell(2, 22).value in ("High", "Medium", "Low"),
      "excel: Evidence Count, Freshness, Key Field Confidence filled for the row")
wbe = load_workbook(wb)
ev_ws = wbe["Evidence"]
_typed_cols = set(excel_upsert.typed.TYPED_COLUMNS + excel_upsert.H1B_COLUMNS)
check([c.value for c in ev_ws[1]] == EV.EVIDENCE_COLUMNS
      and sum(r[1] not in _typed_cols for r in ev_ws.iter_rows(min_row=2, values_only=True)) == ws.cell(2, 20).value,
      "excel: Evidence sheet, one set of rows after 5 upserts of one company (rows replaced, not doubled)")
check(wbe["Contacts"].max_row == 3 and wbe["Company channels"].max_row == 7, "excel: Contacts and Company channels untouched by Evidence")
check(_ssn("Evidence") != "Evidence", "evidence: a list cannot be named Evidence")

mwb, mdb = TMP / "migrate.xlsx", TMP / "migrate.db"
import shutil as _sh  # noqa: E402
_sh.copy(wb, mwb)
mm = load_workbook(mwb)
mm.remove(mm["Evidence"])
for i in range(20, 23):
    mm["Accounts"].cell(1, i).value = None
    mm["Accounts"].cell(2, i).value = None
mm.save(mwb)
con_m = memory_connect(mdb)
memory.record_row(con_m, row(), "Accounts")
before = load_workbook(mwb).sheetnames
_patched, memory.connect = memory.connect, memory_connect
msg = excel_upsert.migrate_evidence(mwb, db=mdb, backups=TMP / "backups")
memory.connect = _patched
after = load_workbook(mwb)
bk = sorted((TMP / "backups").iterdir())
check(bk and (bk[-1] / "migrate.xlsx").exists() and (bk[-1] / "migrate.db").exists(), "migrate: workbook and memory.db backed up first")
check(after.sheetnames == before + ["Evidence"] and after["Accounts"].cell(2, 20).value == 8
      and sum(r[1] not in _typed_cols for r in after["Evidence"].iter_rows(min_row=2, values_only=True)) == 8,
      "migrate: Evidence sheet and derived columns built from memory.db; existing sheets kept in order")
check([after["Accounts"].cell(1, i).value for i in range(1, 20)] == list(excel_upsert.COLUMNS)
      and f"Evidence sheet: {after['Evidence'].max_row - 1} rows" in msg,
      "migrate: Accounts headers unchanged")

# Form D (mocked SEC)
FORM_D_XML = """<?xml version="1.0"?><edgarSubmission><primaryIssuer><cik>0001544522</cik><entityName>Acme Inc.</entityName>
<issuerAddress><city>SAN MATEO</city><stateOrCountry>CA</stateOrCountry></issuerAddress><jurisdictionOfInc>DELAWARE</jurisdictionOfInc>
</primaryIssuer><offeringData><industryGroup><industryGroupType>Other Technology</industryGroupType></industryGroup>
<typeOfFiling><newOrAmendment><isAmendment>false</isAmendment></newOrAmendment><dateOfFirstSale><value>2019-12-16</value></dateOfFirstSale></typeOfFiling>
<typesOfSecuritiesOffered><isEquityType>true</isEquityType></typesOfSecuritiesOffered>
<offeringSalesAmounts><totalOfferingAmount>250000054</totalOfferingAmount><totalAmountSold>250000054</totalAmountSold></offeringSalesAmounts>
<investors><totalNumberAlreadyInvested>3</totalNumberAlreadyInvested></investors></offeringData></edgarSubmission>"""
pd = L.parse_form_d(FORM_D_XML)
check(pd["sold_usd"] == 250000054 and pd["offering_usd"] == 250000054 and pd["first_sale"] == "2019-12-16"
      and pd["issuer"] == "Acme Inc." and pd["city"] == "SAN MATEO" and pd["investors"] == "3" and pd["amendment"] is False
      and pd["securities"] == ["isEquityType"], "formd: primary_doc.xml parsed (amounts, first sale, issuer, investors)")
check(L.parse_form_d(FORM_D_XML.replace("250000054</totalOfferingAmount>", "Indefinite</totalOfferingAmount>"))["offering_usd"] is None,
      "formd: an 'Indefinite' offering amount stays blank (never estimated)")
SUB = {"name": "Acme Inc.", "filings": {"recent": {"form": ["10-K", "D", "D/A", "D"],
       "accessionNumber": ["a-1", "0001-20-000002", "0001-21-000003", "0001-18-000001"],
       "filingDate": ["2024-01-01", "2020-07-08", "2021-01-02", "2018-08-08"]}}}
check([f["date"] for f in L.form_d_filings(SUB)] == ["2021-01-02", "2020-07-08", "2018-08-08"], "formd: Form D and D/A only, newest first")
check(L.rank_filers("postman", {"1": "tokenfoliopostmanfundi", "2": "postmaninc", "3": "postmanventuresfund"}) == ["2", "1", "3"],
      "formd: a filer named like the company ranks before funds and SPVs")
check(L._sec_cik("CIK 1544522", {})[0] == ["0001544522"], "formd: a CIK is used directly")
check("@" in L.SEC_UA and L.SEC_UA.startswith("RecruitDesk"), "formd: SEC User-Agent carries a contact address")
check(0.1 <= L.SEC_GAP < 1.0, "formd: SEC budget at most 10 requests a second")
t0 = _time.time()
L._polite("https://data.sec.gov/x"); L._polite("https://www.sec.gov/y"); L._polite("https://efts.sec.gov/z")
check(_time.time() - t0 >= 2 * L.SEC_GAP - 0.01, "formd: SEC hosts share one rate budget")


class _R:
    def __init__(self, status=200, body=None, text=""):
        self.status_code, self._b, self.text = status, body, text

    def json(self):
        return self._b


_calls: list = []


def _fake_get(url, headers=None, timeout=20):
    _calls.append((url, headers or {}))
    if "submissions" in url:
        return _R(body={"name": "Acme Inc.", "filings": {"recent": {"form": ["D"] * 8, "accessionNumber": [f"0001-2{i}-00000{i}" for i in range(8)],
                                                                      "filingDate": [f"201{i}-01-01" for i in range(8)]}}})
    return _R(text=FORM_D_XML)


_real_get = L.get
L.get = _fake_get
buf = io.StringIO()
with contextlib.redirect_stdout(buf):
    L.formd("0001544522")
L.get = _real_get
fo = buf.getvalue()
check("amount sold $250,000,054 of $250,000,054 offered · first sale 2019-12-16" in fo and "-index.htm" in fo,
      "formd: each filing printed with its amount, first sale date and filing URL")
check(len([u for u, _ in _calls if u.endswith("primary_doc.xml")]) == L.FORMD_MAX and "2 older Form D" in fo,
      "formd: at most FORMD_MAX documents read per lookup")
check(all(h.get("User-Agent") == L.SEC_UA for _, h in _calls), "formd: every SEC request sends the contact User-Agent")

# GLEIF and DNS (offline)
GL = {"data": [{"id": "254900ULRGCUEIK0K191", "attributes": {"lei": "254900ULRGCUEIK0K191", "registration": {"status": "ISSUED"},
       "entity": {"legalName": {"name": "ACME INC."}, "status": "ACTIVE",
                  "legalAddress": {"city": "LEWES", "region": "US-DE", "country": "US"},
                  "headquartersAddress": {"city": "San Mateo", "region": "US-CA", "country": "US"}}}}]}
g = L.gleif_records(GL)[0]
check(g["lei"] == "254900ULRGCUEIK0K191" and g["hq_city"] == "San Mateo" and g["hq_country"] == "US"
      and g["url"].endswith("254900ULRGCUEIK0K191"), "gleif: LEI, legal name, HQ city / country and record URL")


def _dns_reply(qtype, answers):
    q = L.dns_query_packet("acme.com", qtype)
    head = _struct.pack(">HHHHHH", 0x4D44, 0x8180, 1, len(answers), 0, 0)
    body = q[12:]
    for rd in answers:
        body += b"\xc0\x0c" + _struct.pack(">HHIH", qtype, 1, 300, len(rd)) + rd
    return head + body


mx_rd = _struct.pack(">H", 10) + b"\x05aspmx\x01l\x06google\x03com\x00"
mx_rd2 = _struct.pack(">H", 20) + b"\x04alt1\xc0\x0c"
mxr, tc = L.parse_dns_answer(_dns_reply(15, [mx_rd, mx_rd2]), 15)
check(mxr == [(10, "aspmx.l.google.com"), (20, "alt1.acme.com")] and tc is False, "dns: MX answers parsed, compressed names followed")
txt = b"google-site-verification=abc"
txr, _ = L.parse_dns_answer(_dns_reply(16, [bytes([len(txt)]) + txt, b"\x05v=spf" + bytes([28]) + b" include:_spf.salesforce.com"]), 16)
check(txr == ["google-site-verification=abc", "v=spf include:_spf.salesforce.com"], "dns: TXT strings parsed and joined")
sig = L.dns_signals(mxr, txr + ["atlassian-domain-verification=x", "MS=ms123"])
check(sig == {"email_provider": ["Google Workspace"], "tools": ["Atlassian", "Google", "Microsoft 365", "Salesforce"]},
      "dns: email provider and SaaS tokens recognised")
check(EV.source_type("https://dns.google/resolve?name=acme.com&type=TXT") == "DNS", "dns: the citable dns.google link is a DNS source")

# CRM export
check(X.split_name("Dr. Jane van der Berg (CTO)") == ("Jane", "van der Berg") and X.split_name("Cher") == ("Cher", ""),
      "crm: names split into first and last")
check(X.clean("Payments [1]; APIs [2].") == "Payments; APIs." and X.clean("Not found (checked: site, Wikidata)") == "",
      "crm: [n] stripped, placeholders left blank")
q = 'https://acme.com/team — "write to jane@acme.com"'
check(X.email_source("jane@acme.com", q) == "https://acme.com/team"
      and X.email_source("raj@acme.com", f"{C.APOLLO_LABEL} · Apollo person id {PID1}") == "Apollo (verified)"
      and X.email_source("raj@acme.com", "Source: Apollo · Apollo person id x") == ""
      and X.email_source("jane@acme.com", "https://acme.com/team") == ""
      and X.email_source(C.WITHHELD, "") == "" and X.email_source("Not published (checked: team page)", "") == "",
      "crm: only self-published (URL + quote) or Apollo-verified emails qualify")
check(X.phone_ok("+1 (855) 747 6767 (company line)", q) == "" and X.phone_ok("+14155550100", q) == "+14155550100",
      "crm: a company line is never a contact's phone")
people = [{"Domain": "acme.com", "Company": "Acme", "Name": "Jane Doe", "Title": "CTO [1]", "Email": "Jane@acme.com",
           "Email Source": q, "Phone": "+1 (855) 747 6767 (company line)", "Phone Source": q},
          {"Domain": "acme.com", "Company": "Acme", "Name": "Jane D.", "Title": "CTO", "Email": "jane@acme.com", "Email Source": q},
          {"Domain": "acme.com", "Company": "Acme", "Name": "Raj Patel", "Title": "VP", "Email": "Not published (checked: team page)"},
          {"Domain": "acme.com", "Company": "Acme", "Name": "Guess Person", "Title": "VP", "Email": "guess@acme.com", "Email Source": ""}]
comp = [{"Company": "Acme", "Domain": "acme.com", "What They Build (brief)": "Payment APIs for online shops [1]"}]
crow, prow, sk = X.build("hubspot", comp, people)
check(crow == [["Acme", "acme.com", "https://acme.com", "Payment APIs for online shops"]]
      and prow == [["Jane", "Doe", "jane@acme.com", "CTO", "", "Acme", "acme.com", "https://acme.com/team"]] and sk == 3,
      "crm: HubSpot rows keyed by domain and email; duplicates, unpublished and unsourced emails left out")
crow2, prow2, _ = X.build("salesforce", comp, people)
check(crow2 == [["Acme", "https://acme.com", "Payment APIs for online shops"]] and prow2[0][5:7] == ["Acme", "https://acme.com"],
      "crm: Salesforce Account / Contact shape")
outd = TMP / "exports"
with contextlib.redirect_stdout(io.StringIO()):
    rc = X.main(["--format", "hubspot", "--workbook", str(wb), "--out", str(outd)])
hc = list(_csv.reader((outd / "hubspot-companies.csv").open()))
hp = list(_csv.reader((outd / "hubspot-contacts.csv").open()))
check(rc == 0 and hc[0] == X.FORMATS["hubspot"]["companies"] and hc[1][:2] == ["Acme", "acme.com"]
      and hp[0] == X.FORMATS["hubspot"]["contacts"] and [r[2] for r in hp[1:]] == ["jane@acme.com"],
      "crm: export from the workbook writes both CSVs with the published contact only")
check("[1]" not in (outd / "hubspot-companies.csv").read_text() and "Not published" not in (outd / "hubspot-contacts.csv").read_text(),
      "crm: no citation markers or placeholders in CRM files")

# Typed company fields (round 2, scripts/typed.py)
import typed as T  # noqa: E402
check(T.candidates(T.T_FOUNDED, "Founded in 2014 by two engineers") == [2014]
      and T.candidates(T.T_FOUNDED, "a 2014 milestone") == [], "typed: Founded Year read only from founded / established wording")
check(T.candidates(T.T_EMP, "approx. 5,200 employees (Revelio, March 2026)") == [5200]
      and T.candidates(T.T_EMP, "5,000+ employees") == [] and T.candidates(T.T_EMP, "more than 300 employees") == [],
      "typed: Employees is a number; lower bounds such as 5,000+ or 'more than' are left blank")
check(T.candidates(T.T_EMP_ASOF, "approx. 18,600 employees (10-K, as of July 31, 2026)") == ["2026-07-31"]
      and T.candidates(T.T_EMP_ASOF, "approx. 80 employees per Tracxn as of Sep 30, 2026") == ["2026-09-30"]
      and T.candidates(T.T_EMP_ASOF, "approx. 5,200 employees (Revelio Labs data for March 2026)") == ["2026-03"]
      and T.candidates(T.T_EMP_ASOF, "approx. 17 employees (undated)") == [],
      "typed: Employees As Of is an ISO date at the source's precision, blank when undated")
fund = ("Series F, about $375 Mn, December 2021, co-led by TCV; valuation $7.5 Bn; about $740 Mn raised in total [4]. "
        "In June 2026 it confidentially filed draft IPO papers [9]; No newer round found (checked: 3 searches for Series G / pre-IPO)")
lr = T.last_round(fund)
check(lr and (lr["round"], lr["amount"], lr["date"]) == ("Series F", 375000000, "2021-12"),
      "typed: last round, amount and date from one clause; a draft IPO filing and a 'checked' note are not rounds")
check(T.last_round("Nasdaq IPO priced on 2021-09-21 at $36 per share, raising $1.03 billion at a $10.13 billion valuation")
      == {"round": "IPO", "amount": 1030000000, "date": "2021-09-21", "start": 7, "end": 10},
      "typed: an IPO's raise skips the per-share price and the valuation")
check(T.last_round("Seed $3M, Jul 18 2023 [1]; Series A $18M, Dec 4 2024 [2]")["round"] == "Series A",
      "typed: the latest round wins")
check(T.candidates(T.T_TOTAL, fund) == [740000000] and T.candidates(T.T_TOTAL, "more than $430M raised in total") == [],
      "typed: Total Raised read from 'in total / to date' wording, never a 'more than' bound")
check(T.candidates(T.T_HQ_CITY, "HQ: Bengaluru, India, 1st Floor [2]") == ["Bengaluru"]
      and T.candidates(T.T_HQ_COUNTRY, "800 West El Camino Real, Suite 180, Mountain View, CA 94040, listed as headquarters") == ["US"]
      and T.candidates(T.T_HQ_CITY, "San Francisco, CA: 1 Market St") == [],
      "typed: HQ City / Country only from a clause that says HQ or headquarters; country as ISO code")
check(T.candidates(T.T_INDIA_CITY, "Bangalore: Building no 310, Indiranagar") == ["Bengaluru"], "typed: India Office City normalised")
check(T.candidates(T.T_PUBLIC, "Public (INTU, Nasdaq)") == ["Public"] and T.candidates(T.T_TICKER, "Public (INTU, Nasdaq)") == ["INTU"]
      and T.candidates(T.T_PUBLIC, "$433M to date, privately held") == ["Private"] and T.candidates(T.T_PUBLIC, "Series D, $225M") == [],
      "typed: Public/Private and Ticker only when stated; a private round never implies Private")
tv, tev = T.derive_row(row())
check(tv == {T.T_FOUNDED: 2014, T.T_EMP: 300, T.T_ROUND: "Series D", T.T_ROUND_USD: 75000000, T.T_ROUND_DATE: "2026-06"}
      and not T.check_typed(tv, tev) and all(r[3].startswith("https://") for r in tev)
      and {r[1]: r[3] for r in tev}[T.T_FOUNDED] == "https://www.wikidata.org/wiki/Q1",
      "typed: derived from the row's cited clauses, each with an Evidence row citing its own [n]")
paren = row(**{"Company Brief (founded, size, customers, milestones)":
               "Founded 2014 [2]; approx. 300 employees per the 10-K (Wikipedia shows approx. 280 [3]) [1]"})
check({r[1]: r[3] for r in T.derive_row(paren)[1]}[T.T_EMP] == "https://acme.com/about",
      "typed: a parenthetical citation never lends its source to the main clause")
check(T.derive_row(row(**{"Company Brief (founded, size, customers, milestones)": "Founded 2014; approx. 300 employees"}))[0] == {
      T.T_ROUND: "Series D", T.T_ROUND_USD: 75000000, T.T_ROUND_DATE: "2026-06"}, "typed: an uncited clause fills nothing")
check(T.check_typed({T.T_EMP: 301}, tev) and T.check_typed({T.T_EMP: "300"}, tev)
      and T.check_typed({T.T_ROUND_DATE: "June 2026"}, tev), "typed: a value its Evidence does not state, a string number or a non-ISO date is refused")
check(not errs(row(**{"Founded Year": 2014, "Last Round": "Series D"}))
      and any("not stated by a cited clause" in x for x in errs(row(**{"Employees": 500})))
      and any("whole number" in x for x in errs(row(**{"Employees": "300"})))
      and any("ISO date" in x for x in errs(row(**{"Last Round Date": "06/2026"})))
      and any("DOL LCA index" in x for x in errs(row(**{"H-1B Top Roles": "Engineer"}))),
      "validator: explicit typed values typed and cited; H-1B columns never set in a rows file")

# H-1B LCA loader (scripts/h1b.py), fixture scripts/fixtures/lca_sample.csv
import shutil as _sh  # noqa: E402
FIX = Path(__file__).resolve().parent / "fixtures" / "lca_sample.csv"
check(h1b.norm_name("Razorpay Software Pvt. Ltd.") == "RAZORPAY SOFTWARE" and h1b.norm_name("Acme Labs, Inc.") == "ACME LABS"
      and h1b.norm_name("ACME LABS LLC") == "ACME LABS", "h1b: legal names normalised (suffixes and punctuation dropped)")
check({"AMAZON", "AWS", "AMAZON WEB SERVICES"} <= h1b.name_keys("AWS"), "h1b: aliases.json employer groups widen the match")
check(h1b.window(2026, 3) == (date(2025, 7, 1), date(2026, 6, 30)) and h1b.quarter_dates(2026, 1) == (date(2025, 10, 1), date(2025, 12, 31)),
      "h1b: the 4-quarter window follows the federal fiscal year")
html = ('<a href="https://www.dol.gov//media/LCA_Disclosure_Data_FY2026_Q3.xlsx">x</a>'
        '<a href="/sites/dolgov/files/ETA/oflc/pdfs/LCA_Disclosure_Data_FY2025_Q4.xlsx">y</a>'
        '<a href="/sites/dolgov/files/ETA/oflc/pdfs/LCA_Disclosure_Data_FY2025_Q3.xlsx">z</a>'
        '<a href="/sites/dolgov/files/ETA/oflc/pdfs/PERM_Disclosure_Data_FY2026_Q3.xlsx">p</a>')
check(h1b.list_files(html) == [(2026, 3, "https://www.dol.gov/media/LCA_Disclosure_Data_FY2026_Q3.xlsx"),
                               (2025, 4, "https://www.dol.gov/sites/dolgov/files/ETA/oflc/pdfs/LCA_Disclosure_Data_FY2025_Q4.xlsx"),
                               (2025, 3, "https://www.dol.gov/sites/dolgov/files/ETA/oflc/pdfs/LCA_Disclosure_Data_FY2025_Q3.xlsx")],
      "h1b: LCA files found on the OFLC page, newest first, other programmes ignored")
hdir = TMP / "h1b"
hdir.mkdir()
q3 = hdir / "LCA_Disclosure_Data_FY2026_Q3.csv"
_sh.copy(FIX, q3)
from openpyxl import Workbook as _WB  # noqa: E402
import csv as _csv2  # noqa: E402
xw = _WB(); xs = xw.active
for rr in _csv2.reader(FIX.open()):
    xs.append(rr)
q4 = hdir / "LCA_Disclosure_Data_FY2025_Q4.xlsx"
xw.save(q4)
hcon = h1b.connect(hdir / "h1b.sqlite")
n3 = h1b.index_file(hcon, q3, "https://www.dol.gov/media/LCA_Disclosure_Data_FY2026_Q3.xlsx")
n4 = h1b.index_file(hcon, q4, "https://www.dol.gov/x/LCA_Disclosure_Data_FY2025_Q4.xlsx")
check(n3 == 7 and n4 == 0, "h1b: H-1B cases only (E-3 dropped); a case number already indexed from another file is kept once (xlsx streamed)")
hs = h1b.summary(hcon, "Acme Labs")
check(hs["filings"] == 4 and hs["certified"] == 3 and hs["newest"] == "FY2026 Q3"
      and hs["by_file"] == [("FY2026 Q3", "https://www.dol.gov/media/LCA_Disclosure_Data_FY2026_Q3.xlsx", 4)],
      "h1b: filings in the last 4 quarters by legal name and DBA, cited per quarter file; out-of-window and other entities excluded")
check(hs["titles"][0] == ("Software Engineer", 2) and hs["worksites"][0] == ("San Francisco, CA", 2)
      and {w["unit"]: (w["min"], w["max"]) for w in hs["wages"]} == {"Year": (140000.0, 240000.0), "Hour": (80.0, 95.0)},
      "h1b: top titles, US worksites and wage ranges per pay unit")
check(h1b.summary(hcon, "AWS")["filings"] == 0 and h1b.summary(hcon, "AWS", ["Amazon.com Services LLC"])["filings"] == 1,
      "h1b: --alias adds the US legal entity name")
hcon.close()
import contextlib as _cl, io as _io  # noqa: E402
import lookup as _lk  # noqa: E402
_saved_idx, h1b.INDEX_PATH = h1b.INDEX_PATH, hdir / "h1b.sqlite"
out = _lk.h1b_report("Acme Labs", [], h1b)
check("filings in the last 4 quarters: 4 (3 certified)" in out and "LCA_Disclosure_Data_FY2026_Q3.xlsx" in out
      and "never a claim about any person" in out, "lookup h1b: prints the count, the cited quarter file and the signal note")
hv = h1b.signal("Acme Labs")
check(hv and hv[0] == {"H-1B Filings (4Q)": 4, "H-1B Top Roles": "Software Engineer (2); Data Scientist (1); Engineering Manager (1)"}
      and all(u.startswith("https://www.dol.gov/") for _, _, u in hv[1]), "h1b: the sheet signal and its Evidence rows")
hz = h1b.signal("Zeta Corp")
check(hz and hz[0]["H-1B Filings (4Q)"] == 0 and len(hz[1]) == 1 and hz[1][0][2].endswith("FY2026_Q3.xlsx"),
      "h1b: a zero count is cited to the newest file searched")
hrows = TMP / "acmelabs.rows.json"
hrows.write_text(json.dumps({"rows": [row(Company="Acme Labs", Domain="acmelabs.com")]}))
hwb = TMP / "h1b.xlsx"
with _cl.redirect_stdout(_io.StringIO()):
    sys.argv = ["excel_upsert.py", "--rows", str(hrows), "--workbook", str(hwb)]
    excel_upsert.main()
_src = {1: "https://acme.com/about"}
_sig = assemble.h1b_signal("Acme Labs", _src)
h1b.INDEX_PATH = _saved_idx
check(_sig.startswith("US H-1B signal (company-level, DOL LCA disclosure data): 4 filing(s) in the last 4 quarters, top roles Software Engineer (2)")
      and _sig.endswith("[2]") and _src[2].endswith("LCA_Disclosure_Data_FY2026_Q3.xlsx") and assemble.h1b_signal("Acme Labs", {}) == "",
      "assemble: Development Requirement gains the cited H-1B signal; nothing without an index")
hws = load_workbook(hwb)
hacc = hws["Accounts"]
hdr = [c.value for c in hacc[1]]
check(hdr == excel_upsert.COLUMNS + EV.DERIVED_COLUMNS + T.TYPED_COLUMNS + excel_upsert.H1B_COLUMNS,
      "excel: typed and H-1B columns appended after Key Field Confidence, existing headers unchanged")
vals = dict(zip(hdr, [c.value for c in hacc[2]]))
check(vals["Founded Year"] == 2014 and isinstance(vals["Employees"], int) and vals["Last Round Amount USD"] == 75000000
      and vals["Last Round Date"] == "2026-06" and vals["HQ City"] is None and vals["H-1B Filings (4Q)"] == 4,
      "excel: typed cells are numbers and ISO dates, uncited ones blank; H-1B count from the index")
hev = [r for r in hws["Evidence"].iter_rows(min_row=2, values_only=True)]
check({"Founded Year", "Employees", "Last Round", "H-1B Filings (4Q)", "H-1B Top Roles"} <= {r[1] for r in hev}
      and all(r[4] == "Filing" for r in hev if r[1].startswith("H-1B")), "excel: every typed and H-1B value has an Evidence row")

# ---- Companies master sheet (2026-10-18): one row per company across all list sheets ------------------
import entities as ENT  # noqa: E402
import shutil as _sh2  # noqa: E402

ewb = TMP / "entities.xlsx"
for lst, rr in (("Accounts", [row(**{"Last Checked": "2026-10-10"})]),
                ("Q4 India", [row(**{"Last Checked": "2026-10-15", "Funding": "$75M Series D, June 2026 [3]"}),
                              row(Company="Beta", Domain="beta.io", **{"Last Checked": "2026-10-12", "Sources": "[1] https://beta.io/about\n[2] https://www.wikidata.org/wiki/Q2\n[3] https://beta.io/press"})])):
    f = TMP / f"ent-{lst[:2]}.rows.json"
    f.write_text(json.dumps({"meta": {"list": lst}, "rows": rr}))
    sys.argv = ["excel_upsert.py", "--rows", str(f), "--workbook", str(ewb)]
    excel_upsert.main()
ex = load_workbook(ewb)
check(ex.sheetnames[:3] == ["Accounts", "Q4 India", "Companies"],
      "entities: Companies sits after the list sheets")
eh = [c.value for c in ex["Companies"][1]]
er = {r[0]: dict(zip(eh, r)) for r in ex["Companies"].iter_rows(min_row=2, values_only=True)}
check(set(er) == {"acme.com", "beta.io"}, "entities: one row per company across lists")
check(er["acme.com"]["Lists"] == "Accounts; Q4 India" and er["beta.io"]["Lists"] == "Q4 India", "entities: Lists names every sheet a company is in")
check(er["acme.com"]["Last Checked"] == "2026-10-15", "entities: the most recently checked row wins")
check(eh[:5] == ["Domain", "Company", "Lists", "People", "Channels"] and "Founded Year" in eh and "H-1B Filings (4Q)" in eh,
      "entities: front columns, then every list column including typed and H-1B")
check(ex["Accounts"].max_row == 2 and ex["Q4 India"].max_row == 3, "entities: list sheets are unchanged")
cpay2 = json.loads((TMP / "acme.rows.json").read_text())
cpay2["rows"][0]["_contacts"], cpay2["rows"][0]["_channels"] = r1["_contacts"], CH
(TMP / "ent-c.rows.json").write_text(json.dumps(cpay2))
sys.argv = ["excel_upsert.py", "--rows", str(TMP / "ent-c.rows.json"), "--workbook", str(ewb), "--list", "Accounts"]
excel_upsert.main()
cx = load_workbook(ewb)["Companies"]
ch = [c.value for c in cx[1]]
cr_ = {r[0]: dict(zip(ch, r)) for r in cx.iter_rows(min_row=2, values_only=True)}
check(cr_["acme.com"]["People"] == 2 and cr_["acme.com"]["Channels"] == 6 and cr_["beta.io"]["People"] == 0,
      "entities: People and Channels count the Contacts and Company channels rows per domain")
sys.argv = ["excel_upsert.py", "--rows", str(TMP / "ent-c.rows.json"), "--workbook", str(ewb), "--list", "Accounts"]
excel_upsert.main()
check(load_workbook(ewb)["Companies"].max_row == 3 and load_workbook(ewb).sheetnames.count("Companies") == 1, "entities: rebuilding is idempotent")
check(_ssn("Companies") != "Companies", "entities: a list cannot be named Companies")

old_wb = TMP / "old.xlsx"
_sh2.copy2(ewb, old_wb)
ow = load_workbook(old_wb)
del ow["Companies"]
ow.save(old_wb)
check(excel_upsert.migrate_entities(old_wb, db=TMP / "memory.db", backups=TMP / "bk-ent", dry_run=True).startswith("dry run: 2 companies")
      and "Companies" not in load_workbook(old_wb).sheetnames and not (TMP / "bk-ent").exists(), "entities: --dry-run counts and changes nothing")
msg_e = excel_upsert.migrate_entities(old_wb, db=TMP / "memory.db", backups=TMP / "bk-ent")
check(msg_e.startswith("migrated:") and "Companies" in load_workbook(old_wb).sheetnames
      and any((TMP / "bk-ent").rglob("old.xlsx")), "entities: migrate backs up the workbook, then builds the sheet")

ex2 = load_workbook(ewb)["Companies"]
check(ex2.protection.sheet and not ex2.protection.autoFilter and not ex2.protection.sort, "entities: the Companies sheet is marked read-only, sorting and filtering allowed")
fw = load_workbook(ewb)
del fw["Companies"]
old_list = fw.create_sheet("Companies")
old_list.append(["Role", "Pay"])
old_list.append(["Chef", 10])
try:
    ENT.refresh(fw)
    clobbered = True
except SystemExit:
    clobbered = False
check(not clobbered and fw["Companies"]["A2"].value == "Chef", "entities: an old sheet called Companies is refused, never deleted")

# ---- last_record.json: the last NEW company of an upsert (2026-10-19) -----------------------------
import hashlib as _hl  # noqa: E402
import last_record as LR  # noqa: E402

def check2(name, cond):
    check(cond, name)


lrdir = TMP / "lastrec"
lrdir.mkdir()
lrwb, lrfile = lrdir / "Account_Research.xlsx", lrdir / "last_record.json"
beta_src = "[1] https://beta.io/about\n[2] https://www.wikidata.org/wiki/Q2\n[3] https://beta.io/press"
gamma_src = "[1] https://gamma.dev/about\n[2] https://www.wikidata.org/wiki/Q3\n[3] https://gamma.dev/press"


def lr_upsert(rr, lst="Accounts"):
    f = lrdir / "rows.json"
    f.write_text(json.dumps({"meta": {"list": lst}, "rows": rr}))
    sys.argv = ["excel_upsert.py", "--rows", str(f), "--workbook", str(lrwb)]
    excel_upsert.main()


lr_upsert([row(), row(Company="Beta", Domain="beta.io", Sources=beta_src)])
doc = json.loads(lrfile.read_text(encoding="utf-8"))
check2("last record: the last NEW company of the batch", doc["agent"] == "account" and doc["record"]["company"] == "Beta" and doc["record"]["domain"] == "beta.io")
check2("last record: only listed fields, and no people, notes, sources or contact data",
      set(doc["record"]) <= set(LR.FIELDS) and "Jane Doe" not in lrfile.read_text(encoding="utf-8") and "https://beta.io/about" not in lrfile.read_text(encoding="utf-8"))
check2("last record: the key is stored hashed", doc["subject_hash"] == _hl.sha256(b"account:beta.io").hexdigest())
import time as _t  # noqa: E402
before = lrfile.read_text(encoding="utf-8")
_t.sleep(1.1)
lr_upsert([row(), row(Company="Beta", Domain="beta.io", Sources=beta_src)])
check2("last record: a rerun that only updates leaves the file alone", lrfile.read_text(encoding="utf-8") == before)
lr_upsert([row(Company="Gamma", Domain="gamma.dev", Sources=gamma_src)], "Q4 India")
check2("last record: a later new company replaces it", json.loads(lrfile.read_text(encoding="utf-8"))["record"]["company"] == "Gamma")
check2("last record: values that look like contact details are dropped", LR.clean({"company": "Acme", "hq": "Mail a@b.co", "funding": "Call +91 98765 43210", "confidence": "High"}) == {"company": "Acme", "confidence": "High"})

nrfile = lrdir / "new_rows.jsonl"
nrows = [json.loads(line) for line in nrfile.read_text(encoding="utf-8").splitlines()]
check2("history rows: one line per NEW company (Acme, Beta, Gamma), none for a rerun", len(nrows) == 3 and {"Beta", "Gamma"} <= {r["record"]["company"] for r in nrows})
check2("history rows: listed fields and a hashed key only, no people or sources",
      all(set(r["record"]) <= set(LR.FIELDS) and len(r["subject_hash"]) == 64 and r["agent"] == "account" for r in nrows) and "Jane Doe" not in nrfile.read_text(encoding="utf-8")
      and "https://beta.io/about" not in nrfile.read_text(encoding="utf-8"))
check2("history rows: contact-looking or empty records add nothing", LR.append_rows(lrdir, [("k", {"company": "Mail a@b.co"}), ("k2", {})]) == 0 and len(nrfile.read_text(encoding="utf-8").splitlines()) == len(nrows))

passed = sum(ok for ok, _ in RESULTS)
print(f"{passed}/{len(RESULTS)} passed")
sys.exit(0 if passed == len(RESULTS) else 1)
