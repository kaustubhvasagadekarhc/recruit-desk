"""The Companies master sheet: one row per company across every list sheet (2026-10-18, data-structure plan item 7).

A company can sit in several list sheets (Accounts, Q4 India, ...). 'Companies' holds it once, keyed by Domain, with
the row of the list that was checked most recently, the names of all its lists, and how many people and channels the
Contacts and Company channels sheets hold for it. It is a pure export like 'Evidence': lists stay the place where
rows are written, and every excel_upsert rebuilds the sheet. Nothing in a list sheet changes.

`excel_upsert.py --migrate-entities` backs up the workbook and memory.db into backups/<stamp>/ and builds the sheet.
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import progress
from common import COLUMNS, F_CHECKED, F_COMPANY, F_DOMAIN, bare_domain

COMPANIES_SHEET = "Companies"
FRONT = [F_DOMAIN, F_COMPANY, "Lists", "People", "Channels"]
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _is_list_sheet(ws) -> bool:
    return [ws.cell(row=1, column=i).value for i in range(1, len(COLUMNS) + 1)] == COLUMNS


def _domain_counts(wb, name: str) -> dict[str, int]:
    out: dict[str, int] = {}
    if name in wb.sheetnames:
        for r in wb[name].iter_rows(min_row=2, max_col=1, values_only=True):
            d = bare_domain(str(r[0] or ""))
            if d:
                out[d] = out.get(d, 0) + 1
    return out


def collect(wb) -> tuple[list[str], list[list]]:
    """(header, rows) of the Companies sheet, from the list sheets of wb."""
    best: dict[str, tuple[str, dict]] = {}
    lists: dict[str, list[str]] = {}
    extra: list[str] = []
    for ws in wb.worksheets:
        if ws.title == COMPANIES_SHEET or not _is_list_sheet(ws):
            continue
        header = [str(c.value or "") for c in ws[1]]
        extra += [h for h in header if h and h not in FRONT and h not in extra]
        for vals in ws.iter_rows(min_row=2, values_only=True):
            rec = {h: vals[i] if i < len(vals) else None for i, h in enumerate(header)}
            d = bare_domain(str(rec.get(F_DOMAIN) or ""))
            if not d:
                continue
            if ws.title not in lists.setdefault(d, []):
                lists[d].append(ws.title)
            stamp = str(rec.get(F_CHECKED) or "")
            if d not in best or stamp > best[d][0]:
                best[d] = (stamp, rec)
    people = _domain_counts(wb, "Contacts")
    chans = _domain_counts(wb, "Company channels")
    header = FRONT + extra
    rows = [[d, rec.get(F_COMPANY), "; ".join(lists[d]), people.get(d, 0), chans.get(d, 0), *[rec.get(h) for h in extra]]
            for d, (_, rec) in best.items()]
    rows.sort(key=lambda r: str(r[1] or "").lower())
    return header, rows


def refresh(wb) -> int:
    """Rebuild the Companies sheet in wb (not saved), right after the last list sheet. Returns the company count."""
    if COMPANIES_SHEET in wb.sheetnames and [wb[COMPANIES_SHEET].cell(1, i).value for i in range(1, len(FRONT) + 1)] != FRONT:
        progress.fail_exit(f"a sheet named '{COMPANIES_SHEET}' already exists and is not the master sheet (an old list called "
                           f"'{COMPANIES_SHEET}'?); rename it before the master can be built. Nothing was changed")
    header, rows = collect(wb)
    if COMPANIES_SHEET in wb.sheetnames:
        del wb[COMPANIES_SHEET]
    if not rows:
        return 0
    ws = wb.create_sheet(COMPANIES_SHEET)
    ws.append(header)
    for i, h in enumerate(header, 1):
        c = ws.cell(1, i)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = 22 if i <= 2 else 14 if h in FRONT or len(h) < 14 else 44
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "C2"
    dom = 1
    for r in rows:
        ws.append([v if v not in ("", None) else None for v in r])
        cell = ws.cell(ws.max_row, dom)
        cell.hyperlink = f"https://{r[0]}"
        cell.font = Font(color="0563C1", underline="single")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{max(ws.max_row, 2)}"
    # Read-only marker (no password; sorting and filtering stay allowed): the sheet is rebuilt on every upsert, so an edit would be lost.
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.sheet_properties.tabColor = "808080"
    last_list = max((i for i, w in enumerate(wb.worksheets) if _is_list_sheet(w)), default=len(wb.worksheets) - 2)
    wb.move_sheet(ws, offset=last_list + 1 - wb.sheetnames.index(COMPANIES_SHEET))
    return len(rows)
