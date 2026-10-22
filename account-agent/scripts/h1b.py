"""H-1B LCA signal from the DOL OFLC public disclosure files (bulk data, no API, no key).

  uv run python scripts/h1b.py fetch [--files 4]   # download the newest 4 LCA disclosure files to data/ and index them
  uv run python scripts/h1b.py status              # which files are indexed and the 4-quarter window
  uv run python scripts/lookup.py h1b "<Company>" [--alias "<legal name>"]   # filings, top roles, worksites, wages

The quarterly files are listed on the OFLC performance page. `fetch` reads that page, downloads the newest four
LCA_Disclosure_Data_FY<year>_Q<n>.xlsx files (streamed to disk, kept under data/, git-ignored) and streams each
sheet row by row into data/h1b.sqlite, keeping only H-1B cases and the columns the lookup needs. A case number seen
in two files is kept once (some fiscal-year files are cumulative).

The lookup counts cases whose decision date falls in the last 4 quarters: the four fiscal quarters ending with the
newest indexed file's quarter. Employers match on normalised legal names (suffixes such as Inc, LLC, Pvt Ltd dropped)
and through schema/aliases.json "employers". Each count cites the quarter file it came from.

An LCA is an employer's filing to hire for a role on an H-1B visa. It is a company-level hiring signal: it says
nothing about any person, and the files carry no nationality.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sqlite3
import sys
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path
from statistics import median

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import ROOT  # noqa: E402

DATA_DIR = ROOT / "data"
INDEX_PATH = DATA_DIR / "h1b.sqlite"
PERF_PAGE = "https://www.dol.gov/agencies/eta/foreign-labor/performance"
FILE_RE = re.compile(r"""href=["']([^"']*LCA_Disclosure_Data_FY(\d{4})_Q([1-4])\.xlsx)["']""", re.I)
MAX_FILE = 1_500_000_000  # bytes; the largest quarterly file so far is about 250 MB
H1B_CLASS = "H-1B"
COLS = {  # record layout name -> index column
    "CASE_NUMBER": "case_number", "CASE_STATUS": "status", "DECISION_DATE": "decision_date", "VISA_CLASS": "visa",
    "JOB_TITLE": "job_title", "SOC_TITLE": "soc_title", "EMPLOYER_NAME": "employer", "TRADE_NAME_DBA": "dba",
    "WORKSITE_CITY": "ws_city", "WORKSITE_STATE": "ws_state", "WAGE_RATE_OF_PAY_FROM": "wage_from",
    "WAGE_RATE_OF_PAY_TO": "wage_to", "WAGE_UNIT_OF_PAY": "wage_unit", "TOTAL_WORKER_POSITIONS": "positions",
}
SCHEMA = """
CREATE TABLE IF NOT EXISTS files (name TEXT PRIMARY KEY, url TEXT NOT NULL, fy INTEGER NOT NULL, q INTEGER NOT NULL,
                                  rows INTEGER NOT NULL, indexed_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS lca (case_number TEXT PRIMARY KEY, file TEXT NOT NULL, decision_date TEXT, status TEXT,
                                employer TEXT, employer_norm TEXT, dba_norm TEXT, job_title TEXT, soc_title TEXT,
                                ws_city TEXT, ws_state TEXT, wage_from REAL, wage_to REAL, wage_unit TEXT,
                                positions INTEGER);
CREATE INDEX IF NOT EXISTS lca_emp ON lca(employer_norm);
CREATE INDEX IF NOT EXISTS lca_dba ON lca(dba_norm);
"""
SUFFIXES = {"INC", "INCORPORATED", "LLC", "L L C", "LTD", "LIMITED", "CORP", "CORPORATION", "CO", "COMPANY", "PVT",
            "PRIVATE", "LLP", "LP", "PLC", "PC", "GMBH", "AG", "SA", "BV", "NV", "PTE", "PTY", "USA", "US", "THE"}


# ---- names ------------------------------------------------------------------------------------------------------

def norm_name(name: str) -> str:
    """'Razorpay Software Pvt. Ltd.' -> 'RAZORPAY SOFTWARE'; 'Postman, Inc.' -> 'POSTMAN'."""
    s = (name or "").upper().replace("&", " AND ")
    s = re.sub(r"\bD/?B/?A\b.*$", "", s)
    s = re.sub(r"[^A-Z0-9 ]+", " ", s)
    words = s.split()
    while words and words[-1] in SUFFIXES:
        words.pop()
    while words and words[0] == "THE":
        words.pop(0)
    return " ".join(words)


def name_keys(company: str, aliases: list[str] | None = None, alias_path: Path | None = None) -> set[str]:
    """Normalised names to match: the company, any --alias, and its group in schema/aliases.json employers."""
    names = [company, *(aliases or [])]
    try:
        groups = json.loads((alias_path or ROOT / "schema" / "aliases.json").read_text(encoding="utf-8")).get("employers", {})
    except (OSError, ValueError):
        groups = {}
    want = {norm_name(n) for n in names}
    for canon, alts in groups.items():
        group = [canon, *alts]
        if want & {norm_name(g) for g in group}:
            names += group
    return {k for k in (norm_name(n) for n in names) if k}


# ---- quarters ---------------------------------------------------------------------------------------------------

def quarter_dates(fy: int, q: int) -> tuple[date, date]:
    """US federal fiscal quarter: FY2026 Q1 = 2025-10-01..2025-12-31, Q4 = 2026-07-01..2026-09-30."""
    y, m = (fy - 1, 10) if q == 1 else (fy, 1 + 3 * (q - 2))
    start = date(y, m, 1)
    nm = date(y + (m + 3 > 12), (m + 3 - 1) % 12 + 1, 1)
    return start, nm - timedelta(days=1)


def window(fy: int, q: int) -> tuple[date, date]:
    """The 4 fiscal quarters ending with FY<fy> Q<q>."""
    end = quarter_dates(fy, q)[1]
    pfy, pq = (fy, q - 3) if q == 4 else (fy - 1, q + 1)
    return quarter_dates(pfy, pq)[0], end


def file_label(name: str) -> str:
    m = re.search(r"FY(\d{4})_Q([1-4])", name)
    return f"FY{m.group(1)} Q{m.group(2)}" if m else name


# ---- parsing ----------------------------------------------------------------------------------------------------

def _iso(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(s[:19] if "-" in s else s, fmt).date().isoformat()
        except ValueError:
            continue
    return ""


def _num(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^0-9.]", "", str(v or ""))
    try:
        return float(s) if s else None
    except ValueError:
        return None


def _rows(path: Path):
    """Header-keyed rows of an LCA file (.xlsx streamed with openpyxl read-only, or .csv for fixtures)."""
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as fh:
            yield from csv.DictReader(fh)
        return
    from openpyxl import load_workbook  # noqa: PLC0415
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        it = wb.worksheets[0].iter_rows(values_only=True)
        header = [str(h or "").strip().upper() for h in next(it)]
        for r in it:
            yield dict(zip(header, r))
    finally:
        wb.close()


def connect(path: Path | None = None) -> sqlite3.Connection:
    p = path or INDEX_PATH
    p.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(p)
    con.row_factory = sqlite3.Row
    con.executescript(SCHEMA)
    return con


def index_file(con: sqlite3.Connection, path: Path, url: str, batch: int = 5000) -> int:
    """Stream one disclosure file into the index (H-1B cases only). Returns the number of cases kept."""
    name = path.name
    fy_q = re.search(r"FY(\d{4})_Q([1-4])", name)
    fy, q = (int(fy_q.group(1)), int(fy_q.group(2))) if fy_q else (0, 0)
    con.execute("DELETE FROM lca WHERE file=?", (name,))
    buf, n = [], 0
    sql = ("INSERT OR IGNORE INTO lca(case_number,file,decision_date,status,employer,employer_norm,dba_norm,job_title,"
           "soc_title,ws_city,ws_state,wage_from,wage_to,wage_unit,positions) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)")
    for r in _rows(path):
        r = {str(k or "").strip().upper(): v for k, v in r.items()}
        if str(r.get("VISA_CLASS") or "").strip().upper() != H1B_CLASS or not r.get("CASE_NUMBER"):
            continue
        emp = str(r.get("EMPLOYER_NAME") or "").strip()
        pos = _num(r.get("TOTAL_WORKER_POSITIONS"))
        buf.append((str(r["CASE_NUMBER"]).strip(), name, _iso(r.get("DECISION_DATE")),
                    str(r.get("CASE_STATUS") or "").strip(), emp, norm_name(emp), norm_name(str(r.get("TRADE_NAME_DBA") or "")),
                    str(r.get("JOB_TITLE") or "").strip(), str(r.get("SOC_TITLE") or "").strip(),
                    str(r.get("WORKSITE_CITY") or "").strip().title(), str(r.get("WORKSITE_STATE") or "").strip().upper(),
                    _num(r.get("WAGE_RATE_OF_PAY_FROM")), _num(r.get("WAGE_RATE_OF_PAY_TO")),
                    str(r.get("WAGE_UNIT_OF_PAY") or "").strip().title(), int(pos) if pos else 1))
        if len(buf) >= batch:
            n += _flush(con, sql, buf)
    n += _flush(con, sql, buf)
    con.execute("INSERT OR REPLACE INTO files(name,url,fy,q,rows,indexed_at) VALUES(?,?,?,?,?,?)",
                (name, url, fy, q, n, date.today().isoformat()))
    con.commit()
    return n


def _flush(con, sql, buf) -> int:
    before = con.total_changes
    con.executemany(sql, buf)
    buf.clear()
    return con.total_changes - before


# ---- lookup -----------------------------------------------------------------------------------------------------

def latest_window(con: sqlite3.Connection) -> tuple[date, date, str] | None:
    r = con.execute("SELECT fy, q FROM files WHERE fy>0 ORDER BY fy DESC, q DESC LIMIT 1").fetchone()
    if not r:
        return None
    s, e = window(r["fy"], r["q"])
    return s, e, f"FY{r['fy']} Q{r['q']}"


def summary(con: sqlite3.Connection, company: str, aliases: list[str] | None = None, top: int = 5,
            alias_path: Path | None = None) -> dict | None:
    """Filings in the last 4 quarters for the company's employer names, or None when nothing is indexed."""
    w = latest_window(con)
    if not w:
        return None
    start, end, newest = w
    nu = con.execute("SELECT url FROM files WHERE fy>0 ORDER BY fy DESC, q DESC LIMIT 1").fetchone()[0]
    keys = sorted(name_keys(company, aliases, alias_path))
    if not keys:
        return {"filings": 0, "keys": [], "window": (start.isoformat(), end.isoformat()), "newest": newest, "newest_url": nu,
                "employers": [], "by_file": [], "titles": [], "worksites": [], "wages": [], "certified": 0}
    marks = ",".join("?" * len(keys))
    rows = con.execute(
        f"SELECT l.*, f.url FROM lca l JOIN files f ON f.name=l.file WHERE (l.employer_norm IN ({marks}) OR "
        f"l.dba_norm IN ({marks})) AND l.decision_date BETWEEN ? AND ?",
        (*keys, *keys, start.isoformat(), end.isoformat())).fetchall()
    by_file = Counter((r["file"], r["url"]) for r in rows)
    titles = Counter(re.sub(r"\s+", " ", r["job_title"]).strip().title() for r in rows if r["job_title"])
    sites = Counter(f"{r['ws_city']}, {r['ws_state']}" for r in rows if r["ws_city"])
    wages = []
    for unit in sorted({r["wage_unit"] for r in rows if r["wage_unit"]}):
        lo = [r["wage_from"] for r in rows if r["wage_unit"] == unit and r["wage_from"]]
        hi = [r["wage_to"] or r["wage_from"] for r in rows if r["wage_unit"] == unit and r["wage_from"]]
        if lo:
            wages.append({"unit": unit, "min": min(lo), "max": max(hi), "median_from": median(lo), "n": len(lo)})
    return {"filings": len(rows), "certified": sum(r["status"].lower() == "certified" for r in rows),
            "keys": keys, "window": (start.isoformat(), end.isoformat()), "newest": newest, "newest_url": nu,
            "employers": sorted({r["employer"] for r in rows}),
            "by_file": [(file_label(f), u, n) for (f, u), n in sorted(by_file.items(), key=lambda x: x[0][0], reverse=True)],
            "titles": titles.most_common(top), "worksites": sites.most_common(top), "wages": wages}


def top_roles_text(s: dict, n: int = 3) -> str:
    return "; ".join(f"{t} ({c})" for t, c in s["titles"][:n])


def print_summary(company: str, s: dict | None) -> None:
    if s is None:
        print("h1b: no LCA files indexed yet; run `uv run python scripts/h1b.py fetch` (downloads ~500 MB once)")
        return
    print(f"H-1B LCA filings for {company} (DOL OFLC disclosure data; window {s['window'][0]} to {s['window'][1]}, "
          f"newest file {s['newest']}; matched names: {', '.join(s['keys']) or '-'})")
    print(f"  filings in the last 4 quarters: {s['filings']} ({s['certified']} certified)")
    if not s["filings"]:
        print("  none matched (try --alias with the US legal entity name, e.g. from lookup.py edgar or gleif)")
    for label, url, n in s["by_file"]:
        print(f"  {label} file: {n} filing(s) · {url}")
    if s["employers"]:
        print(f"  employer names on the filings: {'; '.join(s['employers'][:5])}")
    if s["titles"]:
        print("  top job titles: " + "; ".join(f"{t} ({c})" for t, c in s["titles"]))
    if s["worksites"]:
        print("  US worksites: " + "; ".join(f"{w} ({c})" for w, c in s["worksites"]))
    for w in s["wages"]:
        print(f"  wages ({w['unit']}): ${w['min']:,.0f} to ${w['max']:,.0f}, median offered ${w['median_from']:,.0f} (n={w['n']})")
    print("  note: an LCA is a company's filing to hire for a role on an H-1B visa; a company-level signal, never a "
          "claim about any person (the files carry no nationality)")


def signal(company: str, path: Path | None = None) -> tuple[dict, list[list]] | None:
    """(H-1B column values, Evidence rows without the domain prefix) for excel_upsert, or None with no index."""
    p = path or INDEX_PATH
    if not p.exists():
        return None
    try:
        con = sqlite3.connect(f"file:{p}?mode=ro", uri=True, timeout=2)
        con.row_factory = sqlite3.Row
        try:
            s = summary(con, company)
        finally:
            con.close()
    except sqlite3.Error:  # an index being rebuilt by `fetch` is skipped, never waited on
        return None
    if s is None:
        return None
    values = {"H-1B Filings (4Q)": s["filings"], "H-1B Top Roles": top_roles_text(s)}
    ev = []
    for label, url, n in s["by_file"]:
        ev.append(["H-1B Filings (4Q)", f"{n} H-1B LCA filing(s) decided {s['window'][0]} to {s['window'][1]} in the "
                   f"{label} disclosure file ({'; '.join(s['employers'][:3])})", url])
    if not s["by_file"]:  # zero is a finding too: cite the newest file that was searched
        ev.append(["H-1B Filings (4Q)", f"0 H-1B LCA filings for {', '.join(s['keys'])} decided {s['window'][0]} to "
                   f"{s['window'][1]} ({s['newest']} and earlier disclosure files)", s["newest_url"]])
    if s["titles"] and s["by_file"]:
        ev.append(["H-1B Top Roles", f"top LCA job titles {s['window'][0]} to {s['window'][1]}: {values['H-1B Top Roles']}",
                   s["by_file"][0][1]])
    return values, ev


# ---- fetch ------------------------------------------------------------------------------------------------------

def list_files(html: str, base: str = "https://www.dol.gov") -> list[tuple[int, int, str]]:
    """(fy, q, absolute url) for every LCA disclosure file linked on the performance page, newest first."""
    seen: dict[tuple[int, int], str] = {}
    for href, fy, q in FILE_RE.findall(html):
        url = href if href.startswith("http") else base + href
        url = re.sub(r"(?<!:)//+", "/", url)
        seen.setdefault((int(fy), int(q)), url)
    return sorted(((fy, q, u) for (fy, q), u in seen.items()), reverse=True)


def download(session, url: str, dest: Path) -> int:
    part = dest.with_suffix(dest.suffix + ".part")
    n = 0
    with session.get(url, stream=True, timeout=120) as r:
        if r.status_code != 200:
            raise RuntimeError(f"{url}: HTTP {r.status_code}")
        with part.open("wb") as fh:
            for chunk in r.iter_content(1 << 20):
                n += len(chunk)
                if n > MAX_FILE:
                    raise RuntimeError(f"{url}: larger than {MAX_FILE} bytes")
                fh.write(chunk)
    part.replace(dest)
    return n


def fetch(files: int = 4, session=None, index: Path | None = None, data_dir: Path | None = None) -> str:
    from lookup import SEC_UA  # noqa: PLC0415 - same contact-carrying User-Agent as the SEC calls
    from netsafe import SafeSession  # noqa: PLC0415
    s = session or SafeSession()
    s.headers.update({"User-Agent": SEC_UA})
    dd = data_dir or DATA_DIR
    dd.mkdir(parents=True, exist_ok=True)
    page = s.get(PERF_PAGE, timeout=60)
    page.raise_for_status()
    want = list_files(page.text)[:files]
    if not want:
        return "no LCA disclosure files found on the OFLC performance page"
    con = connect(index or dd / "h1b.sqlite")
    done = {r["name"] for r in con.execute("SELECT name FROM files")}
    lines = []
    for fy, q, url in want:
        name = url.rsplit("/", 1)[-1]
        dest = dd / name
        if not dest.exists():
            print(f"downloading {name} …", flush=True)
            size = download(s, url, dest)
            lines.append(f"{name}: downloaded {size / 1e6:.0f} MB")
        if name not in done:
            print(f"indexing {name} (streamed row by row) …", flush=True)
            lines.append(f"{name}: {index_file(con, dest, url)} H-1B cases indexed")
        else:
            lines.append(f"{name}: already indexed")
    keep = {u.rsplit("/", 1)[-1] for _, _, u in want}
    for r in con.execute("SELECT name FROM files").fetchall():
        if r["name"] not in keep:
            con.execute("DELETE FROM lca WHERE file=?", (r["name"],))
            con.execute("DELETE FROM files WHERE name=?", (r["name"],))
            lines.append(f"{r['name']}: dropped from the index (older than the newest {files})")
    con.commit()
    w = latest_window(con)
    con.close()
    if w:
        lines.append(f"window: {w[0]} to {w[1]} (4 quarters ending {w[2]})")
    return "\n".join(lines)


def status(index: Path | None = None) -> str:
    p = index or INDEX_PATH
    if not p.exists():
        return "no LCA files indexed (run `uv run python scripts/h1b.py fetch`)"
    con = connect(p)
    out = [f"{r['name']}: {r['rows']} H-1B cases, indexed {r['indexed_at']} · {r['url']}"
           for r in con.execute("SELECT * FROM files ORDER BY fy DESC, q DESC")]
    w = latest_window(con)
    con.close()
    if w:
        out.append(f"window: {w[0]} to {w[1]} (4 quarters ending {w[2]})")
    return "\n".join(out) or "index is empty"


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = ap.add_subparsers(dest="cmd", required=True)
    p = sp.add_parser("fetch"); p.add_argument("--files", type=int, default=4)
    sp.add_parser("status")
    a = ap.parse_args(argv)
    print(fetch(a.files) if a.cmd == "fetch" else status())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
